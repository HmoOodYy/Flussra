"""Bounded, typed file data independent of presentation and payroll mappings."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import posixpath
import re
import zipfile
from contextlib import closing
from datetime import datetime, time, timedelta
from decimal import ROUND_FLOOR, Decimal, DecimalException, localcontext
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format, is_timedelta_format

from app.imports.contracts import CANONICAL_FILE_VERSION
from app.imports.errors import ImportFailure
from app.imports.files import (
    FileCell,
    _enforce_metadata_bound,
    _inspect_csv,
    _inspect_xlsx,
    _metadata_elements,
    _path_for_sheet,
    _selected_cell_metadata,
    _sheet_paths,
    _xml_events,
)
from app.imports.policy import DEFAULT_IMPORT_RESOURCE_POLICY, ImportResourcePolicy
from app.product_policy import DEFAULT_PRODUCT_POLICY, ProductPolicy

_REFERENCE = re.compile(r"^([A-Z]{1,3})([1-9][0-9]*)(?::([A-Z]{1,3})([1-9][0-9]*))?$")
_NUMBER = re.compile(r"^[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[Ee][+-]?[0-9]+)?$")
_NOTICE = "Some data was ignored because no column header was provided. Please review your file."


def _invalid(message: str = "Stored file data is malformed. Reset and upload the files again."):
    return ImportFailure("IMPORT_CANONICAL_DATA_INVALID", message, 409)


class _BoundedJson:
    def __init__(self, maximum: int):
        self.maximum = maximum
        self.data = bytearray()

    def write(self, text: str) -> None:
        encoded = text.encode("utf-8")
        if len(self.data) + len(encoded) > self.maximum:
            raise ImportFailure("IMPORT_CANONICAL_FILE_TOO_LARGE",
                                "Extracted file data exceeds the configured canonical-data limit.")
        self.data.extend(encoded)

    def item(self, value: Any) -> None:
        # A cell is text-bounded; only one cell is serialized at a time.
        self.write(json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False))

    def start(self, file_type: str, name: str, headers: list[dict]) -> None:
        self.write('{"format_version":')
        self.item(CANONICAL_FILE_VERSION)
        self.write(',"file_type":')
        self.item(file_type)
        self.write(',"unit_name":')
        self.item(name)
        self.write(',"unit_index":1,"column_count":')
        self.item(max(item["column"] for item in headers))
        self.write(',"header":[')
        for index, header in enumerate(headers):
            if index:
                self.write(",")
            self.item(header)
        self.write('],"rows":[')
        self.first_row = True

    def row(self, number: int, cells) -> None:
        if not self.first_row:
            self.write(",")
        self.first_row = False
        self.write('{"row":')
        self.item(number)
        self.write(',"cells":[')
        for index, cell in enumerate(cells):
            if index:
                self.write(",")
            self.item(cell)
        self.write("]}")

    def finish(self) -> bytes:
        self.write("]}")
        return bytes(self.data)


def _usable_cache(present: bool, kind: str, raw: str | None) -> bool:
    return present and (kind == "str" or bool(raw))


def _serial_value(raw: str, number_format: str, epoch: datetime) -> dict:
    if is_timedelta_format(number_format):
        # Serial days remain exact rather than being rounded through timedelta.
        return {"kind": "duration", "value": raw, "excel_serial": raw,
                "epoch": epoch.isoformat(), "encoding": "excel_serial"}
    try:
        serial = Decimal(raw)
        minimum = Decimal((datetime.min - epoch).days - 1)
        maximum = Decimal((datetime.max - epoch).days + 1)
        if not serial.is_finite() or not minimum < serial < maximum:
            raise ValueError("Excel date outside representable range")
        with localcontext() as context:
            context.prec = max(32, len(serial.as_tuple().digits) + 16)
            day = int(serial.to_integral_value(rounding=ROUND_FLOOR))
            # The ISO view truncates sub-microseconds; the exact serial is retained.
            # A near-midnight fraction cannot round into another day.
            micros = int((serial - day) * Decimal(86_400_000_000))
        if 0 <= serial < 1:
            value = (datetime(2000, 1, 1) + timedelta(microseconds=micros)).time().isoformat()
            kind = "time"
        else:
            adjustment = 1 if epoch.year == 1899 and 0 < serial < 60 else 0
            value = (epoch + timedelta(days=day + adjustment, microseconds=micros)).isoformat()
            kind = "date"
        return {"kind": kind, "value": value, "excel_serial": raw, "epoch": epoch.isoformat()}
    except (DecimalException, ValueError, OverflowError):
        return {"kind": "error", "value": "Invalid Excel date/time value",
                "excel_serial": raw, "epoch": epoch.isoformat()}


def _cell(column: int, cell, marker: dict, epoch: datetime) -> dict:
    formula = bool(marker.get("formula"))
    cache_type = marker.get("cache_type") or "n"
    if formula and cache_type not in {"n", "b", "s", "str", "d", "e"}:
        raise ImportFailure("IMPORT_FORMULA_CACHE_INVALID", "A formula has an unsupported cached-result type.")
    if cell.data_type not in {"n", "b", "s", "str", "inlineStr", "d", "e"}:
        raise ImportFailure("IMPORT_CELL_TYPE_UNSUPPORTED", "The workbook contains an unsupported cell type.")
    raw = marker.get("raw_cache_value")
    present = bool(marker.get("cache_present"))
    usable = _usable_cache(present, cache_type, raw)
    numeric = marker.get("raw_number")
    value = cell.value
    if formula and not usable:
        typed = {"kind": "unavailable"}
    elif formula and cache_type == "str" and raw in (None, ""):
        typed = {"kind": "text", "value": ""}
    elif value is None and numeric in (None, "") and not marker.get("date"):
        typed = {"kind": "blank"}
    elif marker.get("native_date"):
        native_kind = "time" if isinstance(value, time) else "duration" if isinstance(value, timedelta) else "date"
        typed = {"kind": native_kind, "value": marker["date"]}
    elif marker.get("error") or cell.data_type == "e" and numeric is None:
        typed = {"kind": "error", "value": str(value)}
    elif numeric not in (None, ""):
        typed = (_serial_value(numeric, cell.number_format, epoch)
                 if is_date_format(cell.number_format) else {"kind": "number", "value": numeric})
    elif isinstance(value, bool):
        typed = {"kind": "boolean", "value": value}
    elif value is None:
        typed = {"kind": "blank"}
    else:
        typed = {"kind": "text", "value": str(value)}
    result = {"column": column, "value": typed}
    if formula:
        result["formula"] = {"text": marker.get("formula_text") or "",
                             "attributes": dict(sorted(marker.get("formula_attributes", {}).items())),
                             "cache": {"present": present, "type": cache_type,
                                       "raw_value": raw, "usable": usable}}
    return result


def _headers(cells: list[dict], product: ProductPolicy) -> list[dict]:
    result = []
    for cell in cells:
        value = cell["value"]
        label = str(value.get("value", "")) if value["kind"] not in {"blank", "unavailable"} else ""
        if value["kind"] == "error" or "formula" in cell and not label.strip():
            raise ImportFailure("IMPORT_HEADER_INVALID",
                                "A header formula needs a usable saved result. Replace it with a column name.")
        if label.strip():
            result.append({"column": cell["column"], "label": label, "hidden": False,
                           **{key: item for key, item in cell.items() if key != "column"}})
    if not result:
        raise ImportFailure("IMPORT_HEADER_REQUIRED", "The first row must contain column headers.")
    if len(result) > product.max_import_combined_columns:
        raise ImportFailure("IMPORT_FILE_COLUMN_LIMIT_EXCEEDED",
                            "This file exceeds the configured headed-column allowance.")
    return result


def _bounds(reference: str) -> tuple[int, int, int, int]:
    match = _REFERENCE.fullmatch(reference)
    if match is None:
        raise ImportFailure("IMPORT_XLSX_INVALID", "The workbook contains an invalid table or merge range.")
    left, top, right, bottom = match.groups()
    def column(label):
        number = 0
        for character in label:
            number = number * 26 + ord(character) - 64
        return number
    bounds = column(left), int(top), column(right or left), int(bottom or top)
    if bounds[2] < bounds[0] or bounds[3] < bounds[1]:
        raise ImportFailure("IMPORT_XLSX_INVALID", "Workbook ranges must have increasing coordinates.")
    return bounds


def _structure(data: bytes, name: str, headers: list[dict], last_row: int,
               resource: ImportResourcePolicy) -> set[int]:
    headed = {item["column"] for item in headers}
    hidden = set()
    table_id = None
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        path = _path_for_sheet(_sheet_paths(archive, archive.namelist(), resource), name)
        with archive.open(path) as source:
            for event, element in _xml_events(source, max_depth=resource.max_xlsx_xml_depth):
                if event != "end":
                    continue
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "col" and element.attrib.get("hidden") in {"1", "true"}:
                    first, last = int(element.attrib["min"]), int(element.attrib["max"])
                    hidden.update(column for column in headed if first <= column <= last)
                elif tag == "mergeCell":
                    left, top, right, bottom = _bounds(element.attrib.get("ref", ""))
                    if top <= last_row and bottom >= 1 and any(left <= column <= right for column in headed):
                        raise ImportFailure("IMPORT_MERGED_CELLS_UNSUPPORTED",
                                            "Merged cells cannot intersect the headed import table.")
                elif tag == "tablePart":
                    if table_id is not None:
                        raise ImportFailure("IMPORT_MULTIPLE_TABLES_UNSUPPORTED",
                                            "A worksheet may contain only one table.")
                    table_id = next((value for key, value in element.attrib.items()
                                     if key.rsplit("}", 1)[-1] == "id"), "")
                    if not table_id:
                        raise ValueError("missing table relationship")
        if table_id is not None:
            rel_path = posixpath.join(posixpath.dirname(path), "_rels", posixpath.basename(path) + ".rels")
            targets = []
            for attrs in _metadata_elements(archive, rel_path, "Relationship", resource.max_xlsx_xml_depth):
                if attrs.get("Id") == table_id:
                    if attrs.get("TargetMode") == "External" or not attrs.get("Type", "").endswith("/table"):
                        raise ValueError("invalid table relationship")
                    targets.append(attrs.get("Target", ""))
            if len(targets) != 1:
                raise ValueError("ambiguous table relationship")
            target = targets[0]
            table_path = (posixpath.normpath(target.lstrip("/")) if target.startswith("/")
                          else posixpath.normpath(posixpath.join(posixpath.dirname(path), target)))
            if not table_path.startswith("xl/") or table_path not in archive.namelist():
                raise ValueError("unsafe table relationship")
            declarations = 0
            with archive.open(table_path) as source:
                for event, element in _xml_events(source, max_depth=resource.max_xlsx_xml_depth):
                    if event == "start" and element.tag.rsplit("}", 1)[-1] == "table":
                        declarations += 1
                        bounds = _bounds(element.attrib.get("ref", ""))
                        if bounds[1] != 1 or element.attrib.get("headerRowCount", "1") != "1":
                            raise ImportFailure("IMPORT_HEADER_ROW_INVALID", "An Excel table must use row 1 as its header.")
            if declarations != 1:
                raise ValueError("invalid table declaration")
    return hidden


def _metadata(meta: dict, headers: list[dict], ignored: set[int], last_row: int,
              formula_counts: dict, resource: ImportResourcePolicy) -> dict:
    labels = [""] * meta["column_count"]
    for header in headers:
        labels[header["column"] - 1] = header["label"]
    result = {key: value for key, value in meta.items() if key != "last_data_rows"}
    result.update(header_row_number=1, header_state="selected", header_cells=labels,
                  headed_column_count=len(headers), hidden_columns=[item["column"] for item in headers if item["hidden"]],
                  headerless_columns=sorted(ignored), notices=[_NOTICE] if ignored else [],
                  imported_row_count=max(0, last_row - 1), **formula_counts)
    _enforce_metadata_bound(result, resource.max_inspection_bytes)
    return result


def _csv(data: bytes, resource: ImportResourcePolicy, product: ProductPolicy, *, header_only: bool = False):
    meta = _inspect_csv(data, set(), 1, resource)
    reader = csv.reader(io.StringIO(data.decode(meta["encoding"]), newline=""),
                        delimiter=meta["delimiter"], strict=True)
    raw_headers = next(reader)
    headers = _headers([{"column": index, "value": {"kind": "text", "value": label}}
                        for index, label in enumerate(raw_headers, 1)], product)
    columns = {item["column"] for item in headers}
    if header_only:
        counts = {key: 0 for key in ("formula_count", "formula_cached_value_count",
                                     "formula_empty_cache_count", "formula_missing_cache_count")}
        return _metadata(meta, headers, set(), 1, counts, resource), b"", "CSV", "visible", 0
    writer = _BoundedJson(product.max_import_canonical_file_bytes)
    writer.start("csv", "CSV", headers)
    ignored = set()
    last_row = 1
    for number, row in enumerate(reader, 2):
        ignored.update(index for index, value in enumerate(row, 1) if index not in columns and value)
        if not any(column <= len(row) and row[column - 1] for column in columns):
            continue
        # Empty interior rows keep their coordinates without buffering their
        # values. An ignored-only or empty tail does not extend the headed table.
        for blank_row in range(last_row + 1, number):
            writer.row(blank_row, ({"column": column, "value": {"kind": "blank"}}
                                   for column in sorted(columns)))
        def cells():
            for column in sorted(columns):
                value = row[column - 1] if column <= len(row) else ""
                yield {"column": column, "value": ({"kind": "text", "value": value}
                                                   if value else {"kind": "blank"})}
        writer.row(number, cells())
        last_row = number
    counts = {key: 0 for key in ("formula_count", "formula_cached_value_count",
                                 "formula_empty_cache_count", "formula_missing_cache_count")}
    return _metadata(meta, headers, ignored, last_row, counts, resource), writer.finish(), "CSV", "visible", 0


def _xlsx(data: bytes, resource: ImportResourcePolicy, product: ProductPolicy, *, header_only: bool = False):
    scanned = _inspect_xlsx(data, set(), None, None, resource, include_all_units=True, strict_table=True)
    unit = scanned["source_units"][0]
    meta, name = unit["inspection"], unit["name"]
    if not meta["row_count"] or not meta["column_count"]:
        raise ImportFailure("IMPORT_HEADER_REQUIRED", "The first row must contain column headers.")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        worksheet = workbook[name]
        worksheet.reset_dimensions()
        values = worksheet.iter_rows(min_row=1, max_row=meta["row_count"], min_col=1,
                                     max_col=meta["column_count"], values_only=False)
        with closing(values), closing(_selected_cell_metadata(data, name, resource, None)) as markers:
            marker_row, row_markers = next(markers, (None, {}))
            first = next(values)
            header_markers = row_markers if marker_row == 1 else {}
            headers = _headers([_cell(index, cell, header_markers.get(index, {}), workbook.epoch)
                                for index, cell in enumerate(first, 1)], product)
            headed = {item["column"] for item in headers}
            last_row = max((meta["last_data_rows"].get(column, 1) for column in headed), default=1)
            hidden = _structure(data, name, headers, last_row, resource)
            for header in headers:
                header["hidden"] = header["column"] in hidden
            ignored = {column for column, row in meta["last_data_rows"].items() if column not in headed and row > 1}
            if header_only:
                counts = {key: 0 for key in ("formula_count", "formula_cached_value_count",
                                             "formula_empty_cache_count", "formula_missing_cache_count")}
                for header in headers:
                    if "formula" in header:
                        counts["formula_count"] += 1
                        cache = header["formula"]["cache"]
                        key = ("formula_missing_cache_count" if not cache["present"] else
                               "formula_empty_cache_count" if not cache["raw_value"] else
                               "formula_cached_value_count")
                        counts[key] += 1
                return (_metadata(meta, headers, ignored, 1, counts, resource), b"", name,
                        unit["visibility"], scanned["decompressed_bytes"])
            writer = _BoundedJson(product.max_import_canonical_file_bytes)
            writer.start("xlsx", name, headers)
            counts = {key: 0 for key in ("formula_count", "formula_cached_value_count",
                                         "formula_empty_cache_count", "formula_missing_cache_count")}
            def count_formula(item):
                if "formula" in item:
                    counts["formula_count"] += 1
                    cache = item["formula"]["cache"]
                    key = ("formula_missing_cache_count" if not cache["present"] else
                           "formula_empty_cache_count" if not cache["raw_value"] else "formula_cached_value_count")
                    counts[key] += 1
                return item
            for header in headers:
                count_formula(header)
            for number, row in enumerate(values, 2):
                while marker_row is not None and marker_row < number:
                    marker_row, row_markers = next(markers, (None, {}))
                current = row_markers if marker_row == number else {}
                if number > last_row:
                    break
                writer.row(number, (count_formula(_cell(column, row[column - 1], current.get(column, {}), workbook.epoch))
                                    for column in sorted(headed)))
            return _metadata(meta, headers, ignored, last_row, counts, resource), writer.finish(), name, unit["visibility"], scanned["decompressed_bytes"]
    finally:
        workbook.close()


def extract_canonical_file(data: bytes, filename: str, *,
                           resource_policy: ImportResourcePolicy = DEFAULT_IMPORT_RESOURCE_POLICY,
                           product_policy: ProductPolicy = DEFAULT_PRODUCT_POLICY) -> dict:
    """Validate the whole package, then serialize headed rows without a matrix."""
    if len(data) > resource_policy.max_file_bytes:
        raise ImportFailure("IMPORT_FILE_TOO_LARGE", "File exceeds its configured upload limit.")
    if not data:
        raise ImportFailure("IMPORT_FILE_EMPTY", "The uploaded file is empty.")
    if not isinstance(filename, str) or not filename.strip():
        raise ImportFailure("IMPORT_FILE_NAME_REQUIRED", "A file name is required.")
    suffix = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    try:
        if suffix == "csv":
            result = _csv(data, resource_policy, product_policy)
        elif suffix == "xlsx":
            result = _xlsx(data, resource_policy, product_policy)
        else:
            raise ImportFailure("IMPORT_FILE_TYPE_UNSUPPORTED", "Only .xlsx and .csv files are supported.")
    except ImportFailure:
        raise
    except Exception as exc:
        raise ImportFailure("IMPORT_FILE_INVALID", "The file's table or cell content is malformed.") from exc
    meta, payload, name, visibility, decompressed = result
    return {"file_type": suffix, "inspection": meta, "decompressed_bytes": decompressed,
            "source_units": [{"name": name, "index": 1, "visibility": visibility, "inspection": meta}],
            "canonical_payload": payload, "canonical_byte_count": len(payload),
            "canonical_sha256": hashlib.sha256(payload).hexdigest()}


def inspect_canonical_headers(data: bytes, filename: str, *,
                              resource_policy: ImportResourcePolicy = DEFAULT_IMPORT_RESOURCE_POLICY,
                              product_policy: ProductPolicy = DEFAULT_PRODUCT_POLICY) -> dict:
    """Run bounded safety checks and inspect only canonical headers, without serializing data rows."""
    if len(data) > resource_policy.max_file_bytes:
        raise ImportFailure("IMPORT_FILE_TOO_LARGE", "File exceeds its configured upload limit.")
    if not data:
        raise ImportFailure("IMPORT_FILE_EMPTY", "The uploaded file is empty.")
    if not isinstance(filename, str) or not filename.strip():
        raise ImportFailure("IMPORT_FILE_NAME_REQUIRED", "A file name is required.")
    suffix = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    try:
        if suffix == "csv":
            result = _csv(data, resource_policy, product_policy, header_only=True)
        elif suffix == "xlsx":
            result = _xlsx(data, resource_policy, product_policy, header_only=True)
        else:
            raise ImportFailure("IMPORT_FILE_TYPE_UNSUPPORTED", "Only .xlsx and .csv files are supported.")
    except ImportFailure:
        raise
    except Exception as exc:
        raise ImportFailure("IMPORT_FILE_INVALID", "The file's table or cell content is malformed.") from exc
    meta, _unused_payload, name, visibility, decompressed = result
    return {"file_type": suffix, "inspection": meta, "decompressed_bytes": decompressed,
            "source_units": [{"name": name, "index": 1, "visibility": visibility, "inspection": meta}]}


def _integer(value, minimum=1):
    return isinstance(value, int) and not isinstance(value, bool) and value >= minimum


def _unique_object(pairs):
    value = {}
    for key, item in pairs:
        if key in value:
            raise ValueError("duplicate JSON member")
        value[key] = item
    return value


def _constant(value):
    raise ValueError(value)


def _typed(value, resource):
    if not isinstance(value, dict):
        raise _invalid()
    kind = value.get("kind")
    if kind in {"blank", "unavailable"}:
        if set(value) != {"kind"}:
            raise _invalid()
        return
    if kind not in {"text", "number", "boolean", "error", "date", "time", "duration"}:
        raise _invalid()
    allowed = {"kind", "value", "excel_serial", "epoch", "encoding"}
    if set(value) - allowed or "value" not in value:
        raise _invalid()
    raw = value["value"]
    if kind == "boolean":
        if type(raw) is not bool:
            raise _invalid()
    elif not isinstance(raw, str) or len(raw.encode("utf-8")) > resource.max_cell_text_bytes:
        raise _invalid()
    if kind == "number" and not _NUMBER.fullmatch(raw):
        raise _invalid()
    if kind == "date":
        try:
            datetime.fromisoformat(raw)
        except ValueError as exc:
            raise _invalid() from exc
    if kind == "time":
        try:
            time.fromisoformat(raw)
        except ValueError as exc:
            raise _invalid() from exc
    if "excel_serial" in value:
        serial = value["excel_serial"]
        if (kind not in {"date", "time", "duration", "error"} or not isinstance(serial, str)
                or len(serial.encode("utf-8")) > resource.max_cell_text_bytes or not _NUMBER.fullmatch(serial)
                or value.get("epoch") not in {"1899-12-30T00:00:00", "1904-01-01T00:00:00"}):
            raise _invalid()
    elif "epoch" in value:
        raise _invalid()
    if "encoding" in value and (kind != "duration" or value["encoding"] != "excel_serial"):
        raise _invalid()


def _cell_schema(cell, column, resource, *, header=False):
    allowed = {"column", "value", "formula"} | ({"label", "hidden"} if header else set())
    if not isinstance(cell, dict) or set(cell) - allowed or cell.get("column") != column or not _integer(cell["column"]):
        raise _invalid()
    _typed(cell.get("value"), resource)
    if header:
        if (not isinstance(cell.get("label"), str) or not cell["label"].strip()
                or len(cell["label"].encode("utf-8")) > resource.max_cell_text_bytes
                or type(cell.get("hidden")) is not bool or cell["value"]["kind"] in {"blank", "unavailable", "error"}):
            raise _invalid()
        if cell["label"] != str(cell["value"]["value"]):
            raise _invalid()
    formula = cell.get("formula")
    if formula is not None:
        if not isinstance(formula, dict) or set(formula) != {"text", "attributes", "cache"}:
            raise _invalid()
        if not isinstance(formula["text"], str) or len(formula["text"].encode("utf-8")) > resource.max_cell_text_bytes:
            raise _invalid()
        attrs = formula["attributes"]
        if (not isinstance(attrs, dict) or any(not isinstance(key, str) or not isinstance(value, str)
                                              for key, value in attrs.items())
                or sum(len(key.encode("utf-8")) + len(value.encode("utf-8")) for key, value in attrs.items()) > resource.max_cell_text_bytes):
            raise _invalid()
        cache = formula["cache"]
        if (not isinstance(cache, dict) or set(cache) != {"present", "type", "raw_value", "usable"}
                or type(cache["present"]) is not bool or type(cache["usable"]) is not bool
                or cache["type"] not in {"n", "b", "s", "str", "d", "e"}
                or cache["raw_value"] is not None and not isinstance(cache["raw_value"], str)):
            raise _invalid()
        raw = cache["raw_value"]
        if raw is not None and len(raw.encode("utf-8")) > resource.max_cell_text_bytes:
            raise _invalid()
        if len(formula["text"].encode("utf-8")) + len((raw or "").encode("utf-8")) > resource.max_cell_text_bytes:
            raise _invalid()
        if (cache["usable"] != _usable_cache(cache["present"], cache["type"], raw)
                or (cell["value"]["kind"] == "unavailable") == cache["usable"]):
            raise _invalid()
    elif cell["value"]["kind"] == "unavailable" or "formula" in cell:
        raise _invalid()


def decode_canonical_file(payload: bytes, *, expected_bytes: int, expected_sha256: str,
                          format_version: int, maximum_bytes: int,
                          resource_policy: ImportResourcePolicy = DEFAULT_IMPORT_RESOURCE_POLICY,
                          maximum_columns: int | None = None) -> dict:
    """Validate ownership-independent bytes and their complete typed format."""
    if type(format_version) is not int or format_version != CANONICAL_FILE_VERSION:
        raise ImportFailure("IMPORT_CONTRACT_UNSUPPORTED", "Unsupported file-data version. Reset and upload again.", 409)
    if not isinstance(payload, bytes) or not _integer(expected_bytes) or len(payload) != expected_bytes:
        raise _invalid()
    if len(payload) > maximum_bytes:
        raise ImportFailure("IMPORT_CANONICAL_FILE_TOO_LARGE", "Stored file data exceeds the current canonical-data allowance.", 422)
    if hashlib.sha256(payload).hexdigest() != expected_sha256:
        raise _invalid("Stored file data failed its integrity check. Reset and upload again.")
    try:
        value = json.loads(payload.decode("utf-8"), object_pairs_hook=_unique_object, parse_constant=_constant)
        keys = {"format_version", "file_type", "unit_name", "unit_index", "column_count", "header", "rows"}
        if (not isinstance(value, dict) or set(value) != keys or type(value["format_version"]) is not int
                or value["format_version"] != CANONICAL_FILE_VERSION or value["file_type"] not in {"csv", "xlsx"}
                or not isinstance(value["unit_name"], str) or not value["unit_name"]
                or not _integer(value["unit_index"]) or value["unit_index"] != 1
                or not _integer(value["column_count"]) or value["column_count"] > resource_policy.max_columns_per_unit
                or not isinstance(value["header"], list) or not value["header"] or not isinstance(value["rows"], list)):
            raise _invalid()
        if maximum_columns is not None and len(value["header"]) > maximum_columns:
            raise ImportFailure("IMPORT_FILE_COLUMN_LIMIT_EXCEEDED", "Stored file data exceeds the current column allowance.", 422)
        columns = []
        for header in value["header"]:
            if not isinstance(header, dict) or not _integer(header.get("column")):
                raise _invalid()
            column = header["column"]
            if column > value["column_count"] or columns and column <= columns[-1]:
                raise _invalid()
            _cell_schema(header, column, resource_policy, header=True)
            if value["file_type"] == "csv" and (
                "formula" in header or header["hidden"] or header["value"]["kind"] != "text"
            ):
                raise _invalid()
            columns.append(column)
        if columns[-1] != value["column_count"]:
            raise _invalid()
        previous = 1
        for row in value["rows"]:
            if (not isinstance(row, dict) or set(row) != {"row", "cells"} or not _integer(row["row"], 2)
                    or row["row"] <= previous or row["row"] > resource_policy.max_rows_per_unit
                    or not isinstance(row["cells"], list) or len(row["cells"]) != len(columns)):
                raise _invalid()
            previous = row["row"]
            if previous * value["column_count"] > resource_policy.max_cells_per_unit:
                raise _invalid()
            for cell, column in zip(row["cells"], columns, strict=True):
                _cell_schema(cell, column, resource_policy)
                if value["file_type"] == "csv" and ("formula" in cell or cell["value"]["kind"] not in {"text", "blank"}):
                    raise _invalid()
        return value
    except ImportFailure:
        raise
    except (UnicodeError, ValueError, TypeError, KeyError, RecursionError, OverflowError) as exc:
        raise _invalid() from exc


def canonical_source_unit(value: dict, *, unit_name: str, date_order: str | None = None) -> dict:
    """Project admitted headed cells without ever reopening the original file."""
    if value["unit_name"] != unit_name:
        raise _invalid("Stored file data does not match its source worksheet.")
    rows = []
    nonempty = set()
    for record in value["rows"]:
        cells = {}
        for cell in record["cells"]:
            typed = cell["value"]
            kind = typed["kind"]
            raw = typed.get("value")
            if kind in {"blank", "unavailable"}:
                raw, kind = None, "text"
            if raw is not None or "formula" in cell:
                nonempty.add(cell["column"])
            cells[str(cell["column"])] = FileCell(raw, kind, "formula" in cell, typed["kind"] == "unavailable")
        rows.append({"source_row": record["row"], "cells": cells})
    headers = [""] * value["column_count"]
    for header in value["header"]:
        headers[header["column"] - 1] = header["label"]
        nonempty.add(header["column"])
    return {"unit_name": unit_name, "rows": rows, "nonempty_columns": sorted(nonempty),
            "inspection": {"header_row_number": 1, "header_state": "selected", "header_cells": headers,
                "row_count": max((row["source_row"] for row in rows), default=1),
                "column_count": value["column_count"], "cell_count": (len(rows) + 1) * len(value["header"]),
                "date_order": date_order}}
