"""Bounded CSV/XLSX inspection without mapping rows to payroll concepts."""

from __future__ import annotations

import csv
import io
import json
import posixpath
import re
import zipfile
from collections.abc import Iterable, Iterator, Mapping
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from typing import Any
from xml.etree.ElementTree import ParseError

from defusedxml import ElementTree as SafeET
from openpyxl import load_workbook

from app.imports.errors import ImportFailure
from app.imports.policy import DEFAULT_IMPORT_RESOURCE_POLICY, ImportResourcePolicy

MAX_FILE_BYTES = DEFAULT_IMPORT_RESOURCE_POLICY.max_file_bytes
MAX_XLSX_DECOMPRESSED_BYTES = DEFAULT_IMPORT_RESOURCE_POLICY.max_file_decompressed_bytes
MAX_XLSX_SHEETS = DEFAULT_IMPORT_RESOURCE_POLICY.max_worksheets_per_file
MAX_SELECTED_ROWS = DEFAULT_IMPORT_RESOURCE_POLICY.max_rows_per_unit
MAX_SELECTED_COLUMNS = DEFAULT_IMPORT_RESOURCE_POLICY.max_columns_per_unit
MAX_SELECTED_CELLS = DEFAULT_IMPORT_RESOURCE_POLICY.max_cells_per_unit
MAX_CELL_TEXT_BYTES = DEFAULT_IMPORT_RESOURCE_POLICY.max_cell_text_bytes
MAX_INSPECTION_JSON_BYTES = DEFAULT_IMPORT_RESOURCE_POLICY.max_inspection_bytes

# Internal safety ceiling: >400x the member count of the generated supported
# workbooks, including a million-cell sheet. See the design's measured evidence.
MAX_XLSX_MEMBERS = 4_096

# Ordinary OOXML parts have shallow fixed hierarchies, independent of row count.
# Bound every XML part (including ignored extensions) before openpyxl loads it.
MAX_XLSX_XML_DEPTH = 256

_BASE_HEADER_LABELS = {
    "drivernumber",
    "driver number",
    "employee key",
    "employeekey",
    "drivername",
    "driver name",
    "workdate",
    "work date",
    "date",
}
_CELL_REF = re.compile(r"^([A-Z]{1,3})([1-9][0-9]*)$")
_MAX_HEADER_SCAN_ROWS = 200


@dataclass(frozen=True, slots=True)
class FileCell(Mapping[str, Any]):
    value: Any
    kind: str
    formula: bool = False
    missing_cache: bool = False

    def __getitem__(self, key: str) -> Any:
        if key not in {"value", "kind", "formula", "missing_cache"}:
            raise KeyError(key)
        return getattr(self, key)

    def __iter__(self) -> Iterator[str]:
        return iter(("value", "kind", "formula", "missing_cache"))

    def __len__(self) -> int:
        return 4


def inspect_file(
    data: bytes,
    filename: str,
    *,
    sheet_name: str | None = None,
    header_row_number: int | None = None,
    date_order: str | None = None,
    header_labels: Iterable[str] | None = None,
    policy: ImportResourcePolicy = DEFAULT_IMPORT_RESOURCE_POLICY,
) -> dict[str, Any]:
    """Validate and inspect one accepted file, returning bounded metadata only."""
    if len(data) > policy.max_file_bytes:
        raise _failure("IMPORT_FILE_TOO_LARGE", "File exceeds its configured upload limit.")
    if not data:
        raise _failure("IMPORT_FILE_EMPTY", "The uploaded file is empty.")
    if not isinstance(filename, str) or not filename.strip():
        raise _failure("IMPORT_FILE_NAME_REQUIRED", "A file name is required.")
    suffix = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    normalized_labels = _header_labels(header_labels)
    normalized_date_order = _normalize_date_order(date_order)

    if suffix == "csv":
        if sheet_name not in (None, "CSV"):
            raise _failure("IMPORT_SHEET_NOT_FOUND", "CSV has one sheet named CSV.")
        result = _inspect_csv(data, normalized_labels, header_row_number, policy)
    elif suffix == "xlsx":
        result = _inspect_xlsx(data, normalized_labels, sheet_name, header_row_number, policy)
    elif suffix in {"xls", "xlsm"}:
        raise _failure("IMPORT_FILE_TYPE_UNSUPPORTED", "Only .xlsx and .csv files are supported.")
    else:
        raise _failure("IMPORT_FILE_TYPE_UNSUPPORTED", "Only .xlsx and .csv files are supported.")

    result["file_type"] = suffix
    result["date_order"] = normalized_date_order
    result["date_order_state"] = "selected" if normalized_date_order else "unresolved"
    _enforce_metadata_bound(result, policy.max_inspection_bytes)
    return result


def inspect_source_file(
    data: bytes,
    filename: str,
    policy: ImportResourcePolicy,
    *,
    header_labels: Iterable[str] | None = None,
) -> dict[str, Any]:
    """Inspect every source unit in one file without selecting or mapping it."""
    if len(data) > policy.max_file_bytes:
        raise _failure("IMPORT_FILE_TOO_LARGE", "File exceeds its configured upload limit.")
    if not data:
        raise _failure("IMPORT_FILE_EMPTY", "The uploaded file is empty.")
    if not isinstance(filename, str) or not filename.strip():
        raise _failure("IMPORT_FILE_NAME_REQUIRED", "A file name is required.")
    suffix = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if suffix == "csv":
        inspection = _inspect_csv(data, _header_labels(header_labels), None, policy)
        units = [{
            "name": "CSV",
            "index": 1,
            "visibility": "visible",
            "inspection": inspection,
        }]
        decompressed = 0
    elif suffix == "xlsx":
        result = _inspect_xlsx(data, _header_labels(header_labels), None, None,
                               policy, include_all_units=True)
        units = result["source_units"]
        decompressed = result["decompressed_bytes"]
    else:
        raise _failure("IMPORT_FILE_TYPE_UNSUPPORTED", "Only .xlsx and .csv files are supported.")
    if len(units) > policy.max_source_units:
        raise _failure("IMPORT_SOURCE_UNIT_LIMIT_EXCEEDED", "Import batch exceeds its configured source-unit limit.")
    result = {"file_type": suffix, "decompressed_bytes": decompressed, "source_units": units}
    _enforce_metadata_bound(result, policy.max_inspection_bytes)
    return result


def read_source_units(
    data: bytes,
    filename: str,
    unit_configs: list[dict[str, Any]],
    policy: ImportResourcePolicy,
) -> list[dict[str, Any]]:
    """Read selected units once per workbook and retain only requested columns."""
    if len(data) > policy.max_file_bytes:
        raise _failure("IMPORT_FILE_TOO_LARGE", "File exceeds its configured upload limit.")
    if not data:
        raise _failure("IMPORT_FILE_EMPTY", "The uploaded file is empty.")
    if not isinstance(filename, str) or not filename.strip():
        raise _failure("IMPORT_FILE_NAME_REQUIRED", "A file name is required.")
    if not unit_configs or len(unit_configs) > policy.max_source_units:
        raise _failure("IMPORT_SOURCE_UNIT_LIMIT_EXCEEDED", "Select a valid number of source units.")
    if any(not isinstance(config, dict) for config in unit_configs):
        raise _failure("IMPORT_SOURCE_UNIT_INVALID", "Source-unit selection is malformed.")
    names = [config.get("unit_name") for config in unit_configs]
    if any(not isinstance(name, str) or not name for name in names):
        raise _failure("IMPORT_SOURCE_UNIT_INVALID", "Source-unit names must be non-empty text.")
    if len(set(names)) != len(unit_configs):
        raise _failure("IMPORT_SOURCE_UNIT_INVALID", "Source-unit selections must be unique.")
    suffix = filename.rsplit(".", 1)[-1].casefold() if isinstance(filename, str) and "." in filename else ""
    if suffix == "csv":
        inspection = inspect_source_file(data, filename, policy)
        if len(unit_configs) != 1 or unit_configs[0].get("unit_name") != "CSV":
            raise _failure("IMPORT_SOURCE_UNIT_NOT_FOUND", "CSV contains one source unit named CSV.")
        unit = inspection["source_units"][0]
        return [_read_csv_source_unit(data, unit, unit_configs[0], policy)]
    if suffix != "xlsx":
        raise _failure("IMPORT_FILE_TYPE_UNSUPPORTED", "Only .xlsx and .csv files are supported.")
    if len(data) > policy.max_file_bytes:
        raise _failure("IMPORT_FILE_TOO_LARGE", "File exceeds its configured upload limit.")
    result = _inspect_xlsx(data, _header_labels(None), None, None, policy,
                           include_all_units=True, unit_configs=unit_configs)
    return result["read_units"]


def _validate_unit_config(config: dict[str, Any], inspection: dict[str, Any], policy: ImportResourcePolicy) -> None:
    header = config.get("header_row_number")
    if isinstance(header, bool) or not isinstance(header, int) or not 1 <= header <= max(inspection["row_count"], 1):
        raise _failure("IMPORT_HEADER_ROW_INVALID", "Selected header row is outside the source unit.")
    columns = config.get("columns")
    if not isinstance(columns, list) or not columns:
        raise _failure("IMPORT_SOURCE_COLUMNS_INVALID", "Select at least one source column.")
    if any(isinstance(column, bool) or not isinstance(column, int) or column < 1
           or column > policy.max_columns_per_unit for column in columns):
        raise _failure("IMPORT_SOURCE_COLUMNS_INVALID", "A selected source column is outside the configured limit.")
    if len(set(columns)) != len(columns):
        raise _failure("IMPORT_SOURCE_COLUMNS_INVALID", "Selected source columns must be unique.")
    if max(columns) > inspection["column_count"]:
        raise _failure("IMPORT_SOURCE_COLUMNS_INVALID", "A selected source column is outside the inspected unit.")
    order = config.get("date_order")
    if order is not None:
        _normalize_date_order(order)


def _read_csv_source_unit(data: bytes, unit: dict[str, Any], config: dict[str, Any],
                          policy: ImportResourcePolicy) -> dict[str, Any]:
    inspection = unit["inspection"]
    _validate_unit_config(config, inspection, policy)
    selected = set(config["columns"])
    rows = []
    nonempty: set[int] = set()
    header_cells: list[str] = []
    try:
        text = data.decode(inspection["encoding"], errors="strict")
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=inspection["delimiter"], strict=True)
        for number, values in enumerate(reader, start=1):
            if number <= config["header_row_number"]:
                if number == config["header_row_number"]:
                    header_cells = list(values[:policy.max_columns_per_unit])
                    nonempty.update(column for column, value in enumerate(values, start=1) if value.strip())
                continue
            nonempty.update(column for column, value in enumerate(values, start=1) if value.strip())
            cells = {}
            for column in selected:
                value = values[column - 1] if column <= len(values) else ""
                if value.strip():
                    nonempty.add(column)
                cells[str(column)] = FileCell(value, "text")
            rows.append({"source_row": number, "cells": cells})
    except (UnicodeError, csv.Error) as exc:
        raise _failure("IMPORT_CSV_MALFORMED", "CSV structure is malformed.") from exc
    accepted_inspection = {**inspection, "header_row_number": config["header_row_number"],
                           "header_state": "selected", "header_cells": header_cells,
                           "date_order": _normalize_date_order(config.get("date_order"))}
    return {"unit_name": unit["name"], "inspection": accepted_inspection,
            "rows": rows, "nonempty_columns": sorted(nonempty)}


def _read_xlsx_source_unit(worksheet, unit: dict[str, Any], config: dict[str, Any],
                           markers, policy: ImportResourcePolicy) -> dict[str, Any]:
    inspection = unit["inspection"]
    selected = set(config["columns"])
    header_row = config["header_row_number"]
    columns = inspection["column_count"]
    rows_count = inspection["row_count"]
    rows = []
    nonempty: set[int] = set()
    header_cells: list[str] = []
    marker_iter = iter(markers)
    marker_row, row_markers = next(marker_iter, (None, {}))
    values_iter = worksheet.iter_rows(min_row=1, max_row=rows_count, min_col=1,
                                      max_col=columns, values_only=False)
    for number, cells in enumerate(values_iter, start=1):
        while marker_row is not None and marker_row < number:
            marker_row, row_markers = next(marker_iter, (None, {}))
        row_markers = row_markers if marker_row == number else {}
        values = {offset + 1: cell for offset, cell in enumerate(cells)}
        if number <= header_row:
            if number == header_row:
                header_cells = [_display_cell(cell.value) for cell in cells[:policy.max_columns_per_unit]]
                nonempty.update(column for column, cell in values.items()
                                if _display_cell(cell.value).strip() or row_markers.get(column, {}).get("formula"))
            continue
        nonempty.update(column for column, cell in values.items()
                        if _display_cell(cell.value).strip() or row_markers.get(column, {}).get("formula"))
        output_cells = {}
        for column in selected:
            cell = values.get(column)
            value = cell.value if cell is not None else None
            marker = row_markers.get(column, {})
            if marker.get("native_date"):
                value = marker.get("date")
            formula = bool(marker.get("formula"))
            is_error = bool(marker.get("error")) or bool(cell is not None and cell.data_type == "e")
            if value is not None and _display_cell(value).strip() or formula:
                nonempty.add(column)
            kind = ("error" if is_error else "date" if marker.get("native_date") or isinstance(value, (date, datetime))
                    else "time" if isinstance(value, time) else "duration" if isinstance(value, timedelta)
                    else "boolean" if isinstance(value, bool) else "number" if isinstance(value, (int, float))
                    else "text")
            rendered = value.isoformat() if isinstance(value, (date, datetime, time)) else (
                str(value) if kind in {"number", "duration"} and value is not None else value)
            if kind == "number" and marker.get("raw_number") is not None:
                # openpyxl's binary float can round away an invalid fractional
                # quantity before Decimal validation. Keep numeric XML lexemes;
                # styled dates/times/durations still use workbook interpretation.
                rendered = marker["raw_number"]
            output_cells[str(column)] = FileCell(
                rendered, kind, formula, bool(marker.get("missing_cache"))
            )
        rows.append({"source_row": number, "cells": output_cells})
    # Drain metadata so malformed trailing row/cell coordinates cannot escape
    # validation when openpyxl stops at the verified rectangle.
    for _remaining in marker_iter:
        pass
    accepted_inspection = {**inspection, "header_row_number": header_row,
                           "header_state": "selected", "header_cells": header_cells,
                           "date_order": _normalize_date_order(config.get("date_order"))}
    return {"unit_name": unit["name"], "inspection": accepted_inspection,
            "rows": rows, "nonempty_columns": sorted(nonempty)}


def _selected_cell_metadata(
    archive_path_data: bytes,
    sheet_name: str,
    policy: ImportResourcePolicy,
    selected_columns: set[int] | None,
):
    """Yield one row of formula/cache/date/exact numeric metadata at a time."""
    try:
        with zipfile.ZipFile(io.BytesIO(archive_path_data), "r") as archive:
            paths = _sheet_paths(archive, archive.namelist(), policy)
            path = _path_for_sheet(paths, sheet_name)
            with archive.open(path) as source:
                row_number = 0
                previous_row = 0
                row_markers: dict[int, dict[str, Any]] = {}
                column = 0
                cell: dict[str, Any] = {}
                last_column = 0
                for event, element in _xml_events(source, max_depth=policy.max_xlsx_xml_depth):
                    tag = _local_name(element.tag)
                    if event == "start" and tag == "row":
                        text_row = element.attrib.get("r")
                        row_number = int(text_row) if text_row and text_row.isdigit() else previous_row + 1
                        if row_number <= previous_row:
                            raise _failure("IMPORT_XLSX_INVALID", "Workbook rows must have unique increasing coordinates.")
                        row_markers = {}
                        last_column = 0
                    elif event == "start" and tag == "c":
                        match = _CELL_REF.fullmatch(element.attrib.get("r", ""))
                        if match is None or int(match.group(2)) != row_number:
                            raise _failure("IMPORT_XLSX_INVALID", "Workbook row and cell coordinates disagree.")
                        column = _column_number(match.group(1))
                        if column <= last_column:
                            raise _failure("IMPORT_XLSX_INVALID", "Workbook contains duplicate or unordered cell coordinates.")
                        last_column = column
                        cell = {"formula": False, "date": None, "native_date": element.attrib.get("t") == "d",
                                "error": element.attrib.get("t") == "e", "value_seen": False,
                                "numeric": element.attrib.get("t", "n") == "n"}
                    elif event == "start" and tag == "f":
                        cell["formula"] = True
                    elif event == "start" and tag == "v":
                        cell["value_seen"] = True
                    elif event == "end" and tag == "v" and cell.get("native_date"):
                        cell["date"] = element.text
                    elif event == "end" and tag == "v":
                        cell["raw_value"] = element.text
                    elif event == "end" and tag == "c":
                        if (selected_columns is None or column in selected_columns) and (
                            cell.get("formula") or cell.get("native_date") or cell.get("error")
                            or cell.get("numeric") and cell.get("value_seen")
                        ):
                            row_markers[column] = {
                                "formula": cell.get("formula", False),
                                "missing_cache": bool(cell.get("formula") and not cell.get("raw_value")),
                                "native_date": cell.get("native_date", False),
                                "date": cell.get("date"),
                                "error": cell.get("error", False),
                                "raw_number": cell.get("raw_value") if cell.get("numeric") else None,
                            }
                        cell = {}
                    elif event == "end" and tag == "row":
                        yield row_number, row_markers
                        previous_row = row_number
    except ImportFailure:
        raise
    except (ParseError, OSError, KeyError, zipfile.BadZipFile, ValueError) as exc:
        raise _failure("IMPORT_XLSX_INVALID", "Workbook worksheet XML is malformed.") from exc


def _inspect_csv(
    data: bytes,
    labels: set[str],
    selected_header_row: int | None,
    policy: ImportResourcePolicy = DEFAULT_IMPORT_RESOURCE_POLICY,
) -> dict[str, Any]:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise _failure("IMPORT_CSV_ENCODING_UNSUPPORTED", "UTF-16 CSV files are not supported.")
    if (re.search(rb"[\x00-\x08\x0b\x0c\x0e-\x1f]", data)
            or data.startswith((b"PK\x03\x04", b"PK\x05\x06", b"%PDF-", b"\xd0\xcf\x11\xe0", b"GIF87a", b"GIF89a"))):
        raise _failure("IMPORT_CSV_BINARY_CONTENT", "The file is not valid text CSV.")
    try:
        text = data.decode("utf-8-sig", errors="strict")
        encoding = "utf-8-sig" if data.startswith(b"\xef\xbb\xbf") else "utf-8"
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1252", errors="strict")
            encoding = "cp1252"
        except UnicodeDecodeError as exc:
            raise _failure("IMPORT_CSV_ENCODING_INVALID", "CSV encoding is not supported.") from exc

    delimiter = _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    max_columns = 0
    rows_seen = 0
    best_row_num = 0
    best_score = 0
    first_text_row: tuple[int, list[str]] | None = None
    detected_header: list[str] = []
    try:
        for row_num, row in enumerate(reader, start=1):
            rows_seen = row_num
            if rows_seen > policy.max_rows_per_unit:
                raise _failure("IMPORT_ROW_LIMIT_EXCEEDED", "CSV exceeds the configured number of rows.")
            row = [cell for cell in row]
            max_columns = max(max_columns, len(row))
            if max_columns > policy.max_columns_per_unit:
                raise _failure("IMPORT_COLUMN_LIMIT_EXCEEDED", "CSV exceeds the configured number of columns.")
            if rows_seen * max_columns > policy.max_cells_per_unit:
                raise _failure("IMPORT_CELL_LIMIT_EXCEEDED", "CSV exceeds the configured rectangular cell limit.")
            _check_text_values(row, policy.max_cell_text_bytes)
            text_values = [cell for cell in row if cell.strip()]
            if first_text_row is None and text_values and len(text_values) * 2 >= max(1, len(row)):
                first_text_row = (row_num, row)
            score = sum(_normalize_header(value) in labels for value in row)
            if score > best_score:
                best_score, best_row_num, detected_header = score, row_num, row
    except csv.Error as exc:
        raise _failure("IMPORT_CSV_MALFORMED", "CSV structure is malformed.") from exc

    if not rows_seen:
        raise _failure("IMPORT_CSV_EMPTY", "CSV contains no rows.")
    if max_columns == 0:
        max_columns = 1
    cell_count = rows_seen * max_columns
    if cell_count > policy.max_cells_per_unit:
        raise _failure("IMPORT_CELL_LIMIT_EXCEEDED", "CSV exceeds the configured rectangular cell limit.")

    if selected_header_row is not None:
        if not isinstance(selected_header_row, int) or not 1 <= selected_header_row <= rows_seen:
            raise _failure("IMPORT_HEADER_ROW_INVALID", "Header row must be within the CSV.")
        header_row_number = selected_header_row
        header_cells = _csv_row_at(data, encoding, delimiter, selected_header_row)
        header_state = "selected"
    elif best_score:
        header_row_number = best_row_num
        header_cells = detected_header
        header_state = "detected"
    elif first_text_row is not None:
        header_row_number, header_cells = first_text_row
        header_state = "detected"
    else:
        header_row_number, header_cells, header_state = 1, [], "detected"

    return {
        "sheets": [{"name": "CSV", "index": 1, "visibility": "visible"}],
        "selected_sheet": "CSV",
        "selected_sheet_index": 1,
        "header_row_number": header_row_number,
        "header_state": header_state,
        "header_cells": header_cells[:policy.max_columns_per_unit],
        "row_count": rows_seen,
        "column_count": max_columns,
        "cell_count": cell_count,
        "encoding": encoding,
        "delimiter": delimiter,
        "formula_count": 0,
        "formula_cached_value_count": 0,
        "formula_empty_cache_count": 0,
        "formula_missing_cache_count": 0,
    }


def _inspect_worksheet_header(worksheet, rows: int, columns: int, labels: set[str],
                              policy: ImportResourcePolicy) -> dict[str, Any]:
    scan_rows = min(rows, _MAX_HEADER_SCAN_ROWS)
    scan_columns = min(columns, policy.max_columns_per_unit)
    best_header: tuple[int, int, list[str]] | None = None
    first_text_row: tuple[int, list[str]] | None = None
    if scan_rows and scan_columns:
        iterator = worksheet.iter_rows(min_row=1, max_row=scan_rows, min_col=1,
                                       max_col=scan_columns, values_only=True)
        for row_number, values in enumerate(iterator, start=1):
            display = [_display_cell(value) for value in values]
            text_values = [value for value in values if isinstance(value, str) and value.strip()]
            if first_text_row is None and text_values and len(text_values) * 2 >= max(1, len(display)):
                first_text_row = (row_number, display)
            score = sum(_normalize_header(value) in labels for value in display)
            if score and (best_header is None or score > best_header[0]):
                best_header = (score, row_number, display)
    if best_header is not None:
        _score, header_row, header_cells = best_header
    elif first_text_row is not None:
        header_row, header_cells = first_text_row
    else:
        header_row, header_cells = 1, []
    return {"header_row_number": header_row, "header_state": "detected",
            "header_cells": header_cells[:policy.max_columns_per_unit]}


def _csv_row_at(data: bytes, encoding: str, delimiter: str, row_number: int) -> list[str]:
    text = data.decode(encoding, errors="strict")
    try:
        for index, row in enumerate(
            csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True), start=1
        ):
            if index == row_number:
                return row
    except csv.Error as exc:
        raise _failure("IMPORT_CSV_MALFORMED", "CSV structure is malformed.") from exc
    return []


def _sniff_delimiter(text: str) -> str:
    sample = text[:64_000]
    if not any(mark in sample for mark in (",", ";", "\t")):
        return ","
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        counts = {candidate: sample.count(candidate) for candidate in (",", ";", "\t")}
        delimiter, count = max(counts.items(), key=lambda item: item[1])
        if count == 0:
            return ","
        # A single consistent delimiter is a safe fallback when Sniffer has
        # insufficient sample rows to choose among the supported separators.
        return delimiter


def _inspect_xlsx(
    data: bytes,
    labels: set[str],
    selected_sheet: str | None,
    selected_header_row: int | None,
    policy: ImportResourcePolicy = DEFAULT_IMPORT_RESOURCE_POLICY,
    *,
    include_all_units: bool = False,
    unit_configs: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data), "r")
    except (zipfile.BadZipFile, OSError) as exc:
        raise _failure("IMPORT_XLSX_INVALID", "The .xlsx file is not a valid ZIP workbook.") from exc

    with archive:
        infos = archive.infolist()
        names = [info.filename for info in infos]
        if len(infos) > policy.max_xlsx_members:
            raise _failure("IMPORT_XLSX_MEMBER_LIMIT_EXCEEDED", "Workbook contains too many ZIP members.")
        if len(names) != len(set(names)):
            raise _failure("IMPORT_XLSX_DUPLICATE_MEMBER", "Workbook contains duplicate ZIP entries.")
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise _failure("IMPORT_XLSX_INVALID", "The file is not a valid .xlsx workbook.")
        for info in infos:
            if info.flag_bits & 0x1:
                raise _failure("IMPORT_XLSX_ENCRYPTED", "Encrypted workbooks are not supported.")
            if info.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                raise _failure("IMPORT_XLSX_INVALID", "Workbook ZIP members must use stored or deflate compression.")
            path_parts = info.filename.replace("\\", "/").split("/")
            if info.filename.startswith(("/", "\\")) or ".." in path_parts:
                raise _failure("IMPORT_XLSX_INVALID", "Workbook contains an unsafe ZIP path.")
        if sum(info.file_size for info in infos) > policy.max_file_decompressed_bytes:
            raise _failure("IMPORT_XLSX_DECOMPRESSED_LIMIT_EXCEEDED", "Workbook exceeds its configured decompressed-byte limit.")

        try:
            sheet_paths = _sheet_paths(archive, names, policy)
            if not sheet_paths:
                raise _failure("IMPORT_XLSX_INVALID", "Workbook contains no worksheets.")
            if len(sheet_paths) > policy.max_worksheets_per_file:
                raise _failure("IMPORT_SHEET_LIMIT_EXCEEDED", "Workbook exceeds its configured worksheet limit.")
            xml_parts = _xml_parts(archive, names, sheet_paths, policy.max_xlsx_xml_depth)
            shape_by_path: dict[str, dict[str, int]] = {}
            formulas_by_path: dict[str, dict[str, int]] = {}
            decompressed = _DecompressedCounter(policy.max_file_decompressed_bytes)
            for info in infos:
                if info.is_dir():
                    continue
                # ZipExtFile normally clips output to the declared file_size.
                # Read against a bounded larger allowance so a lying directory
                # cannot hide extra inflated bytes before openpyxl reads it.
                bounded_info = copy(info)
                bounded_info.file_size = policy.max_file_decompressed_bytes + 1
                with archive.open(bounded_info, "r") as member:
                    counted_member = _CountingReader(member, decompressed.add)
                    if info.filename in xml_parts:
                        xml_state: dict[str, Any] = {}
                        for _event, element in _xml_events(counted_member, max_depth=policy.max_xlsx_xml_depth):
                            _scan_xml_element(
                                _event,
                                element,
                                info.filename,
                                sheet_paths,
                                shape_by_path,
                                formulas_by_path,
                                xml_state,
                                policy.max_cell_text_bytes,
                            )
                        while counted_member.read(64 * 1024):
                            pass
                    else:
                        while counted_member.read(64 * 1024):
                            pass
                    if counted_member.bytes_read != info.file_size:
                        raise _failure("IMPORT_XLSX_INVALID", "Workbook ZIP member length disagrees with its content.")
            if not include_all_units:
                _check_xlsx_shape_limits(shape_by_path, sheet_paths, selected_sheet)
            if "vbaproject.bin" in " ".join(names).casefold() or b"macroenabled" in archive.read(
                "[Content_Types].xml"
            ).lower():
                raise _failure("IMPORT_FILE_TYPE_MISMATCH", "Macro-enabled workbooks are not supported as .xlsx.")
        except ImportFailure:
            raise
        except (ParseError, zipfile.BadZipFile, RuntimeError, ValueError, OSError) as exc:
            raise _failure("IMPORT_XLSX_INVALID", "Workbook XML or ZIP structure is malformed.") from exc

        try:
            values_book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        except Exception as exc:
            raise _failure("IMPORT_XLSX_INVALID", "The .xlsx workbook could not be opened safely.") from exc
        try:
            if len(values_book.worksheets) > policy.max_worksheets_per_file:
                raise _failure("IMPORT_SHEET_LIMIT_EXCEEDED", "Workbook exceeds its configured worksheet limit.")
            if not values_book.worksheets:
                raise _failure("IMPORT_XLSX_INVALID", "Workbook contains no worksheets.")
            sheet_names = [worksheet.title for worksheet in values_book.worksheets]
            if include_all_units:
                visibility = {
                    ws.title: "visible" if ws.sheet_state == "visible" else "hidden"
                    for ws in (*values_book.worksheets, *values_book.chartsheets)
                }
                units = []
                for index, name in enumerate(sheet_names, start=1):
                    path = _path_for_sheet(sheet_paths, name)
                    shape = shape_by_path.get(path, {"rows": 0, "columns": 0})
                    rows, columns = shape["rows"], shape["columns"]
                    worksheet = values_book[name]
                    if not hasattr(worksheet, "reset_dimensions"):
                        raise _failure("IMPORT_SHEET_UNSUPPORTED", "A workbook unit is not a data worksheet.")
                    worksheet.reset_dimensions()
                    header = _inspect_worksheet_header(worksheet, rows, columns, labels, policy)
                    formulas = formulas_by_path.get(path, {
                        "formula_count": 0,
                        "formula_cached_value_count": 0,
                        "formula_empty_cache_count": 0,
                        "formula_missing_cache_count": 0,
                    })
                    units.append({
                        "name": name,
                        "index": index,
                        "visibility": visibility[name],
                        "inspection": {
                            **header,
                            "row_count": rows,
                            "column_count": columns,
                            "cell_count": rows * columns,
                            "encoding": None,
                            "delimiter": None,
                            **formulas,
                        },
                    })
                result = {"source_units": units, "decompressed_bytes": decompressed.total}
                if unit_configs is None:
                    return result
                if not unit_configs or len(unit_configs) > policy.max_source_units:
                    raise _failure("IMPORT_SOURCE_UNIT_LIMIT_EXCEEDED", "Select a valid number of source units.")
                by_name = {unit["name"]: unit for unit in units}
                selected = []
                aggregate_cells = 0
                aggregate_rows = 0
                for config in unit_configs:
                    name = config.get("unit_name")
                    unit = by_name.get(name)
                    if unit is None:
                        raise _failure("IMPORT_SOURCE_UNIT_NOT_FOUND", "Selected worksheet does not exist.")
                    meta = unit["inspection"]
                    _validate_unit_config(config, meta, policy)
                    _check_shape(meta["row_count"], meta["column_count"], policy)
                    aggregate_cells += meta["cell_count"]
                    aggregate_rows += max(0, meta["row_count"] - config["header_row_number"])
                    if aggregate_cells > policy.max_selected_cells:
                        raise _failure("IMPORT_SELECTED_CELL_LIMIT_EXCEEDED", "Selected units exceed their aggregate cell limit.")
                    if aggregate_rows > policy.max_data_rows:
                        raise _failure("IMPORT_DATA_ROW_LIMIT_EXCEEDED", "Selected units exceed their aggregate data-row limit.")
                    selected.append((unit, config))
                marker_rows = {}
                for unit, config in selected:
                    marker_rows[unit["name"]] = _selected_cell_metadata(
                        archive_path_data=data, sheet_name=unit["name"], policy=policy,
                        selected_columns=None,
                    )
                read_units = []
                for unit, config in selected:
                    worksheet = values_book[unit["name"]]
                    if not hasattr(worksheet, "reset_dimensions"):
                        raise _failure("IMPORT_SHEET_UNSUPPORTED", "Select a worksheet containing cells, not a chart sheet.")
                    worksheet.reset_dimensions()
                    read_units.append(_read_xlsx_source_unit(
                        worksheet, unit, config, marker_rows[unit["name"]], policy
                    ))
                result["read_units"] = read_units
                return result
            chosen = selected_sheet if selected_sheet is not None else sheet_names[0]
            if chosen not in sheet_names:
                raise _failure("IMPORT_SHEET_NOT_FOUND", "Selected sheet does not exist in this workbook.")
            value_ws = values_book[chosen]
            if not hasattr(value_ws, "reset_dimensions"):
                raise _failure("IMPORT_SHEET_UNSUPPORTED", "Select a worksheet containing cells, not a chart sheet.")
            value_ws.reset_dimensions()
            path = _path_for_sheet(sheet_paths, chosen)
            shape = shape_by_path.get(path, {"rows": 0, "columns": 0})
            rows = shape["rows"]
            columns = shape["columns"]
            _check_shape(rows, columns, policy)
            best_header: tuple[int, int, list[str]] | None = None
            first_text_row: tuple[int, list[str]] | None = None
            selected_header_cells: list[str] = []
            row_iterator = (
                value_ws.iter_rows(
                    min_row=1, max_row=rows, min_col=1, max_col=columns, values_only=True
                )
                if rows > 0 and columns > 0
                else ()
            )
            for row_index, cells in enumerate(row_iterator, start=1):
                row_values = [_display_cell(value) for value in cells]
                text_values = [value for value in cells if isinstance(value, str) and value.strip()]
                if first_text_row is None and text_values and len(text_values) * 2 >= max(1, len(row_values)):
                    first_text_row = (row_index, row_values)
                score = sum(_normalize_header(value) in labels for value in row_values)
                if score and (best_header is None or score > best_header[0]):
                    best_header = (score, row_index, row_values)
                if selected_header_row == row_index:
                    selected_header_cells = row_values
            if selected_header_row is not None:
                if not isinstance(selected_header_row, int) or not 1 <= selected_header_row <= max(rows, 1):
                    raise _failure("IMPORT_HEADER_ROW_INVALID", "Header row must be within the selected sheet.")
                header_row = selected_header_row
                header_state = "selected"
                if not selected_header_cells:
                    selected_header_cells = _row_at(value_ws, selected_header_row, columns)
            elif best_header is not None:
                _score, header_row, selected_header_cells = best_header
                header_state = "detected"
            elif first_text_row is not None:
                header_row, selected_header_cells = first_text_row
                header_state = "detected"
            else:
                header_row, selected_header_cells, header_state = 1, [], "detected"

            visibility = {
                ws.title: "visible" if ws.sheet_state == "visible" else "hidden"
                for ws in (*values_book.worksheets, *values_book.chartsheets)
            }
            formula_counts = formulas_by_path.get(
                path,
                {
                    "formula_count": 0,
                    "formula_cached_value_count": 0,
                    "formula_empty_cache_count": 0,
                    "formula_missing_cache_count": 0,
                },
            )
            result = {
                "sheets": [
                    {"name": name, "index": index, "visibility": visibility[name]}
                    for index, name in enumerate(sheet_names, start=1)
                ],
                "selected_sheet": chosen,
                "selected_sheet_index": sheet_names.index(chosen) + 1,
                "header_row_number": header_row,
                "header_state": header_state,
                "header_cells": selected_header_cells[:policy.max_columns_per_unit],
                "row_count": rows,
                "column_count": columns,
                "cell_count": rows * columns,
                "encoding": None,
                "delimiter": None,
                **formula_counts,
            }
        except ImportFailure:
            raise
        except Exception as exc:
            raise _failure("IMPORT_XLSX_INVALID", "Workbook cell content or styles are malformed.") from exc
        finally:
            values_book.close()

    return result


def _sheet_paths(
    archive: zipfile.ZipFile,
    names: list[str],
    policy: ImportResourcePolicy = DEFAULT_IMPORT_RESOURCE_POLICY,
) -> dict[str, str]:
    sheets: list[tuple[str, str]] = []
    for attributes in _metadata_elements(archive, "xl/workbook.xml", "sheet"):
        sheets.append((attributes.get("name", ""), next(
            (value for key, value in attributes.items() if key.endswith("}id")), "")))
        if len(sheets) > policy.max_xlsx_members:
            raise _failure("IMPORT_XLSX_MEMBER_LIMIT_EXCEEDED", "Workbook declares too many sheet entries.")
    required = {rel for _, rel in sheets}
    targets: dict[str, tuple[str, str]] = {}
    for attributes in _metadata_elements(archive, "xl/_rels/workbook.xml.rels", "Relationship"):
        rel = attributes.get("Id", "")
        if rel in required:
            if rel in targets or attributes.get("TargetMode") == "External":
                raise _failure("IMPORT_XLSX_INVALID", "Worksheet relationships must be unique and internal.")
            targets[rel] = (attributes.get("Target", ""), attributes.get("Type", ""))
    paths: dict[str, str] = {}
    for name, rel_id in sheets:
        relationship = targets.get(rel_id)
        if not name or not relationship:
            raise _failure("IMPORT_XLSX_INVALID", "Workbook contains an invalid sheet relationship.")
        target, relationship_type = relationship
        if relationship_type.endswith("/chartsheet"):
            continue
        if not relationship_type.endswith("/worksheet"):
            raise _failure("IMPORT_XLSX_INVALID", "Workbook sheet relationship is not a worksheet.")
        path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        if not path.startswith("xl/") or path not in names:
            raise _failure("IMPORT_XLSX_INVALID", "Workbook contains an unsafe sheet relationship.")
        if name in paths:
            raise _failure("IMPORT_XLSX_INVALID", "Workbook contains duplicate worksheet names.")
        paths[name] = path
        if len(paths) > policy.max_worksheets_per_file:
            raise _failure("IMPORT_SHEET_LIMIT_EXCEEDED", "Workbook exceeds the configured worksheet limit.")
    return paths


def _metadata_elements(
    archive: zipfile.ZipFile,
    name: str,
    tag: str,
    max_depth: int = MAX_XLSX_XML_DEPTH,
):
    try:
        with archive.open(name) as source:
            for event, element in _xml_events(source, max_depth=max_depth):
                if event == "end" and _local_name(element.tag) == tag:
                    yield dict(element.attrib)
    except (ParseError, OSError, KeyError, zipfile.BadZipFile) as exc:
        raise _failure("IMPORT_XLSX_INVALID", "Workbook XML is malformed.") from exc


def _xml_parts(
    archive: zipfile.ZipFile,
    names: list[str],
    sheet_paths: dict[str, str],
    max_depth: int = MAX_XLSX_XML_DEPTH,
) -> set[str]:
    """OPC part types, not suffix alone, determine XML parsed by workbook readers."""
    members = set(names)
    xml = {name for name in names if name.lower().endswith((".xml", ".rels"))}
    xml.update(sheet_paths.values())
    extensions: dict[str, set[str]] = {}
    for name in names:
        leaf = name.rsplit("/", 1)[-1]
        if "." in leaf:
            extensions.setdefault(leaf.rsplit(".", 1)[-1].lower(), set()).add(name)
    declared_xml: set[str] = set()
    seen_defaults: set[str] = set()
    seen_overrides: set[str] = set()
    # Retain only names from the already member-count-bounded archive, even
    # when an untrusted manifest declares many nonexistent parts/extensions.
    for tag in ("Default", "Override"):
        for attributes in _metadata_elements(archive, "[Content_Types].xml", tag, max_depth):
            content_type = attributes.get("ContentType", "").split(";", 1)[0].strip().lower()
            is_xml = content_type in ("application/xml", "text/xml") or content_type.endswith("+xml")
            if tag == "Default":
                extension = attributes.get("Extension", "").lower()
                if extension not in extensions:
                    continue
                if extension in seen_defaults:
                    raise _failure("IMPORT_XLSX_INVALID", "Workbook declares duplicate content-type defaults.")
                seen_defaults.add(extension)
                if is_xml:
                    declared_xml.update(extensions[extension])
            else:
                part = attributes.get("PartName", "").lstrip("/")
                if part in members:
                    if part in seen_overrides:
                        raise _failure("IMPORT_XLSX_INVALID", "Workbook declares duplicate content-type overrides.")
                    seen_overrides.add(part)
                    if is_xml:
                        declared_xml.add(part)
                    else:
                        declared_xml.discard(part)
    return xml | declared_xml


def _xml_events(source, *, max_depth: int = MAX_XLSX_XML_DEPTH):
    """Discard finished children as well as text: empty XML nodes also cost memory."""
    stack = []
    for event, element in SafeET.iterparse(source, events=("start", "end"),
                                          forbid_dtd=True, forbid_entities=True, forbid_external=True):
        if event == "start":
            if len(stack) >= max_depth:
                raise _failure("IMPORT_XLSX_XML_DEPTH_LIMIT_EXCEEDED",
                               f"Workbook XML exceeds the internal depth bound of {max_depth} elements.")
            stack.append(element)
        yield event, element
        if event == "end":
            stack.pop()
            element.clear()
            if stack:
                stack[-1].remove(element)


def _check_xlsx_shape_limits(
    shapes: dict[str, dict[str, int]], sheet_paths: dict[str, str], selected_sheet: str | None
) -> None:
    if selected_sheet is not None and selected_sheet not in sheet_paths:
        raise _failure("IMPORT_SHEET_NOT_FOUND", "Selected sheet does not exist in this workbook.")
    chosen = selected_sheet if selected_sheet is not None else next(iter(sheet_paths))
    path = sheet_paths[chosen]
    shape = shapes.get(path, {"rows": 0, "columns": 0})
    _check_shape(shape["rows"], shape["columns"])


class _CountingReader:
    """Read-only wrapper that counts actual decompressed XML bytes."""

    def __init__(self, source, count_bytes) -> None:
        self._source = source
        self._count_bytes = count_bytes
        self.bytes_read = 0

    def read(self, size: int = -1) -> bytes:
        chunk = self._source.read(size)
        self.bytes_read += len(chunk)
        self._count_bytes(len(chunk))
        return chunk


class _DecompressedCounter:
    def __init__(self, max_bytes: int = MAX_XLSX_DECOMPRESSED_BYTES) -> None:
        self.total = 0
        self.max_bytes = max_bytes

    def add(self, size: int) -> None:
        self.total += size
        if self.total > self.max_bytes:
            raise _failure("IMPORT_XLSX_DECOMPRESSED_LIMIT_EXCEEDED", "Workbook exceeds its configured decompressed-byte limit.")


def _scan_xml_element(
    event: str,
    element,
    member_name: str,
    sheet_paths: dict[str, str],
    shapes: dict[str, dict[str, int]],
    formulas: dict[str, dict[str, int]],
    state: dict[str, Any],
    max_cell_text_bytes: int = MAX_CELL_TEXT_BYTES,
) -> None:
    # Iterparse emits start/end events. Accumulate one cell or shared string at
    # a time, preserve worksheet coordinates, and clear completed XML nodes.
    tag = _local_name(element.tag)
    is_sheet = member_name in sheet_paths.values()
    if event == "start":
        if tag == "row" and is_sheet:
            row_text = element.attrib.get("r")
            row_number = int(row_text) if row_text and row_text.isdigit() else state.get("last_row", 0) + 1
            if row_number <= state.get("last_row", 0):
                raise _failure("IMPORT_XLSX_INVALID", "Workbook rows must have unique increasing coordinates.")
            state["current_row"] = row_number
            state["last_row"] = row_number
            state["last_column"] = 0
        elif tag == "c" and is_sheet:
            match = _CELL_REF.fullmatch(element.attrib.get("r", ""))
            if match is None:
                raise _failure("IMPORT_XLSX_INVALID", "Workbook contains an invalid cell reference.")
            column, row = _column_number(match.group(1)), int(match.group(2))
            if row != state.get("current_row") or column <= state.get("last_column", 0):
                raise _failure("IMPORT_XLSX_INVALID", "Workbook cell coordinates must be unique and increasing within each row.")
            state.update(context="c", text_bytes=0, formula=False, value_seen=False,
                         value_nonempty=False, cell_type=element.attrib.get("t"),
                         cell_row=row, cell_column=column)
        elif tag == "si":
            state.clear()
            state.update(context="si", text_bytes=0)
        elif tag == "f" and state.get("context") == "c":
            state["formula"] = True
        elif tag == "v" and state.get("context") == "c":
            state["value_seen"] = True
        return

    if tag in {"t", "v", "f"}:
        value = element.text or ""
        if state.get("context"):
            state["text_bytes"] += len(value.encode("utf-8"))
            if state["text_bytes"] > max_cell_text_bytes:
                raise _failure("IMPORT_CELL_TEXT_LIMIT_EXCEEDED", "A cell exceeds its configured text limit.")
            if tag == "v" and value:
                state["value_nonempty"] = True
            if tag == "v":
                state["raw_value"] = value
    if tag == "row" and is_sheet:
        row = state.get("current_row", 0)
        shape = shapes.setdefault(member_name, {"rows": 0, "columns": 0})
        shape["rows"] = max(shape["rows"], row)
    elif tag == "c" and is_sheet:
        shape = shapes.setdefault(member_name, {"rows": 0, "columns": 0})
        column = state.get("cell_column", 0)
        row = state.get("cell_row", 0)
        shape["rows"] = max(shape["rows"], row)
        shape["columns"] = max(shape["columns"], column)
        if state.get("formula"):
            counts = formulas.setdefault(
                member_name,
                {
                    "formula_count": 0,
                    "formula_cached_value_count": 0,
                    "formula_empty_cache_count": 0,
                    "formula_missing_cache_count": 0,
                },
            )
            counts["formula_count"] += 1
            if not state["value_seen"]:
                counts["formula_missing_cache_count"] += 1
            elif state["value_nonempty"]:
                counts["formula_cached_value_count"] += 1
            else:
                counts["formula_empty_cache_count"] += 1
        state["last_column"] = column
        state.pop("context", None)
    elif tag == "si":
        state.clear()
    elif tag == "si":
        state.clear()
    element.clear()


def _path_for_sheet(sheet_paths: dict[str, str], sheet_name: str) -> str:
    try:
        return sheet_paths[sheet_name]
    except KeyError as exc:
        raise _failure("IMPORT_SHEET_NOT_FOUND", "Selected sheet does not exist in this workbook.") from exc


def _check_shape(
    rows: int,
    columns: int,
    policy: ImportResourcePolicy = DEFAULT_IMPORT_RESOURCE_POLICY,
) -> None:
    if rows > policy.max_rows_per_unit:
        raise _failure("IMPORT_ROW_LIMIT_EXCEEDED", "Selected unit exceeds its configured row limit.")
    if columns > policy.max_columns_per_unit:
        raise _failure("IMPORT_COLUMN_LIMIT_EXCEEDED", "Selected unit exceeds its configured column limit.")
    if rows * columns > policy.max_cells_per_unit:
        raise _failure("IMPORT_CELL_LIMIT_EXCEEDED", "Selected unit exceeds its configured rectangular-cell limit.")


def _row_at(worksheet, row_number: int, columns: int) -> list[str]:
    if columns <= 0:
        return []
    return [_display_cell(value) for value in next(
        worksheet.iter_rows(min_row=row_number, max_row=row_number, min_col=1, max_col=columns, values_only=True)
    )]


def _header_labels(header_labels: Iterable[str] | None) -> set[str]:
    labels = {_normalize_header(value) for value in _BASE_HEADER_LABELS}
    if header_labels is not None:
        labels.update(_normalize_header(value) for value in header_labels if isinstance(value, str))
    return labels


def _normalize_header(value: str) -> str:
    return " ".join(value.strip().casefold().split())


def _normalize_date_order(value: str | None) -> str | None:
    if value is None or value == "":
        return None
    normalized = value.strip().upper()
    if normalized not in {"DMY", "MDY", "YMD"}:
        raise _failure("IMPORT_DATE_ORDER_INVALID", "Date order must be DMY, MDY or YMD.")
    return normalized


def _display_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _check_text_values(values: Iterable[str], max_bytes: int = MAX_CELL_TEXT_BYTES) -> None:
    for value in values:
        _check_xml_text(value, max_bytes)


def _check_xml_text(value: str, max_bytes: int = MAX_CELL_TEXT_BYTES) -> None:
    if len(value.encode("utf-8")) > max_bytes:
        raise _failure("IMPORT_CELL_TEXT_LIMIT_EXCEEDED", "A cell exceeds its configured text limit.")


def _column_number(letters: str) -> int:
    result = 0
    for letter in letters:
        result = result * 26 + ord(letter) - ord("A") + 1
    return result


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _enforce_metadata_bound(
    metadata: dict[str, Any], max_bytes: int = MAX_INSPECTION_JSON_BYTES
) -> None:
    try:
        serialized = json.dumps(metadata, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    except (TypeError, ValueError) as exc:
        raise _failure("IMPORT_INSPECTION_INVALID", "File inspection metadata is not serializable.") from exc
    if len(serialized) > max_bytes:
        raise _failure("IMPORT_INSPECTION_TOO_LARGE", "Inspection metadata exceeds its configured limit.")


def _failure(code: str, message: str) -> ImportFailure:
    return ImportFailure(code, message)
