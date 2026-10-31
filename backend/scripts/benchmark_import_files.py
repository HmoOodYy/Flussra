"""Measure compact multi-file/layout preview on disposable production-schema databases.

Run from backend with Python 3.12: python -m scripts.benchmark_import_files --scenario current
Named historical presets remain reference workloads and may exceed current product rules.
Each workload runs in its own process; peak working set includes fixture overhead.
No application database or public apply endpoint is used.
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import io
import json
import os
import platform
import statistics
import subprocess
import sys
import tempfile
import time
from collections import Counter
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from scripts.import_benchmark_support import (
    application,
    memory_snapshot,
    private_database,
    seed_period,
)

ROOT = Path(__file__).resolve().parents[2]
SCENARIOS = ("representative_csv", "monthly_daily_csv", "monthly_daily_xlsx", "rare_csv", "issue_heavy_csv")
BASE_CURRENT_SCENARIOS = ("current_daily_csv", "current_daily_xlsx", "current_daily_issue_csv")
CAPABILITY_SCENARIOS = ("current_target_mapping_csv", "current_identity_csv", "current_date_column_csv",
                        "current_filename_dates_csv", "current_sheet_dates_xlsx", "current_repeated_rows_csv")
ACCEPTANCE_SCENARIOS = ("current_completed_repeat_csv", "current_partial_overlap_csv",
                        "current_canonical_byte_boundary_csv")
CURRENT_SCENARIOS = (*BASE_CURRENT_SCENARIOS, *CAPABILITY_SCENARIOS, *ACCEPTANCE_SCENARIOS)
SINGLE_DAY_SCENARIOS = ("single_day_csv", "single_day_xlsx")
CURRENT_POLICY_SCENARIOS = (*CURRENT_SCENARIOS, *SINGLE_DAY_SCENARIOS)
SCENARIO_GROUPS = {"current": CURRENT_SCENARIOS, "single_day": SINGLE_DAY_SCENARIOS, "all": SCENARIOS}
# Drivers in the period Branch, Pay Item definitions, period days, Company Drivers outside that Branch.
CURRENT_DIMENSIONS = {"current_target_mapping_csv": (150, 10, 31, 0), "current_identity_csv": (150, 4, 31, 1500),
                      "current_date_column_csv": (150, 6, 31, 0), "current_filename_dates_csv": (150, 6, 31, 0),
                      "current_sheet_dates_xlsx": (150, 6, 31, 0), "current_repeated_rows_csv": (150, 3, 31, 0),
                      "current_completed_repeat_csv": (150, 14, 31, 0),
                      "current_partial_overlap_csv": (150, 6, 31, 0),
                      # One uploaded day is measured within a normal 31-day setup period.
                      "current_canonical_byte_boundary_csv": (150, 1, 31, 0),
                      "single_day_csv": (150, 14, 31, 0), "single_day_xlsx": (150, 14, 31, 0)}
# One day is deliberately absent so the Date review reports missing coverage.
DATE_AUTHORITIES = {"current_date_column_csv": "column", "current_filename_dates_csv": "filename",
                    "current_sheet_dates_xlsx": "sheet"}
SINGLE_DAY_DATE_AUTHORITIES = {"single_day_csv": "filename", "single_day_xlsx": "sheet"}
DATE_AUTHORITIES_BY_SCENARIO = {**DATE_AUTHORITIES, **SINGLE_DAY_DATE_AUTHORITIES}
MISSING_COVERAGE_DAY_INDEX = 10
COLUMN_DATE_FILE_DAYS = 7
STRUCTURAL_HEADERS = ("DriverNumber", "DriverName", "WorkDate")
SUM_BLANK_EVERY = 25
IDENTITY_PERIOD = 50
IDENTITY_FINDINGS = {"name_variation": "IMPORT_DRIVER_NAME_DISCREPANCY", "unknown_code": "IMPORT_DRIVER_UNRESOLVED",
                     "name_only": "IMPORT_DRIVER_CONFIRMATION_REQUIRED"}
REPEATS = 3
CATEGORY_COUNT = 120
CATEGORY_HEADER = "Category"
NOT_COVERED = ("Successful Import completion/application and active rule-head promotion remain unimplemented; "
               "the completed-repeat workload uses a controlled test-style Completed receipt fixture; "
               "current-head reuse remains unexercised.")


def scenario_dimensions(scene):
    return (CURRENT_DIMENSIONS.get(scene, (150, 14, 31, 0)) if scene in CURRENT_POLICY_SCENARIOS else
            (20, 4, 7, 0) if scene == "representative_csv" else
            (1000, 20, 10, 0) if scene in ("rare_csv", "issue_heavy_csv") else (150, 20, 31, 0))


def scenario_covered_days(scene, days):
    """Days with source rows, including the single-day canonical-boundary file."""
    if scene == "current_canonical_byte_boundary_csv" or scene in SINGLE_DAY_SCENARIOS:
        return list(days[:1])
    return [day for index, day in enumerate(days)
            if not (scene in DATE_AUTHORITIES and index == MISSING_COVERAGE_DAY_INDEX)]


def written_date(day):
    """A written month and day without a year, such as Oct 16."""
    return f"{day:%b} {day.day}"


def target_mapping_plan(definitions):
    """Direct inputs plus Sum targets over distinct and duplicate-header inputs."""
    codes = [definition.code for definition in definitions]
    if len(codes) < 4:
        raise ValueError("The mapping workload needs at least four definitions.")
    targets = [{"code": code, "transform": "Direct", "headers": [code], "operands": ["1"], "blankable": False}
               for code in codes[:-3]]
    regular, triple, repeated = codes[-3:]
    targets += [
        {"code": regular, "transform": "Sum", "headers": [f"{regular}_Regular", f"{regular}_Bonus"],
         "operands": ["1", "2"], "blankable": True},
        {"code": triple, "transform": "Sum", "headers": [f"{triple}_A", f"{triple}_B", f"{triple}_C"],
         "operands": ["1", "2", "3"], "blankable": False},
        {"code": repeated, "transform": "Sum", "headers": [repeated, repeated],
         "operands": ["0.5", "0.25"], "blankable": False}]
    return {"shared": ["DriverNumber"], "ignored": ["Region", "Notes"], "targets": targets,
            "header": ["DriverNumber", "Region", "Notes", *[header for item in targets for header in item["headers"]]]}


def mapping_operand_blank(driver_index):
    return driver_index % SUM_BLANK_EVERY == 0


def mapping_row(plan, driver_index, driver_number):
    cells = [driver_number, f"R{driver_index % 7}", f"note {driver_index}"]
    for target in plan["targets"]:
        operands = list(target["operands"])
        if target["blankable"] and mapping_operand_blank(driver_index):
            operands[-1] = ""
        cells.extend(operands)
    return cells


def mapping_expected(target, driver_index):
    if target["blankable"] and mapping_operand_blank(driver_index):
        return None
    return sum((Decimal(item) for item in target["operands"]), Decimal(0))


def identity_kind(driver_index):
    return {0: "name_variation", 1: "unknown_code", 2: "name_only"}.get(driver_index % IDENTITY_PERIOD, "exact")


def identity_row(driver, driver_index, day_index, quantities):
    number, name = driver["drivernumber"], driver["fullname"]
    kind = identity_kind(driver_index)
    if kind == "name_variation":
        name = f"{name} (alias)"
    elif kind == "unknown_code":
        number, name = f"UNKNOWN-{day_index:02d}-{driver_index:04d}", f"Unlisted Driver {driver_index}"
    elif kind == "name_only":
        number = ""
    return [number, name, *quantities]


def repeated_category(driver_index, day_index, repeat):
    return (driver_index * 7 + day_index * 3 + repeat) % CATEGORY_COUNT


def category_label(category):
    return f"Route {category + 1:03d}"


def category_excluded(category):
    return category % 4 == 3


def repeated_quantities(repeat, count):
    return [(1 + repeat, 2, 1)[position % 3] for position in range(count)]


def repeated_expected(driver_index, day_index, count, *, filtered):
    """Per-definition totals for one Driver/day, or None when every row is excluded."""
    totals, included = [Decimal(0)] * count, 0
    for repeat in range(REPEATS):
        if filtered and category_excluded(repeated_category(driver_index, day_index, repeat)):
            continue
        included += 1
        totals = [total + quantity for total, quantity in zip(totals, repeated_quantities(repeat, count), strict=True)]
    return totals if included else None


def repeated_totals(driver_count, day_count):
    """Row, group and excluded-row totals of the repeated-row workload."""
    excluded = visible = 0
    for driver_index in range(driver_count):
        for day_index in range(day_count):
            kept = [not category_excluded(repeated_category(driver_index, day_index, repeat))
                    for repeat in range(REPEATS)]
            excluded += kept.count(False)
            visible += any(kept)
    return {"rows": driver_count * day_count * REPEATS, "groups": driver_count * day_count,
            "excluded_rows": excluded, "visible_groups": visible}


def _csv_bytes(header, rows):
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def canonical_boundary_payloads(fixture):
    """Build accepted/over-limit CSVs from production canonical-byte measurements."""
    from app.imports.canonical import extract_canonical_file
    from app.imports.policy import DEFAULT_IMPORT_RESOURCE_POLICY
    from app.product_policy import DEFAULT_PRODUCT_POLICY

    target_bytes = DEFAULT_PRODUCT_POLICY.max_import_canonical_file_bytes
    drivers = fixture["drivers"]
    if not drivers:
        raise ValueError("The canonical-boundary workload needs at least one Driver row.")
    max_cell_chars = DEFAULT_IMPORT_RESOURCE_POLICY.max_cell_text_bytes
    probe_policy = replace(DEFAULT_PRODUCT_POLICY, max_import_canonical_file_bytes=target_bytes + 1)

    def payload_with_padding(total_chars):
        per_row, remainder = divmod(total_chars, len(drivers))
        if per_row + (remainder > 0) > max_cell_chars:
            raise ValueError("Canonical boundary cannot fit within the per-cell text limit.")
        rows = []
        for index, driver in enumerate(drivers):
            padding = per_row + (index < remainder)
            rows.append([driver["drivernumber"], "1", "x" * padding])
        return _csv_bytes(["DriverNumber", "Q1", "Notes"], rows)

    def canonical_count(total_chars):
        data = payload_with_padding(total_chars)
        extracted = extract_canonical_file(data, "2026-01-01.csv", product_policy=probe_policy)
        return extracted["canonical_byte_count"]

    baseline_chars = len(drivers)
    baseline_count = canonical_count(baseline_chars)
    padding_chars = baseline_chars + (target_bytes - baseline_count)
    maximum_chars = len(drivers) * max_cell_chars
    if target_bytes <= baseline_count or padding_chars + 1 > maximum_chars:
        raise ValueError("Canonical boundary cannot fit within the per-cell text limit.")
    exact = payload_with_padding(padding_chars)
    exact_count = canonical_count(padding_chars)
    if exact_count != target_bytes:
        raise ValueError(
            f"Could not construct exact canonical boundary: expected {target_bytes}, got {exact_count}.")
    over = payload_with_padding(padding_chars + 1)
    over_count = canonical_count(padding_chars + 1)
    if over_count != target_bytes + 1:
        raise ValueError(f"One additional padding byte changed canonical size to {over_count}, not {target_bytes + 1}.")
    return exact, over, {"limit_bytes": target_bytes, "accepted_canonical_bytes": exact_count,
                         "rejected_canonical_bytes": over_count, "padding_characters": padding_chars}


def _xlsx_bytes(name, header, rows):
    from openpyxl import Workbook
    book = Workbook(write_only=True)
    try:
        sheet = book.create_sheet(name)
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()
    finally:
        book.close()


def capability_files(scene, fixture):
    definitions, drivers, days = fixture["definitions"], fixture["drivers"], fixture["days"]
    codes = [definition.code for definition in definitions]
    covered = scenario_covered_days(scene, days)
    day_index = {day: index for index, day in enumerate(days)}
    if scene == "current_target_mapping_csv":
        plan = target_mapping_plan(definitions)
        return [(f"{day.isoformat()}.csv", _csv_bytes(plan["header"], (
            mapping_row(plan, index, driver["drivernumber"]) for index, driver in enumerate(drivers))))
            for day in days]
    if scene == "current_identity_csv":
        return [(f"{day.isoformat()}.csv", _csv_bytes(["DriverNumber", "DriverName", *codes], (
            identity_row(driver, index, day_index[day], ["1"] * len(codes)) for index, driver in enumerate(drivers))))
            for day in days]
    if scene == "current_date_column_csv":
        # Weekly files carry many dates; the last one writes them as a month and day without a year.
        chunks = [covered[start:start + COLUMN_DATE_FILE_DAYS]
                  for start in range(0, len(covered), COLUMN_DATE_FILE_DAYS)]
        return [(f"week-{number}.csv", _csv_bytes(["DriverNumber", "WorkDate", *codes], (
            [driver["drivernumber"], written_date(day) if number == len(chunks) else day.isoformat(),
             *["1"] * len(codes)] for day in chunk for driver in drivers)))
            for number, chunk in enumerate(chunks, 1)]
    if scene == "current_filename_dates_csv":
        return [(f"{written_date(day)}.csv", _csv_bytes(["DriverNumber", *codes], (
            [driver["drivernumber"], *["1"] * len(codes)] for driver in drivers))) for day in covered]
    if scene == "current_sheet_dates_xlsx":
        return [(f"driver-log-{number:02d}.xlsx", _xlsx_bytes(written_date(day), ["DriverNumber", *codes], (
            [driver["drivernumber"], *[1] * len(codes)] for driver in drivers)))
            for number, day in enumerate(covered, 1)]
    if scene == "current_partial_overlap_csv":
        if len(codes) != 6:
            raise ValueError("The partial-overlap workload requires six target definitions.")
        files = []
        for day_index, day in enumerate(days):
            available = [codes[0], *(codes[1:3] if day_index % 2 == 0 else codes[3:6])]
            files.append((f"{day.isoformat()}.csv", _csv_bytes(
                ["DriverNumber", *available],
                ([driver["drivernumber"], *["1"] * len(available)] for driver in drivers))))
        return files
    if scene == "current_canonical_byte_boundary_csv":
        if len(codes) != 1:
            raise ValueError("The canonical-boundary workload requires one target definition.")
        exact, _over, _evidence = canonical_boundary_payloads(fixture)
        return [(f"{days[0].isoformat()}.csv", exact)]
    if scene == "current_completed_repeat_csv":
        return [(f"{day.isoformat()}.csv", _csv_bytes(["DriverNumber", *codes], (
            [driver["drivernumber"], *["1"] * len(codes)] for driver in drivers))) for day in days]
    return [(f"{day.isoformat()}.csv", _csv_bytes(["DriverNumber", CATEGORY_HEADER, *codes], (
        [driver["drivernumber"], category_label(repeated_category(index, day_index[day], repeat)),
         *map(str, repeated_quantities(repeat, len(codes)))]
        for index, driver in enumerate(drivers) for repeat in range(REPEATS)))) for day in days]


def source_files(scene, fixture):
    if scene in CAPABILITY_SCENARIOS or scene in ACCEPTANCE_SCENARIOS:
        return capability_files(scene, fixture)
    definitions, drivers, days = fixture["definitions"], fixture["drivers"], fixture["days"]
    def csv_bytes(selected_definitions, selected_days, *, date_column):
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerow(["DriverNumber", *(["WorkDate"] if date_column else []),
                         *[definition.code for definition in selected_definitions]])
        for day in selected_days:
            for driver in drivers:
                writer.writerow([driver["drivernumber"], *([day.isoformat()] if date_column else []),
                    *[("invalid" if scene in ("issue_heavy_csv", "current_daily_issue_csv") else "1")
                      for _ in selected_definitions]])
        return buffer.getvalue().encode("utf-8")
    if scene == "monthly_daily_csv":
        # Two different layouts share a Driver/day grid; each supplies half the dynamic definitions.
        midpoint = len(definitions) // 2
        return [(f"{day.isoformat()}-family-{family}.csv", csv_bytes(group, [day], date_column=False))
                for day in days for family, group in enumerate((definitions[:midpoint], definitions[midpoint:]), 1)]
    if scene == "monthly_daily_xlsx":
        from openpyxl import Workbook
        book = Workbook(write_only=True)
        try:
            for day in days:
                sheet = book.create_sheet(day.isoformat())
                sheet.append(["DriverNumber", *[definition.code for definition in definitions]])
                for driver in drivers:
                    sheet.append([driver["drivernumber"], *[1 for _ in definitions]])
            buffer = io.BytesIO()
            book.save(buffer)
            return [("monthly.xlsx", buffer.getvalue())]
        finally:
            book.close()
    if scene in SINGLE_DAY_SCENARIOS:
        day = days[0]
        if scene == "single_day_csv":
            return [(f"{day.isoformat()}.csv", csv_bytes(definitions, [day], date_column=False))]
        return [(f"{day.isoformat()}.xlsx", _xlsx_bytes(day.isoformat(),
                ["DriverNumber", *[definition.code for definition in definitions]],
                ([driver["drivernumber"], *[1 for _ in definitions]] for driver in drivers)))]
    if scene in CURRENT_POLICY_SCENARIOS:
        if scene == "current_daily_xlsx":
            from openpyxl import Workbook
            files = []
            for day in days:
                book = Workbook(write_only=True)
                try:
                    sheet = book.create_sheet(day.isoformat())
                    sheet.append(["DriverNumber", *[definition.code for definition in definitions]])
                    for driver in drivers:
                        sheet.append([driver["drivernumber"], *[1 for _ in definitions]])
                    buffer = io.BytesIO()
                    book.save(buffer)
                    files.append((f"{day.isoformat()}.xlsx", buffer.getvalue()))
                finally:
                    book.close()
            return files
        return [(f"{day.isoformat()}.csv", csv_bytes(definitions, [day], date_column=False))
                for day in days]
    return [("supplied.csv", csv_bytes(definitions, days, date_column=True))]


class Phases:
    def __init__(self):
        self.values = {}
        self.calls = {}
        self.originals = []
        self.lock_started = None
        self.copies = {}

    def _record(self, name, started):
        self.values[name] = self.values.get(name, 0) + time.perf_counter() - started
        self.calls[name] = self.calls.get(name, 0) + 1

    def _wrap_async(self, module, name, label=None):
        original = getattr(module, name)
        phase = label or name
        async def measured(*args, **kwargs):
            started = time.perf_counter()
            try:
                return await original(*args, **kwargs)
            finally:
                self._record(phase, started)
        self.originals.append((module, name, original))
        setattr(module, name, measured)

    def install(self):
        from app.imports import batch_service
        for name in ("_inspect_batch", "_materialize_batch", "build_batch_preview", "prepare_preview"):
            original = getattr(batch_service, name)
            def measured(*args, _original=original, _name=name, **kwargs):
                started = time.perf_counter()
                try:
                    return _original(*args, **kwargs)
                finally:
                    self._record(_name, started)
            self.originals.append((batch_service, name, original))
            setattr(batch_service, name, measured)
        self._wrap_async(batch_service, "_completed_receipt_snapshot", "completed_receipt_snapshot")
        from app.imports import preview_model
        original_match = preview_model.matching_completed_receipts
        def measured_match(*args, **kwargs):
            started = time.perf_counter()
            try:
                return original_match(*args, **kwargs)
            finally:
                self._record("matching_completed_receipts", started)
        self.originals.append((preview_model, "matching_completed_receipts", original_match))
        preview_model.matching_completed_receipts = measured_match
        from app.imports import preview_store
        original_copy = preview_store._copy_records
        async def measured_copy(db, table, columns, records):
            count = 0
            def counted():
                nonlocal count
                for row in records:
                    count += 1
                    yield row
            started = time.perf_counter()
            await original_copy(db, table, columns, counted())
            self.values["copy_" + table] = time.perf_counter() - started
            self.copies[table] = {"records": count, "batches": (count + 999) // 1000}
        self.originals.append((preview_store, "_copy_records", original_copy))
        preview_store._copy_records = measured_copy
        for name in ("load_context", "persist_preview", "clear_derived_state", "_lock_mutable"):
            original = getattr(batch_service, name)
            async def measured(*args, _original=original, _name=name, **kwargs):
                started = time.perf_counter()
                try:
                    result = await _original(*args, **kwargs)
                    if _name == "_lock_mutable":
                        self.lock_started = time.perf_counter()
                        args[0].info["measured_import_lock"] = self.lock_started
                    return result
                finally:
                    self._record(_name, started)
            self.originals.append((batch_service, name, original))
            setattr(batch_service, name, measured)
        # Only functions the service resolves through their module attribute are wrapped, so a timer that
        # never fires means the step did not run (for example, no Company name scan for code-only files).
        from app.imports import payroll_context, preview_model
        for module, name, label in (
            (batch_service, "_prepare_preview_dates", "date_preparation"),
            (batch_service, "_stable_identifiers", "identity_requested_keys"),
            (batch_service, "interpret_date", "date_interpretation"),
            (preview_model, "interpret_date", "date_interpretation"),
            (preview_model, "resolve_row_identity", "identity_row_resolution"),
            (preview_model, "_apply_row_handling", "row_handling_aggregation"),
            (preview_model, "aggregate_quantity_contributions", "aggregate_contributions"),
            (batch_service, "_inspect_category_sources", "category_evidence"),
        ):
            original = getattr(module, name)
            def measured_sync(*args, _original=original, _label=label, **kwargs):
                started = time.perf_counter()
                try:
                    return _original(*args, **kwargs)
                finally:
                    self._record(_label, started)
            self.originals.append((module, name, original))
            setattr(module, name, measured_sync)
        for module, name, label in (
            (payroll_context, "_lookup_company_names", "identity_company_name_lookup"),
            (batch_service, "_save_mapping", "mapping_save"),
            (batch_service, "_prepare_current_rules", "rule_preparation"),
        ):
            original = getattr(module, name)
            async def measured_async(*args, _original=original, _label=label, **kwargs):
                started = time.perf_counter()
                try:
                    return await _original(*args, **kwargs)
                finally:
                    self._record(_label, started)
            self.originals.append((module, name, original))
            setattr(module, name, measured_async)

    def reset(self):
        self.values = {}
        self.calls = {}
        self.lock_started = None
        self.copies = {}

    def restore(self):
        for module, name, original in self.originals:
            setattr(module, name, original)


async def timed_request(client, path, headers, send, phases):
    phases.reset()
    baseline_started = time.perf_counter()
    ordinary = await client.get(path, headers=headers)
    assert ordinary.status_code == 200, ordinary.text
    baseline = time.perf_counter() - baseline_started
    probes = []
    started = time.perf_counter()
    pending = asyncio.create_task(send())
    while not pending.done():
        before = time.perf_counter()
        ordinary = await client.get(path, headers=headers)
        probes.append(time.perf_counter() - before)
        assert ordinary.status_code == 200, ordinary.text
        if not pending.done():
            await asyncio.sleep(0.05)
    response = await pending
    finished = time.perf_counter()
    assert response.status_code == 200, response.text
    return response.json(), {"request_seconds": finished - started, "phases_seconds": dict(phases.values),
        "phase_calls": dict(phases.calls), "response_bytes": len(response.content),
        "copy": dict(phases.copies),
        "locked_to_response_seconds": finished - phases.lock_started if phases.lock_started else None,
        "ordinary_get_baseline_seconds": baseline, "ordinary_get_count": len(probes),
        "ordinary_get_median_seconds": statistics.median(probes) if probes else None,
        "ordinary_get_max_seconds": max(probes) if probes else None, "memory": memory_snapshot()}


def json_request(client, method, path, payload, headers):
    """A bounded JSON request whose exact body size is measured."""
    body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return len(body), lambda: client.request(method, path, content=body,
                                             headers={**headers, "Content-Type": "application/json"})


def require_timers(metrics, *names, absent=()):
    """A timer that reads zero would hide an unwrapped or skipped step."""
    for name in names:
        assert metrics["phases_seconds"].get(name, 0) > 0 and metrics["phase_calls"].get(name, 0) > 0, (
            name, metrics["phases_seconds"])
    for name in absent:
        assert name not in metrics["phases_seconds"], (name, metrics["phases_seconds"])


async def get_pages(client, path, params, headers, *, limit=None):
    """Follow next_offset to the end; returns every item and the page count."""
    items, offset, pages = [], 0, 0
    while True:
        response = await client.get(path, params={**params, "offset": offset,
                                                  **({"limit": limit} if limit else {})}, headers=headers)
        assert response.status_code == 200, response.text
        assert len(response.content) <= 2_000_000
        page = response.json()
        items.extend(page["items"])
        pages += 1
        if page["next_offset"] is None:
            return items, pages
        offset = page["next_offset"]


async def read_rows(client, headers, iid, preview, *, limit=None, offset=0, severity=None):
    params = {"expected_revision": preview["revision"], "plan_fingerprint": preview["preview"]["plan_fingerprint"],
              "offset": offset, **({"limit": limit} if limit else {}), **({"severity": severity} if severity else {})}
    response = await client.get(f"/payroll/imports/{iid}/rows", params=params, headers=headers)
    assert response.status_code == 200, response.text
    assert len(response.content) <= 2_000_000
    return response.json(), len(response.content)


def _file_shapes(files):
    """Describe generated upload tables; this is not canonical extraction evidence."""
    shapes = []
    for name, payload in files:
        if name.endswith(".csv"):
            rows = list(csv.reader(io.StringIO(payload.decode("utf-8"))))
            header_columns = len(rows[0]) if rows else 0
            data_rows = max(0, len(rows) - 1)
            worksheet_count = 0
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            try:
                worksheet_count = len(workbook.sheetnames)
                rows = list(workbook.active.iter_rows(values_only=True))
                header_columns = len(rows[0]) if rows else 0
                data_rows = max(0, len(rows) - 1)
            finally:
                workbook.close()
        shapes.append({"name": name, "upload_bytes": len(payload), "header_columns": header_columns,
                       "data_rows": data_rows, "table_cells": (data_rows + bool(header_columns)) * header_columns,
                       "worksheets": worksheet_count})
    return shapes


async def source_counts(engine, tenant, period_id):
    """Company Driver/Employee identity, payroll draft and current-rule rows that no preview may change."""
    from sqlalchemy import text

    async with engine.connect() as db:
        return dict((await db.execute(text("""SELECT
            (SELECT count(*) FROM core.drivers WHERE companyid=:cid) AS drivers,
            (SELECT count(*) FROM core.employees WHERE companyid=:cid) AS employees,
            (SELECT count(*) FROM core.driverassignments WHERE companyid=:cid) AS assignments,
            (SELECT count(*) FROM payroll.payrolldraftlines WHERE payrollperiodid=:pid) AS draft_lines,
            (SELECT count(*) FROM import.payrollworkruleheads WHERE companyid=:cid) AS rule_heads"""),
            {"cid": tenant.company_id, "pid": period_id})).mappings().one())


async def canonical_workload_snapshot(engine, tenant, period_id):
    """Capture canonical payroll, source evidence and active heads around a preview-only path."""
    from sqlalchemy import text

    async with engine.connect() as db:
        lines = [dict(row) for row in (await db.execute(text("""
            SELECT d.*
            FROM payroll.PayrollDraftLines AS d
            WHERE d.CompanyID=:cid AND d.PayrollPeriodID=:pid ORDER BY d.DraftLineID
        """), {"cid": tenant.company_id, "pid": period_id})).mappings()]
        evidence = [dict(row) for row in (await db.execute(text("""
            SELECT e.*
            FROM payroll.PayrollPeriodAuditEvidenceEvents AS e
            WHERE e.CompanyID=:cid AND e.PayrollPeriodID=:pid AND e.EvidenceDomain='SOURCE'
            ORDER BY e.PayrollPeriodAuditEvidenceEventID
        """), {"cid": tenant.company_id, "pid": period_id})).mappings()]
        heads = [dict(row) for row in (await db.execute(text("""
            SELECT BranchID, PayDefinitionID, RuleVersionID, Revision
            FROM import.PayrollWorkRuleHeads WHERE CompanyID=:cid ORDER BY BranchID, PayDefinitionID
        """), {"cid": tenant.company_id})).mappings()]
    return {"draft_lines": lines, "evidence": evidence, "rule_heads": heads}


async def import_staging_snapshot(engine, tenant):
    """Count Import staging rows so rejected admission cannot leave partial records."""
    from sqlalchemy import text

    async with engine.connect() as db:
        row = (await db.execute(text("""SELECT
            (SELECT count(*) FROM import.Imports WHERE CompanyID=:cid) AS imports,
            (SELECT count(*) FROM import.FileSources WHERE CompanyID=:cid) AS files,
            (SELECT count(*) FROM import.FileCanonicalData WHERE CompanyID=:cid) AS canonical_files,
            (SELECT count(*) FROM import.PayrollWorkRecords WHERE CompanyID=:cid) AS source_records,
            (SELECT count(*) FROM import.PayrollWorkPreviewGroups WHERE CompanyID=:cid) AS preview_groups
        """), {"cid": tenant.company_id})).mappings().one()
    return dict(row)


def renamed_repeat_files(files, days):
    """Keep all payloads but vary filename date spelling for the completed-repeat case."""
    if len(files) != len(days):
        raise ValueError("Completed-repeat file count must match period days.")
    renamed = [(f"{written_date(day)}.csv", payload) for (_name, payload), day in zip(files, days, strict=True)]
    if len({name for name, _payload in renamed}) != len(renamed):
        raise ValueError("Written date names must be unique within the repeat upload.")
    return renamed


def filename_date_selection(_unit):
    """Let the A08 interpreter resolve written-month names against period days."""
    return {"kind": "filename"}


def date_source_for_unit(scene, file, unit):
    """Choose the same date authority path as current workloads for inferred-name cases."""
    authority = DATE_AUTHORITIES.get(scene)
    if authority == "column":
        return {"kind": "column"}
    if authority:
        return {"kind": authority}
    candidates = unit["inspection"]["filename_date_candidates" if file["type"] == "csv"
                                  else "unit_date_candidates"]
    if len(candidates) != 1:
        raise AssertionError(candidates)
    return {"kind": "filename" if file["type"] == "csv" else "sheet", "value": candidates[0]}


async def measure_completed_repeat(*, client, engine, tenant, period, fixture, baseline_import_id, original_files,
                                   ordinary_path, phases, preview_once, resources):
    """Measure completed-receipt matching after a real Manual edit, without an apply route."""
    from sqlalchemy import text

    baseline_id = baseline_import_id
    fixture_started = time.perf_counter()
    async with engine.begin() as db:
        stored = await db.execute(text("""
            SELECT w.ReviewSummary->'semantic_input_receipt' AS receipt
            FROM import.PayrollWorkImports w
            WHERE w.ImportID=:iid AND w.CompanyID=:cid
        """), {"iid": baseline_id, "cid": tenant.company_id})
        receipt = stored.scalar_one()
        assert receipt["version"] == 1 and len(receipt["targets"]) == len(fixture["definitions"])
        await db.execute(text("""UPDATE import.Imports SET Status='Completed', Revision=Revision+1,
            CompletedByUserID=:uid, CompletedAtUtc=NOW(), UpdatedByUserID=:uid, UpdatedAtUtc=NOW()
            WHERE ImportID=:iid AND CompanyID=:cid"""),
            {"uid": tenant.owner, "iid": baseline_id, "cid": tenant.company_id})

    driver = fixture["drivers"][0]
    day = fixture["days"][0]
    definition = fixture["definitions"][0]
    phases.reset()
    manual_started = time.perf_counter()
    manual = await client.post(f"/payroll/periods/{period['payroll_period_id']}/lines",
        headers=tenant.admin,
        json={"driver_assignment_id": driver["driverassignmentid"], "work_date": day.isoformat(),
              "payroll_period_definition_id": definition.payroll_period_definition_id,
              "quantity": "99", "source_type": "Manual"})
    manual_seconds = time.perf_counter() - manual_started
    assert manual.status_code == 201, manual.text
    before_repeat = await canonical_workload_snapshot(engine, tenant, period["payroll_period_id"])
    fixture_seconds = time.perf_counter() - fixture_started

    repeat_files = renamed_repeat_files(original_files, fixture["days"])
    upload, ingress = await timed_request(client, ordinary_path, tenant.admin,
        lambda: client.post(f"/payroll/periods/{period['payroll_period_id']}/imports",
            files=[("file", (name, data, "application/octet-stream")) for name, data in repeat_files],
            headers={**tenant.admin, "Idempotency-Key": str(uuid4())}), phases)
    import_id = upload["import_id"]
    assert len(upload["layouts"]) == 1

    units = []
    for source_file in upload["files"]:
        for unit in source_file["units"]:
            units.append({"source_unit_id": unit["source_unit_id"], "selected": True,
                          "date_source": filename_date_selection(unit)})
    size, send = json_request(client, "PUT", f"/payroll/imports/{import_id}/interpretation",
                              {"expected_revision": upload["revision"], "units": units}, tenant.admin)
    upload, interpretation = await timed_request(client, ordinary_path, tenant.admin, send, phases)
    interpretation["request_body_bytes"] = size
    interpretation["units"] = len(units)

    headers_by_layout = []
    for layout in upload["layouts"]:
        headers = []
        offset = 0
        while True:
            inspected = await client.get(f"/payroll/imports/{import_id}", params={
                "header_layout_id": layout["layout_id"], "header_offset": offset, "header_limit": 25},
                headers=tenant.admin)
            assert inspected.status_code == 200, inspected.text
            page = inspected.json()["header_page"]
            headers.extend(item["label"] for item in page["items"])
            if page["next_offset"] is None:
                break
            offset = page["next_offset"]
        headers_by_layout.append(headers)
    definition_ids = {item.code: item.payroll_period_definition_id for item in fixture["definitions"]}
    shared, ignored, targets = mapping_requests(
        "current_completed_repeat_csv", upload, headers_by_layout, definition_ids, fixture["definitions"])
    size, send = json_request(client, "PUT", f"/payroll/imports/{import_id}/mapping", {
        "expected_revision": upload["revision"], "shared_columns": shared,
        "ignored_columns": ignored, "targets": targets}, tenant.admin)
    upload, mapping = await timed_request(client, ordinary_path, tenant.admin, send, phases)
    mapping.update({"request_body_bytes": size, "targets": len(targets),
                    "max_request_bytes": resources["max_mapping_request_bytes"]})
    assert size <= resources["max_mapping_request_bytes"]

    preview, preview_metrics = await preview_once(import_id, upload["revision"])
    refreshed, refresh_metrics = await preview_once(import_id, preview["revision"])
    require_timers(preview_metrics, "completed_receipt_snapshot", "matching_completed_receipts")
    require_timers(refresh_metrics, "completed_receipt_snapshot", "matching_completed_receipts")
    notice = refreshed["preview"]["summary"]["repeat_notice"]
    target_count = len(fixture["definitions"])
    assert notice["state"] == "Complete" and notice["matched_target_count"] == target_count
    assert notice["supplied_target_count"] == target_count

    expected_groups = len(fixture["drivers"]) * len(fixture["days"])
    total = expected_groups * target_count
    offset = 0
    pages = []
    no_op_count = 0
    edited_value = None
    edited_page = None
    while offset < expected_groups:
        page_started = time.perf_counter()
        page, response_bytes = await read_rows(client, tenant.admin, import_id, refreshed,
            limit=resources["row_page_max"], offset=offset)
        page_seconds = time.perf_counter() - page_started
        assert page["filtered_count"] == expected_groups
        assert len(page["items"]) == min(resources["row_page_max"], expected_groups - offset)
        for group in page["items"]:
            for definition in fixture["definitions"]:
                key = str(definition.payroll_period_definition_id)
                value = group["values"][key]
                assert value["candidate_count"] == 1 and len(value["candidates"]) == 1, value
                candidate = value["candidates"][0]
                assert candidate["outcome"] == "AlreadyImported" and candidate["action"] == "NoOp", candidate
                assert candidate["quantity"] == "1" and candidate["repeat_input_sha256"], candidate
                no_op_count += 1
                if (group["driver"]["driver_number"] == driver["drivernumber"]
                        and group["work_date"] == day.isoformat()
                        and definition.payroll_period_definition_id == fixture["definitions"][0].payroll_period_definition_id):
                    edited_value = candidate
                    edited_page = {"source_record_count": group["source_record_count"],
                                   "item_count": len(page["items"])}
        pages.append({"offset": offset, "items": len(page["items"]), "seconds": page_seconds,
                      "response_bytes": response_bytes})
        offset += len(page["items"])
    assert no_op_count == total == 65_100, no_op_count
    assert edited_value is not None and edited_page is not None
    assert edited_value["current_quantity"] == "99" and edited_value["current_source_exists"] is True
    assert edited_value["current_provenance"]["source_type"] == "Manual"
    after_repeat = await canonical_workload_snapshot(engine, tenant, period["payroll_period_id"])
    assert after_repeat == before_repeat
    return {"controlled_completed_fixture_seconds": fixture_seconds,
            "manual_edit_seconds": manual_seconds, "manual_edit_status": manual.status_code,
            "incoming_upload_seconds": ingress["request_seconds"], "incoming_interpretation": interpretation,
            "incoming_mapping": mapping, "repeat_preview": preview_metrics, "repeat_refresh": refresh_metrics,
            "repeat_notice": {"state": notice["state"], "matched_target_count": notice["matched_target_count"],
                              "supplied_target_count": notice["supplied_target_count"]},
            "groups": expected_groups, "targets": target_count, "candidate_no_ops": no_op_count,
            "candidate_page_count": len(pages), "candidate_pages": pages,
            "manual_candidate": {"incoming_quantity": edited_value["quantity"],
                                 "current_quantity": edited_value["current_quantity"],
                                 "current_source_type": edited_value["current_provenance"]["source_type"]},
            "source_evidence_events": len(before_repeat["evidence"]),
            "canonical_evidence_and_heads_unchanged": True}


async def measure_partial_overlap(*, client, tenant, fixture, import_id, preview, resources):
    """Check target candidate totals across all rows for alternating compatible layouts."""
    expected_groups = len(fixture["drivers"]) * len(fixture["days"])
    target_counts = Counter()
    groups_seen = 0
    page_metrics = []
    day_indexes = {day.isoformat(): index for index, day in enumerate(fixture["days"])}
    offset = 0
    while offset < expected_groups:
        started = time.perf_counter()
        page, response_bytes = await read_rows(client, tenant.admin, import_id, preview,
            limit=resources["row_page_max"], offset=offset)
        elapsed = time.perf_counter() - started
        assert page["filtered_count"] == expected_groups
        assert page["items"]
        for group in page["items"]:
            groups_seen += 1
            day_index = day_indexes[group["work_date"]]
            for definition_index, definition in enumerate(fixture["definitions"]):
                target_id = str(definition.payroll_period_definition_id)
                present = (definition_index == 0
                           or (day_index % 2 == 0 and definition_index in (1, 2))
                           or (day_index % 2 == 1 and definition_index in (3, 4, 5)))
                value = group["values"].get(target_id)
                if not present:
                    assert value is None, (group["work_date"], target_id, value)
                    continue
                assert value is not None and value["candidate_count"] == 1, (group, target_id)
                candidates = value["candidates"]
                assert len(candidates) == 1 and candidates[0]["quantity"] == "1", (group, target_id)
                assert candidates[0]["outcome"] == "New", candidates[0]
                target_counts[target_id] += 1
        page_metrics.append({"offset": offset, "items": len(page["items"]), "seconds": elapsed,
                             "response_bytes": response_bytes})
        offset += len(page["items"])
    assert groups_seen == expected_groups
    expected_per_target = {
        str(definition.payroll_period_definition_id): (expected_groups if index == 0 else
            len(fixture["drivers"]) * sum(day_index % 2 == (0 if index in (1, 2) else 1)
                                          for day_index in range(len(fixture["days"]))))
        for index, definition in enumerate(fixture["definitions"])
    }
    assert dict(target_counts) == expected_per_target, (target_counts, expected_per_target)
    expected_candidates = sum(expected_per_target.values())
    assert preview["preview"]["counts"]["new"] == expected_candidates
    assert preview["preview"]["counts"]["errors"] == 0
    date_review = preview["preview"]["summary"]["date_review"]
    assert date_review["recognized_day_count"] == len(fixture["days"])
    return {"physical_layouts": 2, "combined_columns": 7, "per_file_physical_columns": [4, 5],
            "driver_day_groups": groups_seen, "candidates": expected_candidates,
            "per_target_candidates": expected_per_target, "recognized_days": date_review["recognized_day_count"],
            "candidate_pages": page_metrics}


def mapping_requests(scene, upload, headers_by_layout, definition_ids, definitions):
    """The full-snapshot mapping body: shared columns, explicit ignores and per-target inputs."""
    if scene == "current_target_mapping_csv":
        plan = target_mapping_plan(definitions)
        signatures = {headers_by_layout[0][item["columns"][0] - 1]: item["signature"]
                      for item in upload["layouts"][0]["inspection"]["duplicate_header_signatures"]}
        def selector(header, ordinal):
            if header not in signatures:
                return {"header": header}
            return {"header": header, "occurrences": [
                {"duplicate_signature": signatures[header], "occurrence_ordinal": ordinal}]}
        targets = []
        for item in plan["targets"]:
            seen = Counter()
            inputs = []
            for header in item["headers"]:
                seen[header] += 1
                inputs.append(selector(header, seen[header]))
            targets.append({"target": f"Definition:{definition_ids[item['code']]}", "transform": item["transform"],
                            "inputs": inputs, "expected_head_revision": 0})
        return ({role: {"header": role} for role in plan["shared"]},
                [{"header": header} for header in plan["ignored"]], targets)
    shared_columns = {}
    target_rules = {}
    for headers in headers_by_layout:
        for role in STRUCTURAL_HEADERS:
            if role in headers:
                shared_columns[role] = {"header": role}
        for header in headers:
            if header in STRUCTURAL_HEADERS or header in target_rules or header not in definition_ids:
                continue
            target_rules[header] = {"target": f"Definition:{definition_ids[header]}",
                "transform": "Direct", "inputs": [{"header": header}], "expected_head_revision": 0}
    if scene == "current_canonical_byte_boundary_csv":
        # The padded headed Notes column is intentionally ignored only after canonical admission.
        return shared_columns, [{"header": "Notes"}], list(target_rules.values())
    return shared_columns, [], list(target_rules.values())


def check_target_mapping(scene_plan, fixture, preview, rows_warned, warned_issues, sample, definition_ids):
    drivers = {item["drivernumber"]: index for index, item in enumerate(fixture["drivers"])}
    flagged = sum(mapping_operand_blank(index) for index in range(len(fixture["drivers"])))
    expected_warnings = flagged * len(fixture["days"])
    counts = preview["preview"]["counts"]
    assert counts["errors"] == 0 and counts["warnings"] == expected_warnings, counts
    assert counts["blank"] == expected_warnings, counts
    assert rows_warned["filtered_count"] == expected_warnings
    assert warned_issues["filtered_count"] == expected_warnings
    assert {item["code"] for item in warned_issues["items"]} == {"IMPORT_SUM_INPUT_MISSING"}
    blankable = next(item for item in scene_plan["targets"] if item["blankable"])
    for item in rows_warned["items"]:
        value = item["values"][str(definition_ids[blankable["code"]])]["candidates"][0]
        assert item["warning_count"] == 1 and value["quantity"] is None and value["is_blank"] is True, item
    checked = 0
    for item in sample["items"]:
        index = drivers[item["driver"]["driver_number"]]
        for target in scene_plan["targets"]:
            candidate = item["values"][str(definition_ids[target["code"]])]["candidates"][0]
            expected = mapping_expected(target, index)
            assert candidate["transform"] == target["transform"], candidate
            assert len(candidate["source_columns"]) == len(target["headers"]), candidate
            if expected is None:
                assert candidate["quantity"] is None and candidate["is_blank"] is True, candidate
            else:
                assert Decimal(candidate["quantity"]) == expected, (target["code"], candidate, expected)
            checked += 1
    assert checked > 0
    return {"direct_targets": sum(item["transform"] == "Direct" for item in scene_plan["targets"]),
            "sum_targets": sum(item["transform"] == "Sum" for item in scene_plan["targets"]),
            "sum_input_columns": sum(len(item["headers"]) for item in scene_plan["targets"]
                                     if item["transform"] == "Sum"),
            "duplicate_header_sum_inputs": sum(len(item["headers"]) for item in scene_plan["targets"]
                                               if len(set(item["headers"])) < len(item["headers"])),
            "ignored_columns": len(scene_plan["ignored"]),
            "blank_sum_operand_groups": expected_warnings, "warnings": counts["warnings"],
            "target_values_checked": checked}


def check_identity(fixture, issue_items, issue_pages, matches):
    drivers = fixture["drivers"]
    per_day = Counter(identity_kind(index) for index in range(len(drivers)))
    expected = {code: per_day[kind] * len(fixture["days"]) for kind, code in IDENTITY_FINDINGS.items()}
    codes = Counter(item["code"] for item in issue_items)
    for code, count in expected.items():
        assert codes[code] == count, (code, count, codes)
    assert fixture["company_drivers"] >= 10 * len(drivers), fixture["company_drivers"]
    assert matches.get(("DriverName", "Proposed"), 0) > 0 and matches.get(("DriverName", "Unresolved"), 0) > 0, matches
    return {"company_drivers": fixture["company_drivers"], "branch_drivers": len(drivers),
            "findings_expected": expected, "issue_codes": dict(codes), "issue_pages": issue_pages,
            "name_matches_by_resolution": {f"{kind}:{resolution}": count for (kind, resolution), count in matches.items()}}


def check_date_review(scene, fixture, review, unit_count, sample_limit):
    days = [day.isoformat() for day in fixture["days"]]
    covered = [day.isoformat() for day in scenario_covered_days(scene, fixture["days"])]
    missing = [day for day in days if day not in covered]
    assert review["recognized_day_count"] == len(covered), review
    assert review["recognized_days"] == covered[:sample_limit]
    assert review["recognized_days_truncated"] is (len(covered) > sample_limit)
    assert review["missing_coverage_count"] == len(missing)
    assert review["missing_coverage_days"] == missing[:sample_limit]
    assert review["missing_coverage_days_truncated"] is (len(missing) > sample_limit)
    assert review["unresolved_row_count"] == 0 and review["overlap_unit_day_count"] == 0, review
    assert review["overlap_unit_days"] == [] and review["overlap_unit_days_truncated"] is False
    assert len(review["units"]) == unit_count
    assert sum(item["recognized_day_count"] for item in review["units"]) == len(covered)
    assert all(item["unresolved_row_count"] == 0 for item in review["units"])
    expects_missing = (scene in DATE_AUTHORITIES_BY_SCENARIO or
                       scene == "current_canonical_byte_boundary_csv")
    assert bool(missing) == expects_missing
    return {"authority": DATE_AUTHORITIES_BY_SCENARIO.get(scene), "sample_limit": sample_limit,
            "period_days": len(days), "supplied_days": len(covered),
            "recognized_day_count": review["recognized_day_count"], "recognized_days_sampled": len(review["recognized_days"]),
            "recognized_days_truncated": review["recognized_days_truncated"],
            "missing_coverage_count": review["missing_coverage_count"], "missing_coverage_days": review["missing_coverage_days"],
            "missing_coverage_days_truncated": review["missing_coverage_days_truncated"],
            "overlap_unit_day_count": review["overlap_unit_day_count"], "units": len(review["units"]),
            "unit_recognized_day_counts": sorted({item["recognized_day_count"] for item in review["units"]})}


async def measure(scene, dsn, *, profile_cleanup=False):
    from sqlalchemy import event, text

    from app.imports.retention import purge_payroll_work_import_details
    from tests.payroll_runtime_fixtures import force_status

    current = scene in CURRENT_POLICY_SCENARIOS
    dimensions = scenario_dimensions(scene)
    async with application(dsn) as (client, engine, _app):
        fixture = await seed_period(client, engine, dsn, drivers=dimensions[0], definitions=dimensions[1],
            days=dimensions[2], outside_drivers=dimensions[3], distinct_names=scene == "current_identity_csv")
        tenant, period = fixture["tenant"], fixture["period"]
        canonical_boundary = None
        if scene == "current_canonical_byte_boundary_csv":
            exact, over, canonical_boundary = canonical_boundary_payloads(fixture)
            files = [(f"{fixture['days'][0].isoformat()}.csv", exact)]
        else:
            exact = over = None
            files = source_files(scene, fixture)
        policy_metadata = None
        if current:
            response = await client.get(
                f"/payroll/periods/{period['payroll_period_id']}/imports/policy", headers=tenant.admin,
            )
            assert response.status_code == 200, response.text
            policy_metadata = response.json()
            assert policy_metadata["product_policy"]["max_company_pay_items"] == 50
            assert policy_metadata["product_policy"]["max_branch_input_concepts"] == 15
            assert policy_metadata["product_policy"]["max_import_combined_columns"] == 40
            assert policy_metadata["product_policy"]["max_import_canonical_file_bytes"] == 1_000_000
            resources = policy_metadata["import_resource_policy"]
            assert resources["max_file_bytes"] == 10_000_000
            assert resources["max_response_bytes"] == 2_000_000
            assert resources["max_numeric_work_digits"] == 16_000
            assert resources["max_mapping_request_bytes"] == 2_000_000
            assert resources["max_identity_lookup_rows"] == 100_000
            assert resources["identity_lookup_batch_rows"] == 512
            assert resources["category_page_default"] == 50 and resources["category_page_max"] == 200
            assert resources["date_review_day_sample"] == 512
            assert resources["row_page_max"] == 100
        async with engine.connect() as db:
            database_temp_before = dict((await db.execute(text("""SELECT temp_files,temp_bytes
                FROM pg_stat_database WHERE datname=current_database()"""))).mappings().one())
        temporary_before = set(Path(tempfile.gettempdir()).iterdir())
        source_counts_before = await source_counts(engine, tenant, period["payroll_period_id"])
        initial_memory = memory_snapshot()
        phases = Phases()
        phases.install()
        sql_count = 0
        sql_seconds = 0
        sql_writes = {}
        cleanup_plans = []
        def transaction_begin(conn):
            conn.info["measured_transaction_begin"] = time.perf_counter()
        def transaction_commit(conn):
            locked = conn.info.pop("measured_import_lock", None)
            started = conn.info.pop("measured_transaction_begin", None)
            if locked is not None and started is not None:
                now = time.perf_counter()
                phases.values["publication_transaction_to_commit_dispatch"] = now - started
                phases.values["period_lock_to_commit_dispatch"] = now - locked
        def before(conn, _cursor, statement, parameters, _context, _many):
            if conn.info.get("cleanup_probe"):
                return statement, parameters
            conn.info["measurement_started"] = time.perf_counter()
            if profile_cleanup and statement.strip().lower().startswith("delete from import.payrollworkpreviewgroups"):
                conn.info["cleanup_probe"] = True
                try:
                    child = conn.exec_driver_sql("""EXPLAIN (ANALYZE, BUFFERS, FORMAT JSON)
                        SELECT 1 FROM ONLY import.payrollworkrecords r
                        WHERE r.previewgroupid=(SELECT previewgroupid FROM import.payrollworkpreviewgroups
                            WHERE importid=$1 AND companyid=$2 LIMIT 1)
                          AND r.importid=$1 AND r.companyid=$2 FOR KEY SHARE OF r""", parameters).scalar_one()
                    cleanup_plans.append({"child_lookup": child})
                finally:
                    conn.info.pop("cleanup_probe", None)
                return "EXPLAIN (ANALYZE, BUFFERS, FORMAT JSON) " + statement, parameters
            return statement, parameters
        def after(conn, _cursor, statement, _params, _context, _many):
            nonlocal sql_count, sql_seconds
            if conn.info.get("cleanup_probe"):
                return
            sql_count += 1
            elapsed = time.perf_counter() - conn.info.pop("measurement_started")
            sql_seconds += elapsed
            normalized = statement.strip().lower()
            if profile_cleanup and normalized.startswith("explain") and "delete from import.payrollworkpreviewgroups" in normalized:
                cleanup_plans.append({"group_delete": [row[0] for row in _cursor.fetchall()]})
            if normalized.startswith(("delete from import.", "update import.", "insert into import.")):
                key = " ".join(normalized.split()[:3])
                sql_writes[key] = sql_writes.get(key, 0) + elapsed
        event.listen(engine.sync_engine, "before_cursor_execute", before, retval=True)
        event.listen(engine.sync_engine, "after_cursor_execute", after)
        event.listen(engine.sync_engine, "begin", transaction_begin)
        event.listen(engine.sync_engine, "commit", transaction_commit)
        try:
            ordinary_path = f"/payroll/periods/{period['payroll_period_id']}"
            async def preview_once(iid, revision):
                nonlocal sql_count, sql_seconds
                sql_count = 0
                sql_seconds = 0
                sql_writes.clear()
                body, metrics = await timed_request(client, ordinary_path, tenant.admin,
                    lambda: client.post(f"/payroll/imports/{iid}/preview", json={"expected_revision": revision},
                                        headers=tenant.admin), phases)
                metrics["sqlalchemy_calls_including_probe_reads"] = sql_count
                metrics["sqlalchemy_seconds_including_probe_reads"] = sql_seconds
                metrics["sql_write_seconds_by_statement"] = dict(sql_writes)
                return body, metrics
            upload, ingress = await timed_request(client, ordinary_path, tenant.admin,
                lambda: client.post(f"/payroll/periods/{period['payroll_period_id']}/imports",
                    files=[("file", (name, data, "application/octet-stream")) for name, data in files],
                    headers={**tenant.admin, "Idempotency-Key": str(uuid4())}), phases)
            iid = upload["import_id"]
            header_structure = upload.get("header_structure") if current else None
            unit_count = sum(len(file["units"]) for file in upload["files"])
            if current:
                assert header_structure and header_structure["available"] is True
                assert len(header_structure["headed_column_counts"]) == unit_count
                physical_counts = {item["count"] for item in header_structure["headed_column_counts"]}
                if scene == "current_partial_overlap_csv":
                    assert header_structure["combined_column_count"] == 7
                    assert physical_counts == {4, 5}
                    assert len(upload["layouts"]) == 2
                else:
                    first_inspection = upload["files"][0]["units"][0]["inspection"]
                    expected_columns = first_inspection["headed_column_count"]
                    assert header_structure["combined_column_count"] == expected_columns
                    assert physical_counts == {expected_columns}
                    assert len(upload["layouts"]) == 1
                assert header_structure["combined_column_count"] <= policy_metadata["product_policy"]["max_import_combined_columns"]
            authority = DATE_AUTHORITIES.get(scene)
            interpretation = None
            if scene.startswith("monthly_daily") or current:
                units = []
                for file in upload["files"]:
                    for unit in file["units"]:
                        source = date_source_for_unit(scene, file, unit)
                        units.append({"source_unit_id": unit["source_unit_id"], "selected": True, "date_source": source})
                size, send = json_request(client, "PUT", f"/payroll/imports/{iid}/interpretation",
                    {"expected_revision": upload["revision"], "units": units}, tenant.admin)
                upload, interpretation = await timed_request(client, ordinary_path, tenant.admin, send, phases)
                interpretation["request_body_bytes"] = size
                interpretation["units"] = len(units)
            definition_ids = {definition.code: definition.payroll_period_definition_id for definition in fixture["definitions"]}
            headers_by_layout = []
            for layout in upload["layouts"]:
                headers = []
                offset = 0
                while True:
                    inspected = await client.get(f"/payroll/imports/{iid}", params={
                        "header_layout_id": layout["layout_id"], "header_offset": offset, "header_limit": 25},
                        headers=tenant.admin)
                    assert inspected.status_code == 200, inspected.text
                    page = inspected.json()["header_page"]
                    headers.extend(item["label"] for item in page["items"])
                    if page["next_offset"] is None:
                        break
                    offset = page["next_offset"]
                headers_by_layout.append(headers)
            shared_columns, ignored_columns, targets = mapping_requests(
                scene, upload, headers_by_layout, definition_ids, fixture["definitions"])
            size, send = json_request(client, "PUT", f"/payroll/imports/{iid}/mapping", {
                "expected_revision": upload["revision"], "shared_columns": shared_columns,
                "ignored_columns": ignored_columns, "targets": targets}, tenant.admin)
            upload, mapping = await timed_request(client, ordinary_path, tenant.admin, send, phases)
            mapping.update({"request_body_bytes": size, "targets": len(targets),
                            "sum_targets": sum(item["transform"] == "Sum" for item in targets),
                            "input_selectors": sum(len(item["inputs"]) for item in targets),
                            "shared_columns": len(shared_columns), "ignored_columns": len(ignored_columns)})
            if current:
                mapping["max_request_bytes"] = resources["max_mapping_request_bytes"]
                assert size <= resources["max_mapping_request_bytes"]
                require_timers(mapping, "mapping_save")
            checks = {}
            definition_count = len(fixture["definitions"])
            totals = repeated_totals(len(fixture["drivers"]), len(fixture["days"])) if scene == "current_repeated_rows_csv" else None
            if totals:
                # RejectRepeated is the default: repeated Driver/day rows block preparation of every quantity.
                rejected, rejected_metrics = await preview_once(iid, upload["revision"])
                reject_counts = rejected["preview"]["counts"]
                assert reject_counts["errors"] == totals["rows"] and reject_counts["new"] == 0, reject_counts
                assert reject_counts["invalid"] == totals["rows"] * definition_count, reject_counts
                require_timers(rejected_metrics, "row_handling_aggregation")
                assert "aggregate_contributions" not in rejected_metrics["phases_seconds"]
                category_pages, unique_categories = [], {}
                page_size = resources["category_page_default"]
                offset = 0
                while True:
                    size, send = json_request(client, "POST", f"/payroll/imports/{iid}/categories/inspect", {
                        "expected_revision": rejected["revision"], "column": {"header": CATEGORY_HEADER},
                        "offset": offset, "limit": page_size}, tenant.admin)
                    page, page_metrics = await timed_request(client, ordinary_path, tenant.admin, send, phases)
                    page_metrics.update({"request_body_bytes": size, "items": len(page["items"])})
                    require_timers(page_metrics, "category_evidence", "_materialize_batch")
                    if not category_pages:
                        first_page_head = page["items"][:1]
                    category_pages.append(page_metrics)
                    fresh = [item for item in page["items"]
                             if (item["category"]["type"], item["category"]["value"]) not in unique_categories]
                    unique_categories.update({(item["category"]["type"], item["category"]["value"]): item
                                              for item in page["items"]})
                    if page["next_offset"] is None:
                        break
                    assert fresh, "category paging repeated items without progress"
                    offset = page["next_offset"]
                categories = list(unique_categories.values())
                assert page["suitable"] is True and page["invalid_value_count"] == 0
                assert len(categories) == CATEGORY_COUNT == page["distinct_count"]
                assert sum(item["count"] for item in categories) == totals["rows"] == page["inspected_row_count"]
                # offset and limit must narrow each page: full pages first, then the remainder.
                probe, probe_metrics = await timed_request(client, ordinary_path, tenant.admin, json_request(
                    client, "POST", f"/payroll/imports/{iid}/categories/inspect", {
                        "expected_revision": rejected["revision"], "column": {"header": CATEGORY_HEADER},
                        "offset": page_size, "limit": page_size}, tenant.admin)[1], phases)
                paging_probe = {"requested_offset": page_size, "requested_limit": page_size,
                                "items_returned": len(probe["items"]), "limit_honored": len(probe["items"]) <= page_size,
                                "offset_applied": probe["items"][:1] != first_page_head,
                                "request_seconds": probe_metrics["request_seconds"],
                                "response_bytes": probe_metrics["response_bytes"]}
                expected_pages = [page_size] * (CATEGORY_COUNT // page_size) + (
                    [CATEGORY_COUNT % page_size] if CATEGORY_COUNT % page_size else [])
                assert [item["items"] for item in category_pages] == expected_pages, category_pages
                assert paging_probe["limit_honored"] and paging_probe["offset_applied"], paging_probe
                assert all(item["observed"] for item in categories)
                size, send = json_request(client, "PUT", f"/payroll/imports/{iid}/row-handling",
                    {"expected_revision": rejected["revision"], "mode": "SumRows"}, tenant.admin)
                summed_state, sum_handling = await timed_request(client, ordinary_path, tenant.admin, send, phases)
                sum_handling["request_body_bytes"] = size
                require_timers(sum_handling, "mapping_save", "_materialize_batch")
                summed, summed_metrics = await preview_once(iid, summed_state["revision"])
                sum_counts = summed["preview"]["counts"]
                assert sum_counts["errors"] == 0 and sum_counts["new"] == totals["groups"] * definition_count, sum_counts
                require_timers(summed_metrics, "row_handling_aggregation", "aggregate_contributions")
                summed_rows, summed_rows_bytes = await read_rows(client, tenant.admin, iid, summed, limit=100)
                days = {day.isoformat(): index for index, day in enumerate(fixture["days"])}
                numbers = {item["drivernumber"]: index for index, item in enumerate(fixture["drivers"])}
                def check_groups(page_json, *, filtered):
                    checked = 0
                    for item in page_json["items"]:
                        expected = repeated_expected(numbers[item["driver"]["driver_number"]], days[item["work_date"]],
                                                     definition_count, filtered=filtered)
                        assert expected is not None, item
                        records = item["source_records"]
                        # Excluded rows leave the Driver/day group; only included rows contribute.
                        rows_in_group = sum(not (filtered and category_excluded(repeated_category(
                            numbers[item["driver"]["driver_number"]], days[item["work_date"]], repeat)))
                                            for repeat in range(REPEATS))
                        assert item["source_record_count"] == rows_in_group == len(records), item
                        assert item["source_records_sampled"] is False
                        for position, definition in enumerate(fixture["definitions"]):
                            key = str(definition.payroll_period_definition_id)
                            prepared = [candidate for candidate in item["values"][key]["candidates"]
                                        if candidate["quantity"] is not None]
                            assert len(prepared) == 1, item["values"][key]
                            assert Decimal(prepared[0]["quantity"]) == expected[position], (prepared, expected)
                            kept = [record["contributions"][key] for record in records if not record["proposed_exclusion"]]
                            assert sum(Decimal(part["contribution_quantity"]) for part in kept) == expected[position]
                            assert sum(part["is_candidate"] for part in kept) == 1, kept
                        checked += 1
                    assert checked > 0
                    return checked
                sum_checked = check_groups(summed_rows, filtered=False)
                # FilterThenSum needs an explicit Include/Exclude decision for every observed category.
                choices = [{"category": item["category"], "decision": "Exclude" if category_excluded(
                    int(item["category"]["value"].split()[1]) - 1) else "Include"} for item in categories]
                size, send = json_request(client, "PUT", f"/payroll/imports/{iid}/row-handling", {
                    "expected_revision": summed["revision"], "mode": "FilterThenSum",
                    "filter_column": {"header": CATEGORY_HEADER}, "choices": choices}, tenant.admin)
                upload, filter_handling = await timed_request(client, ordinary_path, tenant.admin, send, phases)
                filter_handling["request_body_bytes"] = size
                filter_handling["max_request_bytes"] = resources["max_mapping_request_bytes"]
                assert size <= resources["max_mapping_request_bytes"]
                require_timers(filter_handling, "mapping_save", "_materialize_batch", "category_evidence")
                checks["repeated_rows"] = {
                    "rows": totals["rows"], "driver_day_groups": totals["groups"], "categories": len(categories),
                    "choices": len(choices), "excluded_categories": sum(item["decision"] == "Exclude" for item in choices),
                    "reject_repeated": {"preview": rejected_metrics, "counts": reject_counts},
                    "categories_inspect": {"pages": len(category_pages), "page_size": page_size,
                                           "items_per_page": [item["items"] for item in category_pages],
                                           "paging_probe": paging_probe,
                                           "request_seconds": [item["request_seconds"] for item in category_pages],
                                           "total_request_seconds": sum(item["request_seconds"] for item in category_pages),
                                           "response_bytes": [item["response_bytes"] for item in category_pages],
                                           "page_metrics": category_pages},
                    "sum_rows": {"row_handling": sum_handling, "preview": summed_metrics, "counts": sum_counts,
                                 "rows_page_bytes": summed_rows_bytes, "groups_checked": sum_checked},
                    "filter_then_sum": {"row_handling": filter_handling}}
            sql_count = 0
            sql_seconds = 0
            sql_writes.clear()
            preview, initial = await preview_once(iid, upload["revision"])
            refreshed, refresh = await preview_once(iid, preview["revision"])
            for key in ("mapping_fingerprint", "content_fingerprint", "plan_fingerprint", "counts"):
                assert refreshed["preview"][key] == preview["preview"][key], key
            scenario_checks = {}
            if scene == "current_partial_overlap_csv":
                scenario_checks["partial_column_overlap"] = await measure_partial_overlap(
                    client=client, tenant=tenant, fixture=fixture, import_id=iid,
                    preview=refreshed, resources=resources)
            elif scene == "current_canonical_byte_boundary_csv":
                assert canonical_boundary is not None and over is not None
                assert policy_metadata["product_policy"]["max_import_canonical_file_bytes"] == canonical_boundary["limit_bytes"]
                staging_before_rejection = await import_staging_snapshot(engine, tenant)
                canonical_before_rejection = await canonical_workload_snapshot(
                    engine, tenant, period["payroll_period_id"])
                rejected_started = time.perf_counter()
                rejected = await client.post(
                    f"/payroll/periods/{period['payroll_period_id']}/imports",
                    files=[("file", ("accepted-boundary.csv", exact, "application/octet-stream")),
                           ("file", ("one-byte-over.csv", over, "application/octet-stream"))],
                    headers={**tenant.admin, "Idempotency-Key": str(uuid4())})
                rejected_seconds = time.perf_counter() - rejected_started
                assert rejected.status_code == 422, rejected.text
                failure = rejected.json()["detail"]
                assert failure["code"] == "IMPORT_CANONICAL_FILE_TOO_LARGE", failure
                staging_after_rejection = await import_staging_snapshot(engine, tenant)
                canonical_after_rejection = await canonical_workload_snapshot(
                    engine, tenant, period["payroll_period_id"])
                assert staging_after_rejection == staging_before_rejection
                assert canonical_after_rejection == canonical_before_rejection
                scenario_checks["canonical_boundary"] = {
                    **canonical_boundary, "effective_product_limit_bytes": policy_metadata["product_policy"][
                        "max_import_canonical_file_bytes"],
                    "period_days": len(fixture["days"]),
                    "supplied_days": len(scenario_covered_days(scene, fixture["days"])),
                    "missing_coverage_count": (len(fixture["days"])
                        - len(scenario_covered_days(scene, fixture["days"]))),
                    "accepted_preview_counts": refreshed["preview"]["counts"],
                    "accepted_preview_seconds": initial["request_seconds"],
                    "accepted_refresh_seconds": refresh["request_seconds"],
                    "valid_plus_over_limit_batch_status": rejected.status_code,
                    "valid_plus_over_limit_batch_error": failure["code"],
                    "rejected_request_seconds": rejected_seconds,
                    "staging_unchanged_on_rejection": True,
                    "payroll_evidence_and_heads_unchanged_on_rejection": True}
            if current:
                for metrics in (initial, refresh):
                    require_timers(metrics, "date_preparation", "date_interpretation", "identity_requested_keys",
                                   "identity_row_resolution", "row_handling_aggregation",
                                   absent=(() if scene == "current_identity_csv" else ("identity_company_name_lookup",)))
                    if scene == "current_identity_csv":
                        require_timers(metrics, "identity_company_name_lookup")
                    if totals:
                        require_timers(metrics, "aggregate_contributions")
            params = {"cid": tenant.company_id, "iid": iid, "pid": period["payroll_period_id"]}
            async with engine.connect() as db:
                counts = dict((await db.execute(text("""SELECT
                    (SELECT count(*) FROM import.payrollworkrecords WHERE importid=:iid) AS normalized_rows,
                    (SELECT count(*) FROM import.payrollworkpreviewgroups WHERE importid=:iid) AS grid_rows,
                    (SELECT count(*) FROM import.importissues WHERE importid=:iid) AS relational_issues,
                    (SELECT count(*) FROM payroll.payrolldraftlines WHERE payrollperiodid=:pid) AS canonical_lines,
                    (SELECT count(*) FROM payroll.payrollperiodauditevidenceevents WHERE payrollperiodid=:pid AND evidencedomain='SOURCE') AS source_events,
                    (SELECT count(*) FROM import.filesources WHERE importid=:iid) AS files,
                    (SELECT count(*) FROM import.filesourceunits WHERE importid=:iid) AS units,
                    (SELECT count(*) FROM import.sourcelayouts WHERE importid=:iid) AS layouts"""), params)).mappings().one())
                canonical_files = []
                if current:
                    counts["pending_rules"] = await db.scalar(text(
                        "SELECT count(*) FROM import.PayrollWorkPendingRules WHERE ImportID=:iid"), params)
                    canonical_files = [dict(row) for row in (await db.execute(text("""
                        SELECT f.fileordinal AS ordinal, f.originalfilename AS filename,
                               f.filesizebytes AS uploaded_bytes, c.FormatVersion AS format_version,
                               c.ByteCount AS canonical_bytes, c.HashSha256 AS sha256
                        FROM import.FileSources f
                        LEFT JOIN import.FileCanonicalData c
                          ON c.FileSourceID=f.FileSourceID AND c.ImportID=f.ImportID AND c.CompanyID=f.CompanyID
                        WHERE f.ImportID=:iid AND f.CompanyID=:cid ORDER BY f.FileOrdinal
                    """), params)).mappings()]
                    counts["canonical_file_data"] = sum(row["format_version"] is not None for row in canonical_files)
                    counts["canonical_bytes"] = sum(row["canonical_bytes"] or 0 for row in canonical_files)
                    assert len(canonical_files) == len(files)
                    assert all(row["format_version"] == 1 and row["canonical_bytes"] is not None and row["canonical_bytes"] > 0 and row["sha256"]
                               for row in canonical_files)
                    assert all(row["canonical_bytes"] <= policy_metadata["product_policy"]["max_import_canonical_file_bytes"]
                               for row in canonical_files)
                    if scene == "current_canonical_byte_boundary_csv":
                        assert [row["canonical_bytes"] for row in canonical_files] == [
                            canonical_boundary["accepted_canonical_bytes"]]
                identity_matches = {}
                if scene == "current_identity_csv":
                    identity_matches = {(row[0], row[1]): row[2] for row in (await db.execute(text("""
                        SELECT identifierkind, resolution, count(*) FROM import.payrollworkdrivermatches
                        WHERE importid=:iid GROUP BY 1, 2"""), params))}
            assert counts["canonical_lines"] == counts["source_events"] == 0
            if current:
                source_counts_after = await source_counts(engine, tenant, period["payroll_period_id"])
                assert source_counts_after == source_counts_before, (source_counts_before, source_counts_after)
                assert source_counts_after["draft_lines"] == 0
                assert source_counts_after["rule_heads"] == source_counts_before["rule_heads"]
            query = {"expected_revision": refreshed["revision"], "plan_fingerprint": refreshed["preview"]["plan_fingerprint"]}
            response_sizes = {}
            for name in ("rows", "issues", "source-records"):
                response = await client.get(f"/payroll/imports/{iid}/{name}", params=query, headers=tenant.admin)
                assert response.status_code == 200, response.text
                response_sizes[name] = len(response.content)
                assert len(response.content) <= 2_000_000
            if current:
                max_rows, response_sizes["rows_max_page"] = await read_rows(
                    client, tenant.admin, iid, refreshed, limit=resources["row_page_max"])
                response_sizes["rows_max_page_items"] = len(max_rows["items"])
            if scene == "current_target_mapping_csv":
                plan = target_mapping_plan(fixture["definitions"])
                warned_rows, _ = await read_rows(client, tenant.admin, iid, refreshed, severity="Warning", limit=100)
                warned_issues = await client.get(f"/payroll/imports/{iid}/issues", params={
                    **query, "severity": "Warning", "limit": 200}, headers=tenant.admin)
                assert warned_issues.status_code == 200, warned_issues.text
                checks["target_mapping"] = check_target_mapping(
                    plan, fixture, refreshed, warned_rows, warned_issues.json(), max_rows, definition_ids)
                checks["target_mapping"]["mapping_request_body_bytes"] = mapping["request_body_bytes"]
            if scene == "current_identity_csv":
                issue_items, issue_pages = await get_pages(client, f"/payroll/imports/{iid}/issues", query,
                                                           tenant.admin, limit=resources["issue_page_max"])
                checks["identity"] = check_identity(fixture, issue_items, issue_pages, identity_matches)
                checks["identity"]["name_lookup_seconds"] = {
                    "preview": initial["phases_seconds"]["identity_company_name_lookup"],
                    "refresh": refresh["phases_seconds"]["identity_company_name_lookup"]}
            if scene == "current_repeated_rows_csv":
                filtered_rows, filtered_bytes = await read_rows(client, tenant.admin, iid, refreshed, limit=100)
                assert refreshed["preview"]["counts"]["errors"] == 0
                assert refreshed["preview"]["counts"]["prepared_excluded_rows"] == totals["excluded_rows"]
                assert refreshed["preview"]["counts"]["new"] == totals["visible_groups"] * definition_count, \
                    refreshed["preview"]["counts"]
                filter_checked = check_groups(filtered_rows, filtered=True)
                assert filtered_rows["filtered_count"] == totals["visible_groups"], filtered_rows["filtered_count"]
                checks["repeated_rows"]["filter_then_sum"].update({
                    "counts": refreshed["preview"]["counts"], "rows_page_bytes": filtered_bytes,
                    "rows_page_items": len(filtered_rows["items"]), "groups_checked": filter_checked,
                    "visible_groups_listed": filtered_rows["filtered_count"], "excluded_rows": totals["excluded_rows"],
                    "visible_groups": totals["visible_groups"]})
            if current:
                review = refreshed["preview"]["summary"]["date_review"]
                checks["date_review"] = check_date_review(scene, fixture, review, unit_count, resources["date_review_day_sample"])
                first_rows, _ = await read_rows(client, tenant.admin, iid, refreshed, limit=1)
                last_rows, _ = await read_rows(client, tenant.admin, iid, refreshed, limit=1,
                                               offset=max(first_rows["filtered_count"] - 1, 0))
                reasons = {"first": first_rows["items"][0]["source_records"][0]["date_interpretation"]["reason"],
                           "last": last_rows["items"][0]["source_records"][0]["date_interpretation"]["reason"]}
                inferred = "period_year_inferred"
                if authority:
                    assert reasons == ({"first": "explicit_date", "last": inferred} if authority == "column" else
                                       {"first": inferred, "last": inferred}), reasons
                checks["date_review"]["row_interpretation_reasons"] = reasons
                checks["date_review"]["date_interpretation_calls"] = {"preview": initial["phase_calls"]["date_interpretation"],
                                                                      "refresh": refresh["phase_calls"]["date_interpretation"]}
            # All baseline pages and checks must use the editable preview's exact
            # revision before the controlled fixture seals it as Completed.
            if scene == "current_completed_repeat_csv":
                scenario_checks["completed_repeat"] = await measure_completed_repeat(
                    client=client, engine=engine, tenant=tenant, period=period, fixture=fixture,
                    baseline_import_id=iid, original_files=files, ordinary_path=ordinary_path,
                    phases=phases, preview_once=preview_once, resources=resources)
            force_status(tenant, period["payroll_period_id"], "Locked")
            sql_writes.clear()
            started = time.perf_counter()
            async with engine.begin() as db:
                await purge_payroll_work_import_details(db, company_id=tenant.company_id,
                    period_id=period["payroll_period_id"], actor_id=tenant.owner)
            purge = time.perf_counter() - started
            purge_statements = dict(sql_writes)
            async with engine.connect() as db:
                assert await db.scalar(text("SELECT count(*) FROM import.payrollworkrecords WHERE importid=:iid"), params) == 0
                assert await db.scalar(text("SELECT count(*) FROM import.filesources WHERE importid=:iid AND filecontent IS NOT NULL"), params) == 0
                database_temp_after = dict((await db.execute(text("""SELECT temp_files,temp_bytes
                    FROM pg_stat_database WHERE datname=current_database()"""))).mappings().one())
            file_shapes = _file_shapes(files) if current else None
            return {"scenario": scene, "python": platform.python_version(), "platform": platform.platform(),
                "transport": "HTTPX ASGITransport, real disposable PostgreSQL, production guards enabled",
                "dimensions": {"drivers": dimensions[0], "definitions": dimensions[1],
                               "ordinary_definitions": dimensions[1],
                               "days": dimensions[2], "files": len(files),
                               "company_drivers": fixture["company_drivers"],
                               "worksheets_per_file": (max(item["worksheets"] for item in file_shapes) if current else None)},
                "file_bytes": sum(len(data) for _, data in files), "initial_memory": initial_memory,
                "file_shapes": file_shapes,
                "ingress": ingress, "interpretation": interpretation, "mapping": mapping if current else None,
                "preview": initial, "refresh": refresh, "counts": counts,
                "header_structure": header_structure,
                "canonical_files": canonical_files if current else None,
                "preview_counts": preview["preview"]["counts"], "purge_seconds": purge,
                "purge_sql_write_seconds": purge_statements,
                "page_response_bytes": response_sizes, "final_memory": memory_snapshot(),
                "capability_checks": checks if current else None,
                "database_temp_before": database_temp_before, "database_temp_after": database_temp_after,
                "cleanup_plans": cleanup_plans,
                "acceptance_checks": scenario_checks,
                "temporary_files_left": sorted(str(path) for path in set(Path(tempfile.gettempdir()).iterdir()) - temporary_before),
                "effective_policy": policy_metadata,
                "coverage": scenario_coverage(scene) if current else None}
        finally:
            phases.restore()
            event.remove(engine.sync_engine, "before_cursor_execute", before)
            event.remove(engine.sync_engine, "after_cursor_execute", after)
            event.remove(engine.sync_engine, "begin", transaction_begin)
            event.remove(engine.sync_engine, "commit", transaction_commit)


def scenario_coverage(scene):
    """What this workload exercises; final application and deployment acceptance stay false."""
    return {"legal_policy_shape": True, "canonical_extraction_and_admission": True,
            "canonical_byte_measurement": True, "combined_column_enforcement": True,
            "per_target_mapping_direct_and_sum": scene == "current_target_mapping_csv",
            "driver_identity_discovery": scene == "current_identity_csv",
            "date_authority_inference_and_review": scene in DATE_AUTHORITIES_BY_SCENARIO,
            "repeated_rows_and_categories": scene == "current_repeated_rows_csv",
            "completed_input_repeat_after_manual_edit": scene == "current_completed_repeat_csv",
            "partial_column_overlap": scene == "current_partial_overlap_csv",
            "canonical_byte_boundary_admission": scene == "current_canonical_byte_boundary_csv",
            "semantic_repeat_recognition": scene == "current_completed_repeat_csv",
            "final_revised_contract_acceptance": False,
            "not_covered": NOT_COVERED}


def run_scenarios(scenes, output):
    """Run each isolated workload and preserve both live progress and JSON artifacts."""
    results = []
    for index, scene in enumerate(scenes, start=1):
        path = output.with_name(output.stem + "-" + scene + ".json")
        print(f"Starting {scene} ({index}/{len(scenes)})", flush=True)
        subprocess.run([sys.executable, "-m", "scripts.benchmark_import_files", "--scenario", scene,
                        "--output", str(path)], check=True, env=os.environ.copy())
        results.append(json.loads(path.read_text(encoding="utf-8")))
        print(f"Completed {scene} ({index}/{len(scenes)})", flush=True)
    output.write_text(json.dumps(results, indent=2), encoding="utf-8")


def scenarios_for_selection(selection):
    """Expand only named scenario groups; keep the current acceptance set stable."""
    return SCENARIO_GROUPS.get(selection, (selection,))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--scenario", choices=(*SCENARIOS, *CURRENT_SCENARIOS,
                                                 *SINGLE_DAY_SCENARIOS, "current", "single_day", "all"),
                        default="current", help="Current supported shapes by default; named legacy presets are reference workloads.")
    parser.add_argument("--profile-cleanup", action="store_true",
                        help="Instrument real group deletion/FK lookup with EXPLAIN ANALYZE (extra overhead).")
    parser.add_argument("--output", type=Path, default=ROOT / ".artifacts/import-batch-benchmark.json")
    args = parser.parse_args()
    if args.profile_cleanup and args.scenario in SCENARIO_GROUPS:
        parser.error("Use --profile-cleanup with one explicit scenario.")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.scenario in SCENARIO_GROUPS:
        run_scenarios(scenarios_for_selection(args.scenario), args.output)
    else:
        with private_database() as dsn:
            result = asyncio.run(measure(args.scenario, dsn, profile_cleanup=args.profile_cleanup))
        args.output.write_text(json.dumps(result, indent=2), encoding="utf-8")
        print(json.dumps(result), flush=True)


if __name__ == "__main__":
    main()
