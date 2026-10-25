from __future__ import annotations

import io
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from tests.builders.workforce import create_workforce_driver
from tests.compensation_api_fixtures import (  # noqa: F401 - register isolated API fixtures
    p3c_application,
    p3c_database_engine,
    p3c_http_client,
)
from tests.compensation_database_fixtures import p3b_database  # noqa: F401
from tests.payroll_runtime_fixtures import (
    assign_weekly_setup,
    build_payroll_tenant,
    create_period,
    get_grid,
    make_driver_self_user,
    make_user,
    open_period_with_grid,
    rated_definition,
    seed_status_key,
)

pytestmark = pytest.mark.asyncio


@pytest.fixture(name="tenant")
def payroll_tenant(p3b_dsn):
    return build_payroll_tenant(p3b_dsn)


def _file_body(data: bytes, filename: str, *, boundary: str = "import-a2-preview") -> bytes:
    content_type = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    if filename.endswith(".xlsx") else "text/csv")
    return (
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; "
        f"filename=\"{filename}\"\r\nContent-Type: {content_type}\r\n\r\n".encode()
        + data
        + f"\r\n--{boundary}--\r\n".encode()
    )


async def _period(client, engine, tenant):
    period, _grid = await open_period_with_grid(client, engine, tenant)
    return period


async def _period_with_mask(client, engine, tenant, mask: int):
    await rated_definition(client, tenant, rate="25")
    await assign_weekly_setup(engine, tenant, tenant.branch_a, mask=mask)
    period = await create_period(client, tenant, tenant.branch_a)
    await get_grid(client, tenant, period)
    return period


async def _upload(client, tenant, period, data: bytes, *, filename="payroll.csv", headers=None):
    return await client.post(
        f"/payroll/periods/{period['payroll_period_id']}/imports",
        content=_file_body(data, filename),
        headers={
            **(headers or tenant.admin),
            "Content-Type": "multipart/form-data; boundary=import-a2-preview",
            "Idempotency-Key": str(uuid4()),
        },
    )


async def _definition_target(client, tenant, import_id: int) -> str:
    response = await client.get(f"/payroll/imports/{import_id}", headers=tenant.admin)
    assert response.status_code == 200, response.text
    options = [item["target"] for item in response.json()["mapping_targets"]
               if item["target"].startswith("Definition:")]
    assert len(options) == 1, options
    return options[0]


async def _map(client, tenant, import_id: int, target: str, *, revision: int = 1):
    return await client.put(
        f"/payroll/imports/{import_id}/mapping",
        json={"expected_revision": revision,
              "column_mapping": {"1": "DriverNumber", "2": "WorkDate", "3": target}},
        headers=tenant.admin,
    )


async def _payroll_inputs(engine, tenant, period, driver_assignment_id=None):
    async with engine.connect() as conn:
        row = (await conn.execute(text("""
            SELECT d.drivernumber, e.employeekey, p.drivernamesnapshot, p.eligibleworkdates
            FROM payroll.payrollperioddrivereligibility p
            JOIN core.driverassignments a ON a.driverassignmentid = p.driverassignmentid
                AND a.companyid = p.companyid AND a.branchid = p.branchid
            JOIN core.drivers d ON d.driverid = a.driverid AND d.companyid = a.companyid
            JOIN core.employees e ON e.employeeid = a.employeeid AND e.companyid = a.companyid
            WHERE p.companyid = :cid AND p.payrollperiodid = :pid AND a.driverassignmentid = :aid
        """), {"cid": tenant.company_id, "pid": period["payroll_period_id"],
              "aid": driver_assignment_id or tenant.driver_a})).mappings().one()
        dates = sorted(row["eligibleworkdates"] or [])
        assert dates, "The period fixture must include at least one eligible date for Driver A."
        return {**dict(row), "dates": dates}


async def _payroll_snapshot(engine, period_id: int):
    tables = (
        "payroll.payrollperiods",
        "payroll.payrollperioddays",
        "payroll.payrollperioddefinitions",
        "payroll.payrollperioddrivereligibility",
        "payroll.payrolldraftlines",
        "payroll.payrollperioddriverdayentrystate",
        "payroll.payrollbonusevents",
        "payroll.payrollperiodauditevidenceevents",
        "payroll.payrollperiodauditevidencecoverage",
        "payroll.payrollcalculationsnapshots",
        "payroll.payrollfinallines",
    )
    async with engine.connect() as conn:
        snapshot = {}
        for table in tables:
            snapshot[table] = await conn.scalar(text(f"""
                SELECT COALESCE(jsonb_agg(to_jsonb(t) ORDER BY to_jsonb(t)::text), '[]'::jsonb)::text
                FROM {table} t WHERE payrollperiodid = :pid
            """), {"pid": period_id})
        snapshot["payroll.payrollcalculationdrivertotals"] = await conn.scalar(text("""
            SELECT COALESCE(jsonb_agg(to_jsonb(t) ORDER BY to_jsonb(t)::text), '[]'::jsonb)::text
            FROM payroll.payrollcalculationdrivertotals t
            JOIN payroll.payrollcalculationsnapshots s USING (payrollcalculationsnapshotid)
            WHERE s.payrollperiodid = :pid
        """), {"pid": period_id})
        snapshot["payroll.payrollcalculationsnapshotcomponents"] = await conn.scalar(text("""
            SELECT COALESCE(jsonb_agg(to_jsonb(c) ORDER BY to_jsonb(c)::text), '[]'::jsonb)::text
            FROM payroll.payrollcalculationsnapshotcomponents c
            JOIN payroll.payrollcalculationsnapshots s USING (payrollcalculationsnapshotid)
            WHERE s.payrollperiodid = :pid
        """), {"pid": period_id})
        snapshot["payroll.payrollcalculationsnapshotusedratedefinitions"] = await conn.scalar(text("""
            SELECT COALESCE(jsonb_agg(to_jsonb(r) ORDER BY to_jsonb(r)::text), '[]'::jsonb)::text
            FROM payroll.payrollcalculationsnapshotusedratedefinitions r
            JOIN payroll.payrollcalculationsnapshots s USING (payrollcalculationsnapshotid)
            WHERE s.payrollperiodid = :pid
        """), {"pid": period_id})
        return snapshot


async def test_mapping_read_preview_and_refresh_are_revision_and_fingerprint_bound(
    p3c_client, p3c_engine, tenant,
):
    period = await _period(p3c_client, p3c_engine, tenant)
    source = await _payroll_inputs(p3c_engine, tenant, period)
    before_payroll = await _payroll_snapshot(p3c_engine, period["payroll_period_id"])
    created = await _upload(
        p3c_client, tenant, period,
        f"Driver Number,Work Date,Regular\n{source['drivernumber']},{source['dates'][0].isoformat()},4\n".encode(),
    )
    assert created.status_code == 200, created.text
    import_id = created.json()["import_id"]
    path = f"/payroll/imports/{import_id}"
    target = await _definition_target(p3c_client, tenant, import_id)

    mapped = await _map(p3c_client, tenant, import_id, target)
    assert mapped.status_code == 200, mapped.text
    body = mapped.json()
    assert body["status"] == "NeedsReview"
    assert body["revision"] == 2
    assert body["mapping"]["version"] == 1
    assert body["mapping"]["columns"] == [
        {"source_column": 1, "target": "DriverNumber"},
        {"source_column": 2, "target": "WorkDate"},
        {"source_column": 3, "target": target,
         "payroll_period_definition_id": int(target.split(":")[1]),
         "pay_definition_id": body["mapping"]["columns"][2]["pay_definition_id"],
         "rate_definition_id": body["mapping"]["columns"][2]["rate_definition_id"]},
    ]
    assert body["mapping"]["columns"][2]["pay_definition_id"] > 0
    assert body["mapping"]["columns"][2]["rate_definition_id"] > 0
    assert body["fingerprints"]["mapping"]
    assert body["fingerprints"]["content"]
    plan = body["fingerprints"]["plan"]
    assert len(plan) == 64
    assert body["summary"]["rows"] == 1
    async with p3c_engine.connect() as connection:
        no_decisions = (await connection.execute(text("""
            SELECT
                (SELECT count(*) FROM import.payrollworkrecords r
                 WHERE r.companyid = :cid AND r.importid = :iid AND r.isexcluded) AS excluded_rows,
                (SELECT count(*) FROM import.payrollworkdrivermatches m
                 WHERE m.companyid = :cid AND m.importid = :iid
                   AND (m.confirmedbyuserid IS NOT NULL OR m.confirmedatutc IS NOT NULL)) AS confirmed_matches,
                (SELECT count(*) FROM import.payrollworkcells c
                 WHERE c.companyid = :cid AND c.importid = :iid AND
                   (c.decision IS NOT NULL OR c.decisionbasisquantity IS NOT NULL OR
                    c.statusdecision IS NOT NULL OR c.statusdecisionbasiskeyid IS NOT NULL OR
                    c.dayoffdecision IS NOT NULL OR c.applieddraftlineid IS NOT NULL OR
                    c.appliedaction IS NOT NULL OR c.appliedevidenceeventid IS NOT NULL)) AS decided_or_applied_cells,
                (SELECT count(*) FROM import.payrollworkstatusremovals s
                 WHERE s.companyid = :cid AND s.importid = :iid) AS status_removals
        """), {"cid": tenant.company_id, "iid": import_id})).mappings().one()
    assert dict(no_decisions) == {
        "excluded_rows": 0, "confirmed_matches": 0,
        "decided_or_applied_cells": 0, "status_removals": 0,
    }

    page = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 2, "plan_fingerprint": plan}, headers=tenant.admin,
    )
    assert page.status_code == 200, page.text
    page_body = page.json()
    assert page_body["revision"] == 2
    assert page_body["plan_fingerprint"] == plan
    assert page_body["filtered_count"] == 1
    assert page_body["items"][0]["source_row"] == 2
    assert page_body["items"][0]["cells"][0]["quantity"] == "4"
    assert page_body["items"][0]["cells"][0]["outcome"] == "New"
    assert page_body["items"][0]["cells"][0]["driver_assignment_id"] == tenant.driver_a
    assert page_body["items"][0]["identifiers"][0]["resolution"] == "Exact"
    assert await _payroll_snapshot(p3c_engine, period["payroll_period_id"]) == before_payroll

    stale_mapping = await _map(p3c_client, tenant, import_id, target, revision=1)
    current_metadata = await p3c_client.get(path, headers=tenant.admin)
    assert stale_mapping.status_code == 409
    assert stale_mapping.json()["detail"]["code"] == "IMPORT_REVISION_CONFLICT"
    assert current_metadata.json()["revision"] == 2
    assert current_metadata.json()["fingerprints"]["plan"] == plan

    stale_page = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 1, "plan_fingerprint": plan}, headers=tenant.admin,
    )
    stale_fingerprint = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 2, "plan_fingerprint": "0" * 64}, headers=tenant.admin,
    )
    assert stale_page.status_code == stale_fingerprint.status_code == 409
    assert stale_page.json()["detail"]["code"] == "IMPORT_PREVIEW_STALE"
    assert stale_fingerprint.json()["detail"]["code"] == "IMPORT_PREVIEW_STALE"

    line = await p3c_client.post(
        f"/payroll/periods/{period['payroll_period_id']}/lines",
        json={"driver_assignment_id": tenant.driver_a,
              "payroll_period_definition_id": int(target.split(":")[1]),
              "work_date": source["dates"][0].isoformat(), "quantity": "2"},
        headers=tenant.admin,
    )
    assert line.status_code == 201, line.text

    refresh = await p3c_client.post(
        f"{path}/refresh", json={"expected_revision": 2}, headers=tenant.admin,
    )
    assert refresh.status_code == 200, refresh.text
    assert refresh.json()["revision"] == 3
    assert refresh.json()["status"] == "NeedsReview"
    refreshed_plan = refresh.json()["fingerprints"]["plan"]
    assert refreshed_plan != plan
    mixed_page = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 2, "plan_fingerprint": plan}, headers=tenant.admin,
    )
    refreshed_page = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 3, "plan_fingerprint": refreshed_plan}, headers=tenant.admin,
    )
    assert mixed_page.status_code == 409
    assert refreshed_page.status_code == 200, refreshed_page.text
    refreshed_cell = refreshed_page.json()["items"][0]["cells"][0]
    assert refreshed_cell["outcome"] == "Conflict"
    assert refreshed_cell["current_value"]["quantity"] == "2"
    assert refreshed_cell["quantity"] == "4"


async def test_preview_resolves_each_identity_name_proposals_unknown_values_and_disagreement(
    p3c_client, p3c_engine, tenant,
):
    workforce_manager = make_user(tenant, ["employees.manage"])
    token = workforce_manager["Authorization"].split(" ", 1)[1]
    first_identity = await create_workforce_driver(
        p3c_client, token, tenant.branch_a, "A2 Import Alpha", f"A2A-{uuid4().hex[:8]}"
    )
    second_identity = await create_workforce_driver(
        p3c_client, token, tenant.branch_a, "A2 Import Beta", f"A2B-{uuid4().hex[:8]}"
    )
    await create_workforce_driver(
        p3c_client, token, tenant.branch_a, "A2 Import Alpha", f"A2C-{uuid4().hex[:8]}"
    )
    period = await _period(p3c_client, p3c_engine, tenant)
    first = await _payroll_inputs(p3c_engine, tenant, period, first_identity["driver_assignment_id"])
    second = await _payroll_inputs(p3c_engine, tenant, period, second_identity["driver_assignment_id"])
    async with p3c_engine.connect() as conn:
        foreign_branch_driver = (await conn.execute(text("""
            SELECT d.drivernumber, e.employeekey, e.fullname
            FROM core.drivers d JOIN core.employees e
              ON e.employeeid = d.employeeid AND e.companyid = d.companyid
            WHERE d.companyid = :cid AND d.driverid = :did
        """), {"cid": tenant.company_id, "did": tenant.driver_b_id})).mappings().one()
    data = (
        "Driver Number,Employee Key,Driver Name,Work Date,Regular\n"
        f"{first['drivernumber']},{first['employeekey']},{first['drivernamesnapshot']},{first['dates'][0].isoformat()},4\n"
        f"{second['drivernumber']},{second['employeekey']},{second['drivernamesnapshot']},{second['dates'][1].isoformat()},2\n"
        f"{foreign_branch_driver['drivernumber']},{foreign_branch_driver['employeekey']},"
        f"{foreign_branch_driver['fullname']},{first['dates'][2].isoformat()},1\n"
        f"999999,NO-SUCH-EMPLOYEE,Unknown Driver,{first['dates'][3].isoformat()},2\n"
        f"{first['drivernumber']},{second['employeekey']},{first['drivernamesnapshot']},{first['dates'][4].isoformat()},3\n"
    ).encode()
    created = await _upload(p3c_client, tenant, period, data)
    assert created.status_code == 200, created.text
    import_id = created.json()["import_id"]
    path = f"/payroll/imports/{import_id}"
    target = await _definition_target(p3c_client, tenant, import_id)
    exact_mapping = {
        "1": "DriverNumber", "2": "EmployeeKey", "3": "DriverName",
        "4": "WorkDate", "5": target,
    }
    mapped = await p3c_client.put(
        f"{path}/mapping", json={"expected_revision": 1, "column_mapping": exact_mapping},
        headers=tenant.admin,
    )
    assert mapped.status_code == 200, mapped.text
    plan = mapped.json()["fingerprints"]["plan"]
    page = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 2, "plan_fingerprint": plan}, headers=tenant.admin,
    )
    assert page.status_code == 200, page.text
    exact, unique_name, foreign_branch, unknown, disagreement = page.json()["items"]
    assert {(item["identifier_kind"], item["resolution"]) for item in exact["identifiers"]} == {
        ("DriverNumber", "Exact"), ("EmployeeKey", "Exact"), ("DriverName", "Unresolved")}, (
            first, exact["identifiers"])
    assert exact["cells"][0]["outcome"] == "Invalid"
    assert {(item["identifier_kind"], item["resolution"]) for item in unique_name["identifiers"]} == {
        ("DriverNumber", "Exact"), ("EmployeeKey", "Exact"), ("DriverName", "Proposed")}
    assert unique_name["cells"][0]["outcome"] == "New"
    assert {item["resolution"] for item in foreign_branch["identifiers"]} == {"Unresolved"}
    assert {item["resolution"] for item in unknown["identifiers"]} == {"Unresolved"}
    assert {item["identifier_kind"] for item in unknown["identifiers"]} == {
        "DriverNumber", "EmployeeKey", "DriverName"}

    issue_page = await p3c_client.get(
        f"{path}/issues", params={"expected_revision": 2, "plan_fingerprint": plan}, headers=tenant.admin,
    )
    assert issue_page.status_code == 200, issue_page.text
    issues = issue_page.json()["items"]
    assert any(item["source_row"] == 4 and item["code"] == "IMPORT_DRIVER_UNRESOLVED" for item in issues)
    assert any(item["source_row"] == 6 and item["code"] == "IMPORT_DRIVER_DISAGREEMENT" for item in issues)

    name_only = await p3c_client.put(
        f"{path}/mapping",
        json={"expected_revision": 2, "column_mapping": {
            "1": "Ignore", "2": "Ignore", "3": "DriverName", "4": "WorkDate", "5": target}},
        headers=tenant.admin,
    )
    assert name_only.status_code == 200, name_only.text
    name_plan = name_only.json()["fingerprints"]["plan"]
    proposed_page = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 3, "plan_fingerprint": name_plan}, headers=tenant.admin,
    )
    assert proposed_page.status_code == 200, proposed_page.text
    proposed, unique_proposal, _, _, _ = proposed_page.json()["items"]
    assert proposed["identifiers"][0]["resolution"] == "Unresolved"
    assert proposed["cells"][0]["outcome"] == "Invalid"
    assert unique_proposal["identifiers"][0]["resolution"] == "Proposed"
    assert unique_proposal["cells"][0]["outcome"] == "Unresolved"
    proposed_issues = await p3c_client.get(
        f"{path}/issues", params={"expected_revision": 3, "plan_fingerprint": name_plan}, headers=tenant.admin,
    )
    proposal_items = proposed_issues.json()["items"]
    assert any(item["source_row"] == 3 and item["code"] == "IMPORT_DRIVER_CONFIRMATION_REQUIRED"
               for item in proposal_items)
    assert not any(item["source_row"] == 2 and item["code"] == "IMPORT_DRIVER_CONFIRMATION_REQUIRED"
                   for item in proposal_items)


async def test_mapping_rejects_incomplete_duplicate_and_foreign_targets_without_advancing_revision(
    p3c_client, p3c_engine, tenant,
):
    period = await _period(p3c_client, p3c_engine, tenant)
    source = await _payroll_inputs(p3c_engine, tenant, period)
    created = await _upload(
        p3c_client, tenant, period,
        f"Driver Number,Work Date,Regular,Comment\n{source['drivernumber']},{source['dates'][0].isoformat()},4,manual\n".encode(),
    )
    assert created.status_code == 200, created.text
    import_id = created.json()["import_id"]
    path = f"/payroll/imports/{import_id}"
    definition = await _definition_target(p3c_client, tenant, import_id)
    base = {"expected_revision": 1}
    attempts = [
        ({**base, "column_mapping": {"1": "DriverNumber", "2": "WorkDate", "3": definition}},
         "IMPORT_MAPPING_INCOMPLETE"),
        ({**base, "column_mapping": {"1": "DriverNumber", "2": "WorkDate", "3": "Ignore", "4": "Ignore"}},
         "IMPORT_MAPPING_REQUIRED_TARGET"),
        ({**base, "column_mapping": {"1": "DriverNumber", "2": "DriverNumber", "3": definition, "4": "Ignore"}},
         "IMPORT_MAPPING_DUPLICATE_TARGET"),
        ({**base, "column_mapping": {"1": "DriverNumber", "2": "WorkDate", "3": "Definition:999999", "4": "Ignore"}},
         "IMPORT_MAPPING_INVALID"),
    ]
    for body, code in attempts:
        response = await p3c_client.put(f"{path}/mapping", json=body, headers=tenant.admin)
        assert response.status_code == 422, response.text
        assert response.json()["detail"]["code"] == code
    refresh = await p3c_client.post(f"{path}/refresh", json=base, headers=tenant.admin)
    read_back = await p3c_client.get(path, headers=tenant.admin)
    assert refresh.status_code == 422
    assert refresh.json()["detail"]["code"] == "IMPORT_MAPPING_REQUIRED"
    assert read_back.status_code == 200
    assert read_back.json()["revision"] == 1
    assert read_back.json()["status"] == "NeedsMapping"
    assert read_back.json()["mapping"] is None


async def test_mapping_rejects_source_column_outside_sheet_without_partial_preview_writes(
    p3c_client, p3c_engine, tenant,
):
    period = await _period(p3c_client, p3c_engine, tenant)
    source = await _payroll_inputs(p3c_engine, tenant, period)
    content = (
        f"Driver Number,Work Date,Regular\n"
        f"{source['drivernumber']},{source['dates'][0].isoformat()},4\n"
    ).encode()
    created = await _upload(p3c_client, tenant, period, content)
    assert created.status_code == 200, created.text
    import_id = created.json()["import_id"]
    target = await _definition_target(p3c_client, tenant, import_id)
    before = await _payroll_snapshot(p3c_engine, period["payroll_period_id"])

    refused = await p3c_client.put(
        f"/payroll/imports/{import_id}/mapping",
        json={"expected_revision": 1, "column_mapping": {
            "1": "DriverNumber", "2": "WorkDate", "3": target, "4": "Ignore",
        }},
        headers=tenant.admin,
    )
    assert refused.status_code == 422, refused.text
    assert refused.json()["detail"]["code"] == "IMPORT_MAPPING_INVALID"
    assert await _payroll_snapshot(p3c_engine, period["payroll_period_id"]) == before

    async with p3c_engine.connect() as connection:
        state = (await connection.execute(text("""
            SELECT i.revision, i.status, w.columnmapping, w.mappingfingerprint,
                   w.contentfingerprint, w.planfingerprint,
                   (SELECT count(*) FROM import.importsourcerecords r
                    WHERE r.companyid = i.companyid AND r.importid = i.importid) AS source_rows,
                   (SELECT count(*) FROM import.payrollworkrecords r
                    WHERE r.companyid = i.companyid AND r.importid = i.importid) AS work_rows,
                   (SELECT count(*) FROM import.payrollworkcells c
                    WHERE c.companyid = i.companyid AND c.importid = i.importid) AS work_cells,
                   (SELECT count(*) FROM import.importissues e
                    WHERE e.companyid = i.companyid AND e.importid = i.importid) AS issues
            FROM import.imports i
            JOIN import.payrollworkimports w USING (companyid, importid)
            WHERE i.companyid = :cid AND i.importid = :iid
        """), {"cid": tenant.company_id, "iid": import_id})).mappings().one()
    assert state["revision"] == 1
    assert state["status"] == "NeedsMapping"
    assert state["columnmapping"] is None
    assert state["mappingfingerprint"] is state["contentfingerprint"] is state["planfingerprint"] is None
    assert (state["source_rows"], state["work_rows"], state["work_cells"], state["issues"]) == (0, 0, 0, 0)


async def test_preview_rows_issues_filters_and_pages_cover_the_full_import_deterministically(
    p3c_client, p3c_engine, tenant,
):
    period = await _period(p3c_client, p3c_engine, tenant)
    source = await _payroll_inputs(p3c_engine, tenant, period)
    rows = [
        (source["drivernumber"], source["dates"][0], "4"),
        (source["drivernumber"], source["dates"][1], "5"),
        ("999999", source["dates"][2], "1"),
        (source["drivernumber"], source["dates"][0], "4"),
    ]
    contents = "Driver Number,Work Date,Regular\n" + "".join(
        f"{driver},{work_date.isoformat()},{quantity}\n" for driver, work_date, quantity in rows)
    created = await _upload(p3c_client, tenant, period, contents.encode())
    assert created.status_code == 200, created.text
    import_id = created.json()["import_id"]
    definition = await _definition_target(p3c_client, tenant, import_id)
    mapped = await _map(p3c_client, tenant, import_id, definition)
    assert mapped.status_code == 200, mapped.text
    plan = mapped.json()["fingerprints"]["plan"]
    path = f"/payroll/imports/{import_id}"

    row_pages = []
    offset = 0
    while True:
        response = await p3c_client.get(
            f"{path}/rows", params={"expected_revision": 2, "plan_fingerprint": plan,
                                   "offset": offset, "limit": 2}, headers=tenant.admin,
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["revision"] == 2 and body["plan_fingerprint"] == plan
        assert body["summary"] == mapped.json()["summary"]
        row_pages.extend(body["items"])
        if body["next_offset"] is None:
            break
        offset = body["next_offset"]
    assert [item["source_row"] for item in row_pages] == [2, 3, 4, 5]
    assert row_pages[0]["cells"][0]["outcome"] == "Invalid"
    assert row_pages[-1]["cells"][0]["outcome"] == "Invalid"

    issue_pages = []
    offset = 0
    while True:
        response = await p3c_client.get(
            f"{path}/issues", params={"expected_revision": 2, "plan_fingerprint": plan,
                                     "offset": offset, "limit": 2}, headers=tenant.admin,
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["summary"] == mapped.json()["summary"]
        issue_pages.extend(body["items"])
        if body["next_offset"] is None:
            break
        offset = body["next_offset"]
    assert {item["code"] for item in issue_pages} >= {
        "IMPORT_DRIVER_UNRESOLVED", "IMPORT_DUPLICATE_FACT"}
    assert sorted(item["source_row"] for item in issue_pages if item["code"] == "IMPORT_DUPLICATE_FACT") == [2, 5]
    assert all(item["severity"] == "Error" for item in issue_pages)

    error_records = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 2, "plan_fingerprint": plan,
                               "severity": "Error"}, headers=tenant.admin,
    )
    assert error_records.status_code == 200, error_records.text
    assert error_records.json()["filtered_count"] == 3
    assert error_records.json()["summary"] == mapped.json()["summary"]
    default_rows = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 2, "plan_fingerprint": plan}, headers=tenant.admin,
    )
    default_issues = await p3c_client.get(
        f"{path}/issues", params={"expected_revision": 2, "plan_fingerprint": plan}, headers=tenant.admin,
    )
    assert default_rows.status_code == default_issues.status_code == 200
    assert default_rows.json()["limit"] == 25 and len(default_rows.json()["items"]) == 4
    assert default_issues.json()["limit"] == 50
    assert len(default_issues.json()["items"]) == len(issue_pages) <= 50
    invalid_severity = await p3c_client.get(
        f"{path}/issues", params={"expected_revision": 2, "plan_fingerprint": plan,
                                 "severity": "Invalid"}, headers=tenant.admin,
    )
    invalid_limits = []
    for route, limit in (("rows", 0), ("rows", -1), ("rows", 101),
                         ("issues", 0), ("issues", -1), ("issues", 201)):
        invalid_limits.append(await p3c_client.get(
            f"{path}/{route}", params={"expected_revision": 2, "plan_fingerprint": plan, "limit": limit},
            headers=tenant.admin,
        ))
    missing_revision = await p3c_client.get(f"{path}/rows", params={"plan_fingerprint": plan}, headers=tenant.admin)
    out_of_range_offset = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 2, "plan_fingerprint": plan,
                               "offset": 2_147_483_648}, headers=tenant.admin,
    )
    malformed_fingerprint = await p3c_client.get(
        f"{path}/rows", params={"expected_revision": 2, "plan_fingerprint": "a" * 63},
        headers=tenant.admin,
    )
    assert invalid_severity.status_code == 422
    assert all(response.status_code == 422 for response in invalid_limits)
    assert (missing_revision.status_code == out_of_range_offset.status_code
            == malformed_fingerprint.status_code == 422)


async def test_preview_reports_configured_day_off_and_real_held_full_day_off_status_without_decisions(
    p3c_client, p3c_engine, tenant,
):
    period = await _period_with_mask(p3c_client, p3c_engine, tenant, mask=1)
    source = await _payroll_inputs(p3c_engine, tenant, period)
    async with p3c_engine.connect() as conn:
        off_day = await conn.scalar(text("""
            SELECT workdate FROM payroll.payrollperioddays
            WHERE payrollperiodid = :pid AND isconfiguredoffday
            ORDER BY workdate LIMIT 1
        """), {"pid": period["payroll_period_id"]})
        regular_day = await conn.scalar(text("""
            SELECT workdate FROM payroll.payrollperioddays
            WHERE payrollperiodid = :pid AND NOT isconfiguredoffday
              AND workdate = ANY(CAST(:eligible AS date[]))
            ORDER BY workdate LIMIT 1
        """), {"pid": period["payroll_period_id"], "eligible": source["dates"]})
    assert off_day in source["dates"]
    assert regular_day in source["dates"]
    grid = await get_grid(p3c_client, tenant, period)
    target = f"Definition:{grid['columns'][0]['payroll_period_definition_id']}"
    existing_line = await p3c_client.post(
        f"/payroll/periods/{period['payroll_period_id']}/lines",
        json={"driver_assignment_id": tenant.driver_a,
              "payroll_period_definition_id": int(target.split(":")[1]),
              "work_date": off_day.isoformat(), "quantity": "2"},
        headers=tenant.admin,
    )
    assert existing_line.status_code == 201, existing_line.text

    off_key = seed_status_key(tenant, "IMPORT_REST_TEST", off=True)
    held_status = await p3c_client.post(
        f"/payroll/periods/{period['payroll_period_id']}/day-grid",
        json={"work_date": regular_day.isoformat(), "rows": [{
            "driver_assignment_id": tenant.driver_a, "values": {}, "status_key": "IMPORT_REST_TEST",
        }]},
        headers=tenant.admin,
    )
    assert held_status.status_code == 200, held_status.text
    before = await _payroll_snapshot(p3c_engine, period["payroll_period_id"])

    imported = await _upload(
        p3c_client, tenant, period,
        (f"Driver Number,Work Date,Regular\n{source['drivernumber']},{off_day.isoformat()},4\n"
         f"{source['drivernumber']},{regular_day.isoformat()},4\n").encode(),
    )
    assert imported.status_code == 200, imported.text
    import_id = imported.json()["import_id"]
    mapped = await _map(p3c_client, tenant, import_id, target)
    assert mapped.status_code == 200, mapped.text
    plan = mapped.json()["fingerprints"]["plan"]
    rows = await p3c_client.get(
        f"/payroll/imports/{import_id}/rows",
        params={"expected_revision": 2, "plan_fingerprint": plan}, headers=tenant.admin,
    )
    issues = await p3c_client.get(
        f"/payroll/imports/{import_id}/issues",
        params={"expected_revision": 2, "plan_fingerprint": plan}, headers=tenant.admin,
    )
    assert rows.status_code == issues.status_code == 200
    assert {cell["work_date"] for row in rows.json()["items"] for cell in row["cells"]} == {
        off_day.isoformat(), regular_day.isoformat()}
    assert any(cell["is_day_off"] is True for row in rows.json()["items"] for cell in row["cells"])
    assert any(cell["has_status"] is True and cell["held_status_key_id"] == off_key
               for row in rows.json()["items"] for cell in row["cells"])
    codes = {item["code"] for item in issues.json()["items"]}
    assert "IMPORT_DAY_OFF_DECISION_REQUIRED" in codes
    assert "IMPORT_STATUS_DECISION_REQUIRED" in codes
    assert "IMPORT_VALUE_CONFLICT" in codes
    assert await _payroll_snapshot(p3c_engine, period["payroll_period_id"]) == before


async def test_preview_current_comparison_uses_existing_zero_as_none_and_exact_quantity_matrix(
    p3c_client, p3c_engine, tenant,
):
    period = await _period(p3c_client, p3c_engine, tenant)
    source = await _payroll_inputs(p3c_engine, tenant, period)
    assert len(source["dates"]) >= 6
    grid = await get_grid(p3c_client, tenant, period)
    target = f"Definition:{grid['columns'][0]['payroll_period_definition_id']}"
    existing = {0: "5", 1: "0", 3: "6", 4: "5"}
    for offset, quantity_value in existing.items():
        response = await p3c_client.post(
            f"/payroll/periods/{period['payroll_period_id']}/lines",
            json={"driver_assignment_id": tenant.driver_a,
                  "payroll_period_definition_id": int(target.split(":")[1]),
                  "work_date": source["dates"][offset].isoformat(), "quantity": quantity_value},
            headers=tenant.admin,
        )
        assert response.status_code == 201, response.text
    incoming = ("5", "0", "0", "0", "4", "4")
    contents = "Driver Number,Work Date,Regular\n" + "".join(
        f"{source['drivernumber']},{source['dates'][offset].isoformat()},{quantity_value}\n"
        for offset, quantity_value in enumerate(incoming))
    created = await _upload(p3c_client, tenant, period, contents.encode())
    assert created.status_code == 200, created.text
    mapped = await _map(p3c_client, tenant, created.json()["import_id"], target)
    assert mapped.status_code == 200, mapped.text
    page = await p3c_client.get(
        f"/payroll/imports/{created.json()['import_id']}/rows",
        params={"expected_revision": 2, "plan_fingerprint": mapped.json()["fingerprints"]["plan"]},
        headers=tenant.admin,
    )
    assert page.status_code == 200, page.text
    cells = [row["cells"][0] for row in page.json()["items"]]
    assert [cell["outcome"] for cell in cells] == [
        "Unchanged", "Unchanged", "Unchanged", "Conflict", "Conflict", "New"]
    assert cells[0]["current_value"]["quantity"] == "5"
    assert cells[1]["current_value"]["quantity"] is None
    assert cells[2]["current_value"]["quantity"] is None
    assert cells[3]["current_value"]["quantity"] == "6"
    assert page.json()["summary"]["conflicts"] >= 2


async def test_numeric_driver_identifier_from_xlsx_warns_and_never_infers_leading_zeroes(
    p3c_client, p3c_engine, tenant,
):
    workforce_manager = make_user(tenant, ["employees.manage"])
    token = workforce_manager["Authorization"].split(" ", 1)[1]
    numeric_driver = await create_workforce_driver(
        p3c_client, token, tenant.branch_a, "A2 Leading Zero", "001234"
    )
    period = await _period(p3c_client, p3c_engine, tenant)
    source = await _payroll_inputs(p3c_engine, tenant, period, numeric_driver["driver_assignment_id"])
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Driver Number", "Work Date", "Regular"])
    worksheet.append([1234, source["dates"][0].isoformat(), 4])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    created = await _upload(p3c_client, tenant, period, output.getvalue(), filename="numeric-driver.xlsx")
    assert created.status_code == 200, created.text
    import_id = created.json()["import_id"]
    target = await _definition_target(p3c_client, tenant, import_id)
    mapped = await _map(p3c_client, tenant, import_id, target)
    assert mapped.status_code == 200, mapped.text
    page = await p3c_client.get(
        f"/payroll/imports/{import_id}/rows",
        params={"expected_revision": 2, "plan_fingerprint": mapped.json()["fingerprints"]["plan"]},
        headers=tenant.admin,
    )
    issues = await p3c_client.get(
        f"/payroll/imports/{import_id}/issues",
        params={"expected_revision": 2, "plan_fingerprint": mapped.json()["fingerprints"]["plan"]},
        headers=tenant.admin,
    )
    assert page.status_code == issues.status_code == 200
    identifier = page.json()["items"][0]["identifiers"][0]
    assert identifier["normalized_value"] == "1234"
    assert identifier["resolution"] == "Unresolved"
    assert any(item["code"] == "IMPORT_NUMERIC_IDENTIFIER" for item in issues.json()["items"])


async def test_preview_read_and_mutation_routes_enforce_company_branch_and_driver_boundaries(
    p3c_client, p3c_engine, tenant, p3b_dsn,
):
    period = await _period(p3c_client, p3c_engine, tenant)
    source = await _payroll_inputs(p3c_engine, tenant, period)
    created = await _upload(
        p3c_client, tenant, period,
        f"Driver Number,Work Date,Regular\n{source['drivernumber']},{source['dates'][0].isoformat()},4\n".encode(),
    )
    assert created.status_code == 200, created.text
    import_id = created.json()["import_id"]
    path = f"/payroll/imports/{import_id}"
    target = await _definition_target(p3c_client, tenant, import_id)
    mapped = await _map(p3c_client, tenant, import_id, target)
    assert mapped.status_code == 200, mapped.text
    plan = mapped.json()["fingerprints"]["plan"]

    viewer = make_user(tenant, ["payroll.view"])
    other_branch = make_user(tenant, ["payroll.view"], scope="SpecificBranch", branch_id=tenant.branch_b)
    driver_self = make_driver_self_user(tenant, ["payroll.view", "payroll.entry"])
    foreign = build_payroll_tenant(p3b_dsn)
    page_params = {"expected_revision": 2, "plan_fingerprint": plan}

    viewer_read = await p3c_client.get(path, headers=viewer)
    viewer_page = await p3c_client.get(f"{path}/rows", params=page_params, headers=viewer)
    viewer_template = await p3c_client.get(
        f"/payroll/periods/{period['payroll_period_id']}/imports/template.xlsx", headers=viewer,
    )
    viewer_map = await p3c_client.put(
        f"{path}/mapping",
        json={"expected_revision": 2,
              "column_mapping": {"1": "DriverNumber", "2": "WorkDate", "3": target}},
        headers=viewer,
    )
    viewer_refresh = await p3c_client.post(
        f"{path}/refresh", json={"expected_revision": 2}, headers=viewer,
    )
    branch_read = await p3c_client.get(path, headers=other_branch)
    branch_template = await p3c_client.get(
        f"/payroll/periods/{period['payroll_period_id']}/imports/template.xlsx", headers=other_branch,
    )
    self_read = await p3c_client.get(f"{path}/issues", params=page_params, headers=driver_self)
    self_refresh = await p3c_client.post(
        f"{path}/refresh", json={"expected_revision": 2}, headers=driver_self,
    )
    foreign_read = await p3c_client.get(path, headers=foreign.admin)

    assert viewer_read.status_code == viewer_page.status_code == viewer_template.status_code == 200
    assert viewer_map.status_code == viewer_refresh.status_code == 403
    assert branch_read.status_code == branch_template.status_code == 403
    assert self_read.status_code == self_refresh.status_code == 403
    assert foreign_read.status_code == 404


async def test_period_template_is_text_safe_auto_maps_and_roundtrips_through_upload_preview(
    p3c_client, p3c_engine, tenant,
):
    period = await _period(p3c_client, p3c_engine, tenant)
    template = await p3c_client.get(
        f"/payroll/periods/{period['payroll_period_id']}/imports/template.xlsx", headers=tenant.admin,
    )
    assert template.status_code == 200, template.text
    book = load_workbook(io.BytesIO(template.content), read_only=True, data_only=False)
    try:
        sheet = book["Payroll work"]
        values = list(sheet.iter_rows(min_row=1, max_row=2, values_only=True))
        cells = list(sheet.iter_rows(min_row=1, max_row=2))
        assert values[0][:3] == ("Driver number", "Driver name", "Date")
        assert values[1][0] is not None and values[1][2] is not None
        assert all(cell.data_type == "s" and cell.number_format == "@" for cell in cells[0])
        assert all(cell.data_type == "s" and cell.number_format == "@" for cell in cells[1][:3])
        assert all(cell.value in (None, "") for cell in cells[1][3:])
    finally:
        book.close()

    frozen_headers = values[0]
    late_definition = await rated_definition(
        p3c_client, tenant, definition_code="AFTER", definition_name="Created After Period",
    )
    assert late_definition["definition_code"] == "AFTER"
    after_late_definition = await p3c_client.get(
        f"/payroll/periods/{period['payroll_period_id']}/imports/template.xlsx", headers=tenant.admin,
    )
    assert after_late_definition.status_code == 200, after_late_definition.text
    after_book = load_workbook(io.BytesIO(after_late_definition.content), read_only=True, data_only=True)
    try:
        after_headers = next(after_book["Payroll work"].iter_rows(min_row=1, max_row=1, values_only=True))
    finally:
        after_book.close()
    assert after_headers == frozen_headers
    assert "Pay Item: AFTER" not in after_headers

    uploaded = await _upload(p3c_client, tenant, period, template.content, filename="period-template.xlsx")
    assert uploaded.status_code == 200, uploaded.text
    import_id = uploaded.json()["import_id"]
    metadata = await p3c_client.get(f"/payroll/imports/{import_id}", headers=tenant.admin)
    assert metadata.status_code == 200, metadata.text
    target = next(item["target"] for item in metadata.json()["mapping_targets"]
                  if item["target"].startswith("Definition:"))
    assert metadata.json()["mapping_suggestions"] == {
        "1": "DriverNumber", "2": "DriverName", "3": "WorkDate", "4": target}
    mapped = await p3c_client.put(
        f"/payroll/imports/{import_id}/mapping",
        json={"expected_revision": 1,
              "column_mapping": {"1": "DriverNumber", "2": "DriverName",
                                 "3": "WorkDate", "4": target}},
        headers=tenant.admin,
    )
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()["status"] == "NeedsReview"
    assert mapped.json()["summary"]["rows"] > 0
    assert mapped.json()["summary"]["skipped"] == mapped.json()["summary"]["rows"]


async def test_completed_import_notice_uses_fixture_state_and_new_import_still_previews_current_facts(
    p3c_client, p3c_engine, tenant,
):
    period = await _period(p3c_client, p3c_engine, tenant)
    source = await _payroll_inputs(p3c_engine, tenant, period)
    content = f"Driver Number,Work Date,Regular\n{source['drivernumber']},{source['dates'][0].isoformat()},4\n".encode()
    completed_candidate = await _upload(p3c_client, tenant, period, content)
    assert completed_candidate.status_code == 200, completed_candidate.text
    completed_id = completed_candidate.json()["import_id"]
    target = await _definition_target(p3c_client, tenant, completed_id)
    mapped = await _map(p3c_client, tenant, completed_id, target)
    assert mapped.status_code == 200, mapped.text

    async with p3c_engine.begin() as conn:
        await conn.execute(text("""
            UPDATE import.imports SET status = 'Completed', revision = revision + 1,
                completedbyuserid = :uid, completedatutc = NOW()
            WHERE importid = :iid AND companyid = :cid
        """), {"uid": tenant.owner, "iid": completed_id, "cid": tenant.company_id})

    current_line = await p3c_client.post(
        f"/payroll/periods/{period['payroll_period_id']}/lines",
        json={"driver_assignment_id": tenant.driver_a,
              "payroll_period_definition_id": int(target.split(":")[1]),
              "work_date": source["dates"][0].isoformat(), "quantity": "2"},
        headers=tenant.admin,
    )
    assert current_line.status_code == 201, current_line.text

    second = await _upload(p3c_client, tenant, period, content)
    assert second.status_code == 200, second.text
    notice = second.json()["completed_import_notice"]
    assert notice["import_id"] == completed_id
    assert notice["byte_match"] is True
    second_id = second.json()["import_id"]
    second_target = await _definition_target(p3c_client, tenant, second_id)
    preview = await _map(p3c_client, tenant, second_id, second_target)
    assert preview.status_code == 200, preview.text
    assert preview.json()["summary"]["rows"] == 1
    assert preview.json()["fingerprints"]["plan"]
    assert preview.json()["completed_import_notice"]["semantic_match"] is True
    second_rows = await p3c_client.get(
        f"/payroll/imports/{second_id}/rows",
        params={"expected_revision": 2, "plan_fingerprint": preview.json()["fingerprints"]["plan"]},
        headers=tenant.admin,
    )
    assert second_rows.status_code == 200, second_rows.text
    assert second_rows.json()["items"][0]["cells"][0]["outcome"] == "Conflict"
    assert second_rows.json()["items"][0]["cells"][0]["current_value"]["quantity"] == "2"

    crlf_content = content.replace(b"\n", b"\r\n")
    semantic_only = await _upload(p3c_client, tenant, period, crlf_content)
    assert semantic_only.status_code == 200, semantic_only.text
    assert semantic_only.json()["completed_import_notice"] is None
    semantic_id = semantic_only.json()["import_id"]
    semantic_target = await _definition_target(p3c_client, tenant, semantic_id)
    semantic_preview = await _map(p3c_client, tenant, semantic_id, semantic_target)
    assert semantic_preview.status_code == 200, semantic_preview.text
    semantic_notice = semantic_preview.json()["completed_import_notice"]
    assert semantic_notice["import_id"] == completed_id
    assert semantic_notice["byte_match"] is False
    assert semantic_notice["semantic_match"] is True
    semantic_rows = await p3c_client.get(
        f"/payroll/imports/{semantic_id}/rows",
        params={"expected_revision": 2, "plan_fingerprint": semantic_preview.json()["fingerprints"]["plan"]},
        headers=tenant.admin,
    )
    assert semantic_rows.status_code == 200, semantic_rows.text
    assert semantic_rows.json()["items"][0]["cells"][0]["outcome"] == "Conflict"
    assert semantic_rows.json()["items"][0]["cells"][0]["current_value"]["quantity"] == "2"


async def test_unsupported_import_contract_fails_closed_for_mapping(
    p3c_client, p3c_engine, tenant,
):
    period = await _period(p3c_client, p3c_engine, tenant)
    created = await _upload(
        p3c_client, tenant, period, b"Driver Number,Work Date,Regular\n0012,2026-10-20,4\n",
    )
    assert created.status_code == 200, created.text
    import_id = created.json()["import_id"]
    target = await _definition_target(p3c_client, tenant, import_id)
    async with p3c_engine.begin() as conn:
        import_id = (await conn.execute(text("""
            INSERT INTO import.imports
                (companyid, branchid, importkind, sourcekind, correlationid, contractversion,
                 idempotencykey, creationfingerprint, createdbyuserid, updatedbyuserid)
            VALUES (:cid, :bid, 'PayrollWork', 'File', :uuid, 2, :key, :hash, :uid, :uid)
            RETURNING importid
        """), {"cid": tenant.company_id, "bid": tenant.branch_a, "uuid": uuid4(), "key": uuid4(),
              "hash": "0" * 64, "uid": tenant.owner})).scalar_one()
        await conn.execute(text("""
            INSERT INTO import.payrollworkimports (importid, companyid, branchid, payrollperiodid)
            VALUES (:iid, :cid, :bid, :pid)
        """), {"iid": import_id, "cid": tenant.company_id, "bid": tenant.branch_a,
              "pid": period["payroll_period_id"]})
        await conn.execute(text("""
            INSERT INTO import.filesources
                (importid, companyid, originalfilename, filetype, filesizebytes, filehashsha256, filecontent)
            VALUES (:iid, :cid, 'future.csv', 'csv', 1, :hash, :content)
        """), {"iid": import_id, "cid": tenant.company_id, "hash": "0" * 64, "content": b"x"})
    unsupported = await p3c_client.put(
        f"/payroll/imports/{import_id}/mapping",
        json={"expected_revision": 1,
              "column_mapping": {"1": "DriverNumber", "2": "WorkDate", "3": target}},
        headers=tenant.admin,
    )
    assert unsupported.status_code == 409, unsupported.text
    assert unsupported.json()["detail"]["code"] == "IMPORT_CONTRACT_UNSUPPORTED"


async def test_finalization_purges_api_created_preview_detail_and_blocks_recreation(
    p3c_client, p3c_engine, tenant,
):
    from tests.calculation_snapshot_helpers import finalize, standard_scenario, submit_and_approve

    scenario = await standard_scenario(p3c_client, p3c_engine, tenant)
    period = scenario.period
    source = await _payroll_inputs(p3c_engine, tenant, period)
    date_value = f"{source['dates'][0].day}/{source['dates'][0].month}/{source['dates'][0].year}"
    data = (f"Driver Number,Work Date,Regular\n{source['drivernumber']},{date_value},4\n"
            f"{source['drivernumber']},{date_value},4\n").encode()
    created = await _upload(p3c_client, tenant, period, data)
    assert created.status_code == 200, created.text
    import_id = created.json()["import_id"]
    target = await _definition_target(p3c_client, tenant, import_id)
    mapped = await _map(p3c_client, tenant, import_id, target)
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()["preview"]["interpretation"]["date_order_state"] == "detected"
    assert mapped.json()["interpretation"]["date_order"] == "DMY"

    detail_tables = (
        "importsourcerecords", "filesourcerecords", "payrollworkrecords",
        "payrollworkdrivermatches", "payrollworkrecordidentifiers", "payrollworkcells", "importissues",
    )
    async with p3c_engine.connect() as conn:
        before = {table: await conn.scalar(text(f"SELECT count(*) FROM import.{table} WHERE importid = :iid"),
                                           {"iid": import_id}) for table in detail_tables}
    assert all(count > 0 for count in before.values()), before
    retained = mapped.json()["fingerprints"]

    await submit_and_approve(p3c_client, tenant, period)
    finalized = await finalize(p3c_client, tenant, period)
    assert finalized.status_code in (200, 204), finalized.text
    async with p3c_engine.connect() as conn:
        after = {table: await conn.scalar(text(f"SELECT count(*) FROM import.{table} WHERE importid = :iid"),
                                          {"iid": import_id}) for table in detail_tables}
    assert all(count == 0 for count in after.values()), after
    read_back = await p3c_client.get(f"/payroll/imports/{import_id}", headers=tenant.admin)
    assert read_back.status_code == 200, read_back.text
    assert read_back.json()["mapping"] is not None
    assert read_back.json()["fingerprints"]["mapping"] == retained["mapping"]
    assert read_back.json()["fingerprints"]["content"] == retained["content"]
    assert read_back.json()["fingerprints"]["plan"] is None
    assert read_back.json()["interpretation"]["date_order"] == mapped.json()["interpretation"]["date_order"]
    revision = read_back.json()["revision"]
    remap = await _map(p3c_client, tenant, import_id, target, revision=revision)
    refresh = await p3c_client.post(
        f"/payroll/imports/{import_id}/refresh", json={"expected_revision": revision}, headers=tenant.admin,
    )
    assert remap.status_code == refresh.status_code == 409
    assert remap.json()["detail"]["code"] == refresh.json()["detail"]["code"] == "IMPORT_NOT_EDITABLE"

