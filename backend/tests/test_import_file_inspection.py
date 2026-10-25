from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook

from app.imports import ImportFailure, inspect_file
from app.imports.files import MAX_XLSX_MEMBERS


def _workbook_bytes(rows: list[list[object]], *, sheet_name: str = "Data") -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _rewrite_xlsx(data: bytes, replacements: dict[str, bytes], additions: dict[str, bytes] | None = None) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info.filename, replacements.get(info.filename, source.read(info.filename)))
        for name, content in (additions or {}).items():
            target.writestr(name, content)
    return output.getvalue()


def test_inspect_csv_reports_bounded_metadata_and_preserves_identifier_text() -> None:
    result = inspect_file(
        b"EmployeeKey;Work Date;Regular\r\n00123;2026-10-20;4.5\r\n",
        "payroll.csv",
        header_labels=["Regular"],
        date_order="dmy",
    )

    assert result["file_type"] == "csv"
    assert result["selected_sheet"] == "CSV"
    assert result["header_row_number"] == 1
    assert result["header_state"] == "detected"
    assert result["header_cells"] == ["EmployeeKey", "Work Date", "Regular"]
    assert (result["row_count"], result["column_count"], result["cell_count"]) == (2, 3, 6)
    assert result["encoding"] == "utf-8"
    assert result["delimiter"] == ";"
    assert result["date_order"] == "DMY"
    assert result["date_order_state"] == "selected"


def test_inspect_csv_detects_cp1252_and_single_column_fallback() -> None:
    result = inspect_file("Name\r\nJos\xe9\r\n".encode("cp1252"), "drivers.csv")

    assert result["encoding"] == "cp1252"
    assert result["delimiter"] == ","
    assert result["header_cells"] == ["Name"]


@pytest.mark.parametrize(
    ("data", "filename", "code"),
    [
        (b"\xff\xfeA\x00", "data.csv", "IMPORT_CSV_ENCODING_UNSUPPORTED"),
        (b"a\x00b", "data.csv", "IMPORT_CSV_BINARY_CONTENT"),
        (b"", "data.csv", "IMPORT_FILE_EMPTY"),
        (b"hello", "data.xlsm", "IMPORT_FILE_TYPE_UNSUPPORTED"),
        (b"hello", "data.xlsx", "IMPORT_XLSX_INVALID"),
        (b"abc", "data.txt", "IMPORT_FILE_TYPE_UNSUPPORTED"),
    ],
)
def test_inspect_file_rejects_unsupported_or_malformed_inputs(
    data: bytes, filename: str, code: str
) -> None:
    with pytest.raises(ImportFailure) as exc_info:
        inspect_file(data, filename)

    assert exc_info.value.code == code
    assert exc_info.value.status_code == 422


def test_inspect_csv_rejects_row_column_cell_and_text_limits() -> None:
    with pytest.raises(ImportFailure, match="rows") as rows_error:
        inspect_file(("x\n" * 20_001).encode(), "rows.csv")
    assert rows_error.value.code == "IMPORT_ROW_LIMIT_EXCEEDED"

    with pytest.raises(ImportFailure) as columns_error:
        inspect_file((",".join(["x"] * 51) + "\n").encode(), "columns.csv")
    assert columns_error.value.code == "IMPORT_COLUMN_LIMIT_EXCEEDED"

    with pytest.raises(ImportFailure) as text_error:
        inspect_file(("x" * 8_001 + "\n").encode(), "text.csv")
    assert text_error.value.code == "IMPORT_CELL_TEXT_LIMIT_EXCEEDED"


def test_inspect_xlsx_lists_hidden_sheets_and_uses_actual_extents() -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Rows"
    first.append(["Driver Number", "Date", "Units"])
    first.append(["0012", "2026-10-20", 4])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    output = io.BytesIO()
    workbook.save(output)
    contents = output.getvalue()
    with zipfile.ZipFile(io.BytesIO(contents)) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").replace(
            b'<dimension ref="A1:C2" />', b'<dimension ref="A1:C1048576" />'
        )
    contents = _rewrite_xlsx(contents, {"xl/worksheets/sheet1.xml": sheet_xml})

    result = inspect_file(contents, "payroll.xlsx", sheet_name="Rows")

    assert result["sheets"] == [
        {"name": "Rows", "index": 1, "visibility": "visible"},
        {"name": "Hidden", "index": 2, "visibility": "hidden"},
    ]
    assert result["selected_sheet"] == "Rows"
    assert result["header_cells"] == ["Driver Number", "Date", "Units"]
    assert (result["row_count"], result["column_count"], result["cell_count"]) == (2, 3, 6)


def test_inspect_xlsx_preserves_formula_cache_distinctions_without_evaluation() -> None:
    data = _workbook_bytes([["Work Date", "Formula"], ["2026-10-20", "=1+1"]])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml")
    # openpyxl writes an empty <v/> for an uncached formula. Add a second
    # formula with no value node to preserve the separate missing-cache case.
    worksheet = worksheet.replace(
        b'<c r="B2"><f>1+1</f><v /></c>',
        b'<c r="B2"><f>1+1</f><v /></c><c r="C2"><f>2+2</f></c>',
    )
    data = _rewrite_xlsx(data, {"xl/worksheets/sheet1.xml": worksheet})

    result = inspect_file(data, "formulas.xlsx")

    assert result["formula_count"] == 2
    assert result["formula_empty_cache_count"] == 1
    assert result["formula_missing_cache_count"] == 1
    assert result["formula_cached_value_count"] == 0
    assert result["header_cells"] == ["Work Date", "Formula", ""]


def test_inspect_xlsx_accepts_selected_header_and_optional_definition_labels() -> None:
    data = _workbook_bytes([["Notes"], ["Actual Hours", "Rate code"]])

    result = inspect_file(
        data,
        "mapped.xlsx",
        header_row_number=2,
        header_labels=["Actual Hours", "R-14"],
    )

    assert result["header_row_number"] == 2
    assert result["header_state"] == "selected"
    assert result["header_cells"] == ["Actual Hours", "Rate code"]


def test_inspect_xlsx_rejects_xml_entities_macro_and_decompression_overflow() -> None:
    data = _workbook_bytes([["Work Date", "Value"]])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml")
    malicious = worksheet.replace(
        b"<worksheet",
        b'<!DOCTYPE worksheet [<!ENTITY x "expanded">]><worksheet',
        1,
    )
    with pytest.raises(ImportFailure) as xml_error:
        inspect_file(_rewrite_xlsx(data, {"xl/worksheets/sheet1.xml": malicious}), "bad.xlsx")
    assert xml_error.value.code == "IMPORT_XLSX_INVALID"

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        content_types = archive.read("[Content_Types].xml") + b"<!-- macroEnabled -->"
    with pytest.raises(ImportFailure) as macro_error:
        inspect_file(_rewrite_xlsx(data, {"[Content_Types].xml": content_types}), "renamed.xlsx")
    assert macro_error.value.code == "IMPORT_FILE_TYPE_MISMATCH"

    padded = _rewrite_xlsx(data, {}, {"custom/large.bin": b"0" * 50_000_001})
    with pytest.raises(ImportFailure) as size_error:
        inspect_file(padded, "large.xlsx")
    assert size_error.value.code == "IMPORT_XLSX_DECOMPRESSED_LIMIT_EXCEEDED"


def test_inspect_xlsx_member_ceiling_and_one_over() -> None:
    data = _workbook_bytes([["Driver Number", "Work Date"]])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        base_members = len(archive.infolist())
    additions = {
        f"padding/member-{index}.bin": b""
        for index in range(MAX_XLSX_MEMBERS - base_members)
    }
    at_limit = _rewrite_xlsx(data, {}, additions)
    assert inspect_file(at_limit, "at-limit.xlsx")["row_count"] == 1

    above_limit = _rewrite_xlsx(
        data,
        {},
        {**additions, f"padding/member-{MAX_XLSX_MEMBERS}.bin": b""},
    )
    with pytest.raises(ImportFailure) as exc_info:
        inspect_file(above_limit, "above-limit.xlsx")
    assert exc_info.value.code == "IMPORT_XLSX_MEMBER_LIMIT_EXCEEDED"
