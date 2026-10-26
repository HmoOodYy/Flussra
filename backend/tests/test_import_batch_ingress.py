from __future__ import annotations

import asyncio
import hashlib
import io
import re
import zipfile
from dataclasses import replace
from datetime import date, time, timedelta
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from starlette.requests import Request

from app.imports.errors import ImportFailure
from app.imports.files import FileCell, inspect_source_file, read_source_units
from app.imports.mapping import quantity
from app.imports.policy import DEFAULT_IMPORT_RESOURCE_POLICY
from app.imports.transport import receive_files


def _xlsx_sheets(count: int = 2, *, first_rows: int = 2) -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "Day 1"
    first.append(["Driver Number", "Work Date", "Regular"])
    first.append(["0012", "2026-10-01", 4])
    for row in range(3, first_rows + 1):
        first.append([f"0012-{row}", f"2026-10-{row:02d}", row])
    for index in range(2, count + 1):
        sheet = workbook.create_sheet(f"Day {index}")
        sheet.append(["Driver Number", "Work Date", "Regular"])
        sheet.append([f"00{index:02d}", f"2026-10-{index:02d}", index])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _rewrite_sheet(data: bytes, transform) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content = transform(content)
            target.writestr(info.filename, content)
    return output.getvalue()


def _multipart(files: list[tuple[str, bytes]], boundary: str = "batch") -> bytes:
    pieces = []
    for filename, data in files:
        pieces.append(
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
            "Content-Type: application/octet-stream\r\n\r\n".encode() + data + b"\r\n"
        )
    pieces.append(f"--{boundary}--\r\n".encode())
    return b"".join(pieces)


def _request(body: bytes, boundary: str = "batch", *, content_length: str | None = None) -> Request:
    messages = [{"type": "http.request", "body": body, "more_body": False}]

    async def receive():
        return messages.pop(0) if messages else {"type": "http.disconnect"}

    headers = [(b"content-type", f"multipart/form-data; boundary={boundary}".encode())]
    if content_length is not None:
        headers.append((b"content-length", content_length.encode()))
    return Request({"type": "http", "method": "POST", "headers": headers}, receive)


def test_inspection_and_read_support_31_worksheets_with_source_provenance(monkeypatch) -> None:
    data = _xlsx_sheets(31)
    inspection = inspect_source_file(data, "month.xlsx", DEFAULT_IMPORT_RESOURCE_POLICY)

    assert len(inspection["source_units"]) == 31
    assert inspection["source_units"][30]["name"] == "Day 31"
    assert inspection["source_units"][0]["inspection"]["header_cells"] == [
        "Driver Number", "Work Date", "Regular"
    ]
    selected = read_source_units(
        data,
        "month.xlsx",
        [
            {"unit_name": "Day 1", "header_row_number": 1, "date_order": None, "columns": [1, 3]},
            {"unit_name": "Day 31", "header_row_number": 1, "date_order": None, "columns": [1, 3]},
        ],
        DEFAULT_IMPORT_RESOURCE_POLICY,
    )

    assert [unit["unit_name"] for unit in selected] == ["Day 1", "Day 31"]
    assert selected[1]["rows"][0] == {
        "source_row": 2,
        "cells": {"1": FileCell("0031", "text"), "3": FileCell("31", "number")},
    }
    assert selected[1]["nonempty_columns"] == [1, 2, 3]


def test_selected_sheets_share_one_openpyxl_workbook(monkeypatch) -> None:
    import app.imports.files as file_reader

    original = file_reader.load_workbook
    calls = 0

    def counted(*args, **kwargs):
        nonlocal calls
        calls += 1
        return original(*args, **kwargs)

    monkeypatch.setattr(file_reader, "load_workbook", counted)
    units = read_source_units(
        _xlsx_sheets(31),
        "month.xlsx",
        [{"unit_name": f"Day {index}", "header_row_number": 1, "date_order": None, "columns": [1]}
         for index in range(1, 32)],
        DEFAULT_IMPORT_RESOURCE_POLICY,
    )
    assert len(units) == 31
    assert calls == 1


def test_unselected_oversized_worksheet_is_inspectable_but_selected_shape_is_refused() -> None:
    policy = replace(DEFAULT_IMPORT_RESOURCE_POLICY, max_rows_per_unit=2)
    data = _xlsx_sheets(2, first_rows=3)
    inspection = inspect_source_file(data, "month.xlsx", policy)
    assert inspection["source_units"][0]["inspection"]["row_count"] == 3
    successful = read_source_units(
        data, "month.xlsx",
        [{"unit_name": "Day 2", "header_row_number": 1, "date_order": None, "columns": [1]}],
        policy,
    )
    assert successful[0]["unit_name"] == "Day 2"

    with pytest.raises(ImportFailure) as error:
        read_source_units(
            data,
            "month.xlsx",
            [{"unit_name": "Day 1", "header_row_number": 1, "date_order": None, "columns": [1]}],
            policy,
        )
    assert error.value.code == "IMPORT_ROW_LIMIT_EXCEEDED"


def test_chart_sheets_do_not_count_as_selectable_data_units() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Driver Number", "Amount"])
    sheet.append(["0012", 4])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    workbook.create_chartsheet("Chart").add_chart(chart)
    output = io.BytesIO()
    workbook.save(output)

    result = inspect_source_file(output.getvalue(), "with-chart.xlsx", DEFAULT_IMPORT_RESOURCE_POLICY)
    assert [unit["name"] for unit in result["source_units"]] == ["Sheet"]


def test_csv_read_retains_only_selected_columns_and_physical_source_rows() -> None:
    data = b"Monthly payroll\r\nDriver Number,Work Date,Amount,Unmapped\r\n0012,2026-10-01,4.00,note\r\n0013,2026-10-02,,\r\n"
    units = read_source_units(
        data,
        "daily.csv",
        [{"unit_name": "CSV", "header_row_number": 2, "date_order": "YMD", "columns": [1, 3]}],
        DEFAULT_IMPORT_RESOURCE_POLICY,
    )
    assert units[0]["rows"] == [
        {"source_row": 3, "cells": {"1": FileCell("0012", "text"), "3": FileCell("4.00", "text")}},
        {"source_row": 4, "cells": {"1": FileCell("0013", "text"), "3": FileCell("", "text")}},
    ]
    assert units[0]["nonempty_columns"] == [1, 2, 3, 4]
    assert units[0]["inspection"]["header_row_number"] == 2
    assert units[0]["inspection"]["header_cells"] == ["Driver Number", "Work Date", "Amount", "Unmapped"]
    assert units[0]["inspection"]["date_order"] == "YMD"


def test_xlsx_nonempty_columns_include_unmapped_formula_without_cached_value() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Driver Number", "Unmapped formula"])
    sheet.append(["0012", "=1+1"])
    output = io.BytesIO()
    workbook.save(output)
    data = _rewrite_sheet(output.getvalue(), lambda xml: re.sub(
        rb'<c r="B2"[^>]*>.*?</c>', b'<c r="B2"><f>1+1</f><v></v></c>', xml
    ))
    units = read_source_units(
        data,
        "unmapped-formula.xlsx",
        [{"unit_name": "Sheet", "header_row_number": 1, "date_order": None, "columns": [1]}],
        DEFAULT_IMPORT_RESOURCE_POLICY,
    )
    assert units[0]["nonempty_columns"] == [1, 2]
    assert list(units[0]["rows"][0]["cells"]) == ["1"]


def test_xlsx_read_preserves_formula_cache_native_date_time_and_duration_kinds() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Driver Number", "Formula missing", "Formula cached", "Native date", "Clock", "Duration"])
    sheet.append(["0012", "=1+1", "=1+1", date(2026, 10, 1), time(10, 30), timedelta(hours=26)])
    sheet["E2"].number_format = "hh:mm"
    sheet["F2"].number_format = "[h]:mm"
    output = io.BytesIO()
    workbook.save(output)
    data = output.getvalue()

    def formulas_and_native_date(xml: bytes) -> bytes:
        xml = re.sub(
            rb'<c r="B2"[^>]*>.*?</c>',
            b'<c r="B2"><f>1+1</f><v></v></c>',
            xml,
        )
        xml = re.sub(
            rb'<c r="C2"[^>]*>.*?</c>',
            b'<c r="C2"><f>1+1</f><v>2</v></c>',
            xml,
        )
        xml = re.sub(
            rb'<c r="D2"[^>]*>.*?</c>',
            b'<c r="D2" t="d"><v>2026-10-01T00:00:00+02:00</v></c>',
            xml,
        )
        return xml

    data = _rewrite_sheet(data, formulas_and_native_date)
    units = read_source_units(
        data,
        "native.xlsx",
        [{"unit_name": "Sheet", "header_row_number": 1, "date_order": None,
          "columns": [2, 3, 4, 5, 6]}],
        DEFAULT_IMPORT_RESOURCE_POLICY,
    )
    cells = units[0]["rows"][0]["cells"]
    assert cells["2"] == FileCell(None, "text", True, True)
    assert cells["3"] == FileCell("2", "number", True, False)
    assert cells["4"] == FileCell("2026-10-01T00:00:00+02:00", "date", False, False)
    assert cells["5"] == FileCell("10:30:00", "time", False, False)
    assert cells["6"] == FileCell("1 day, 2:00:00", "duration", False, False)


def test_xlsx_numeric_lexemes_and_cached_results_retain_exact_decimal_validation():
    book = Workbook()
    book.active.append(["DriverNumber", "WorkDate", "Quantity", "Cached", "Extreme", "Exact"])
    book.active.append(["001", "2026-10-20", 1, "=1+1", 1, 4.25])
    output = io.BytesIO()
    book.save(output)
    book.close()
    def numeric_values(xml):
        for coordinate, value, formula in [
            (b"C2", b"1.0000000000000001", b""),
            (b"D2", b"3.0000000000000001", b"<f>1+1</f>"),
            (b"E2", b"1E1000000000", b""),
            (b"F2", b"4.2500", b""),
        ]:
            xml = re.sub(rb'<c r="' + coordinate + rb'"[^>]*>.*?</c>',
                b'<c r="' + coordinate + b'" t="n">' + formula + b'<v>' + value + b'</v></c>', xml)
        return xml
    data = _rewrite_sheet(output.getvalue(), numeric_values)
    parsed = read_source_units(data, "exact.xlsx", [{"unit_name": "Sheet", "header_row_number": 1,
        "date_order": None, "columns": [1, 2, 3, 4, 5, 6]}], DEFAULT_IMPORT_RESOURCE_POLICY)
    cells = parsed[0]["rows"][0]["cells"]
    definition = SimpleNamespace(input_type="Decimal")
    assert cells["3"].value == "1.0000000000000001"
    assert cells["4"] == FileCell("3.0000000000000001", "number", True, False)
    for column in ("3", "4", "5"):
        value, error = quantity(cells[column], definition)
        assert value is None and error is not None
    assert quantity(cells["6"], definition) == (Decimal("4.2500"), None)


def test_xlsx_refuses_duplicate_cell_coordinates_in_any_inspected_worksheet() -> None:
    data = _xlsx_sheets(1)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        original = archive.read("xl/worksheets/sheet1.xml")
    assert b'r="B2"' in original
    corrupted = _rewrite_sheet(data, lambda xml: xml.replace(b'r="B2"', b'r="A2"', 1))

    with pytest.raises(ImportFailure) as error:
        inspect_source_file(corrupted, "duplicate.xlsx", DEFAULT_IMPORT_RESOURCE_POLICY)
    assert error.value.code == "IMPORT_XLSX_INVALID"


def test_receive_files_accepts_31_parts_and_uses_actual_aggregate_counts() -> None:
    policy = replace(DEFAULT_IMPORT_RESOURCE_POLICY, max_files=31, max_file_bytes=4,
                     max_total_file_bytes=124, max_multipart_overhead_bytes=16_000)
    files = [(f"day-{index}.csv", f"{index:04d}".encode()) for index in range(31)]
    body = _multipart(files)
    received = asyncio.run(receive_files(_request(body, content_length="1"), policy))

    assert len(received) == 31
    assert received[30].data == b"0030"
    assert received[30].file_sha256 == hashlib.sha256(b"0030").hexdigest()
    assert [item.filename for item in received[:2]] == ["day-0.csv", "day-1.csv"]

    too_many_bytes = replace(policy, max_total_file_bytes=123)
    with pytest.raises(ImportFailure) as error:
        asyncio.run(receive_files(_request(body), too_many_bytes))
    assert error.value.code == "IMPORT_BATCH_TOO_LARGE"


def test_receive_files_rejects_extra_fields_empty_parts_and_truncated_batches() -> None:
    policy = replace(DEFAULT_IMPORT_RESOURCE_POLICY, max_files=2)
    boundary = "batch"
    extra_field = (
        b"--batch\r\nContent-Disposition: form-data; name=\"file\"; filename=\"a.csv\"\r\n\r\nx"
        b"\r\n--batch\r\nContent-Disposition: form-data; name=\"other\"\r\n\r\ny"
        b"\r\n--batch--\r\n"
    )
    with pytest.raises(ImportFailure, match="file"):
        asyncio.run(receive_files(_request(extra_field, boundary), policy))
    with pytest.raises(ImportFailure) as empty:
        asyncio.run(receive_files(_request(_multipart([("a.csv", b"")]), boundary), policy))
    assert empty.value.code == "IMPORT_FILE_EMPTY"
    truncated = _multipart([("a.csv", b"x")]).removesuffix(b"--\r\n")
    with pytest.raises(ImportFailure) as incomplete:
        asyncio.run(receive_files(_request(truncated), policy))
    assert incomplete.value.code == "IMPORT_UPLOAD_INCOMPLETE"


def test_policy_defaults_match_the_configured_operational_contract() -> None:
    policy = DEFAULT_IMPORT_RESOURCE_POLICY
    assert (policy.max_file_bytes, policy.max_total_file_bytes, policy.max_files) == (10_000_000, 50_000_000, 64)
    assert (policy.max_worksheets_per_file, policy.max_source_units, policy.max_layouts) == (128, 128, 64)
    assert (policy.max_rows_per_unit, policy.max_columns_per_unit) == (20_000, 128)
    assert (policy.max_cells_per_unit, policy.max_selected_cells, policy.max_data_rows) == (1_000_000, 2_000_000, 50_000)
    assert policy.max_candidates == 250_000


def test_settings_exposes_overridable_import_policy_fields() -> None:
    from app.config import Settings

    settings = Settings(
        DATABASE_URL="postgresql+asyncpg://user@localhost/test",
        SECRET_KEY="test-only",
        IMPORT_MAX_FILES=31,
        IMPORT_MAX_TOTAL_FILE_BYTES=12_000,
    )
    assert settings.import_resource_policy.max_files == 31
    assert settings.import_resource_policy.max_total_file_bytes == 12_000
