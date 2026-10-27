from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import replace

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table

from app.imports.canonical import (
    canonical_source_unit,
    decode_canonical_file,
    extract_canonical_file,
)
from app.imports.errors import ImportFailure
from app.imports.policy import DEFAULT_IMPORT_RESOURCE_POLICY
from app.product_policy import DEFAULT_PRODUCT_POLICY


def _xlsx(rows, sheets=1):
    book = Workbook()
    book.active.title = "Payroll"
    for row in rows:
        book.active.append(row)
    for i in range(1, sheets):
        book.create_sheet(f"Extra {i}")
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _rewrite(data, transform):
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content = transform(content)
            target.writestr(info.filename, content)
    return output.getvalue()


def _decode(result, **kwargs):
    return decode_canonical_file(result["canonical_payload"], expected_bytes=result["canonical_byte_count"],
        expected_sha256=result["canonical_sha256"], format_version=1,
        maximum_bytes=DEFAULT_PRODUCT_POLICY.max_import_canonical_file_bytes, **kwargs)


def test_csv_keeps_duplicate_headers_physical_rows_and_ignores_headerless_values():
    result = extract_canonical_file(b"DriverNumber,,DriverNumber\n0012,discard,0012\n\n0013,,0013\n", "daily.csv")
    value = _decode(result)
    assert [(h["column"], h["label"]) for h in value["header"]] == [(1, "DriverNumber"), (3, "DriverNumber")]
    assert [row["row"] for row in value["rows"]] == [2, 3, 4]
    assert result["inspection"]["headerless_columns"] == [2]
    assert result["inspection"]["notices"] and "discard" not in result["canonical_payload"].decode()
    baseline = extract_canonical_file(b"ID,Amount,\n001,2,\n", "tail.csv")
    tail = extract_canonical_file(b"ID,Amount,\n001,2,\n,,,unheaded tail\n", "tail.csv")
    assert tail["inspection"]["row_count"] == 3
    assert tail["inspection"]["imported_row_count"] == 1
    assert tail["canonical_payload"] == baseline["canonical_payload"]


def test_canonical_serialized_byte_limit_accepts_equality_and_rejects_one_byte_over():
    assert DEFAULT_PRODUCT_POLICY.max_import_canonical_file_bytes == 1_000_000
    data = b"DriverNumber,Amount\n0012,4.00\n"
    baseline = extract_canonical_file(data, "daily.csv")
    exact = replace(DEFAULT_PRODUCT_POLICY, max_import_canonical_file_bytes=baseline["canonical_byte_count"])
    assert len(extract_canonical_file(data, "daily.csv", product_policy=exact)["canonical_payload"]) == baseline["canonical_byte_count"]
    over = replace(DEFAULT_PRODUCT_POLICY, max_import_canonical_file_bytes=baseline["canonical_byte_count"] - 1)
    with pytest.raises(ImportFailure) as exc:
        extract_canonical_file(data, "daily.csv", product_policy=over)
    assert exc.value.code == "IMPORT_CANONICAL_FILE_TOO_LARGE"


def test_xlsx_extraction_uses_one_workbook_and_never_reads_legacy_source_matrix(monkeypatch):
    import app.imports.canonical as canonical
    import app.imports.files as old_files
    original, calls = canonical.load_workbook, 0
    def counted(*args, **kwargs):
        nonlocal calls
        calls += 1
        return original(*args, **kwargs)
    def no_matrix(*args, **kwargs):
        raise AssertionError("legacy source-matrix reader was called")
    monkeypatch.setattr(canonical, "load_workbook", counted)
    monkeypatch.setattr(old_files, "read_source_units", no_matrix)
    result = extract_canonical_file(_xlsx([["ID", "Amount"], ["001", 2]]), "daily.xlsx")
    assert calls == 1
    assert result["source_units"][0]["name"] == "Payroll"
    assert _decode(result)["rows"][0]["row"] == 2


def test_xlsx_rejects_multiple_sheets_and_style_only_extent_does_not_change_canonical_data():
    with pytest.raises(ImportFailure) as exc:
        extract_canonical_file(_xlsx([["ID"], ["001"]], sheets=2), "multi.xlsx")
    assert exc.value.code == "IMPORT_WORKBOOK_SHEET_COUNT_INVALID"
    rows = [["ID", "Amount"], ["001", 2]]
    plain = extract_canonical_file(_xlsx(rows), "styled.xlsx")
    book = Workbook()
    sheet = book.active
    sheet.title = "Payroll"
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(bold=True, color="FF0000")
    sheet["B1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    sheet["XFD1048576"].number_format = "0.000"
    sheet["XFD1048576"].font = Font(italic=True, color="0000FF")
    output = io.BytesIO()
    book.save(output)
    book.close()
    styled = extract_canonical_file(output.getvalue(), "styled.xlsx")
    assert [row["row"] for row in _decode(styled)["rows"]] == [2]
    assert styled["inspection"]["row_count"] == 2
    assert styled["inspection"]["column_count"] == 2
    assert styled["canonical_payload"] == plain["canonical_payload"]


def test_formula_cache_states_and_raw_numeric_lexemes_are_preserved():
    data = _xlsx([["ID", "Zero", "False", "Empty text", "Empty numeric", "Missing", "Exact"],
                  ["001", "=1", "=FALSE", '=""', "=1+1", "=2+2", 4]])
    def formulas(xml):
        entries = ((b"B2", b'<c r="B2"><f>1</f><v>0</v></c>'),
                   (b"C2", b'<c r="C2" t="b"><f>FALSE</f><v>0</v></c>'),
                   (b"D2", b'<c r="D2" t="str"><f>""</f><v></v></c>'),
                   (b"E2", b'<c r="E2" t="n"><f>1+1</f><v></v></c>'),
                   (b"F2", b'<c r="F2"><f>2+2</f></c>'),
                   (b"G2", b'<c r="G2" t="n"><v>4.2500</v></c>'))
        for coord, replacement in entries:
            xml = re.sub(rb'<c r="' + coord + rb'"[^>]*>.*?</c>', replacement, xml)
        return xml
    value = _decode(extract_canonical_file(_rewrite(data, formulas), "formula.xlsx"))
    cells = {item["column"]: item for item in value["rows"][0]["cells"]}
    assert cells[2]["value"] == {"kind": "number", "value": "0"} and cells[2]["formula"]["cache"]["usable"]
    assert cells[3]["value"] == {"kind": "boolean", "value": False} and cells[3]["formula"]["cache"]["usable"]
    assert cells[4]["value"] == {"kind": "text", "value": ""} and cells[4]["formula"]["cache"]["usable"]
    assert cells[5]["value"]["kind"] == cells[6]["value"]["kind"] == "unavailable"
    assert cells[5]["formula"]["cache"]["present"] and not cells[5]["formula"]["cache"]["usable"]
    assert not cells[6]["formula"]["cache"]["present"]
    assert cells[7]["value"] == {"kind": "number", "value": "4.2500"}


def test_current_decoder_column_and_integrity_bounds():
    result = extract_canonical_file(b"A,B\nx,y\n", "daily.csv")
    with pytest.raises(ImportFailure) as exc:
        _decode(result, maximum_columns=1)
    assert exc.value.code == "IMPORT_FILE_COLUMN_LIMIT_EXCEEDED"
    malformed = b'{"format_version":1,"format_version":1}'
    with pytest.raises(ImportFailure):
        decode_canonical_file(malformed, expected_bytes=len(malformed), expected_sha256=hashlib.sha256(malformed).hexdigest(),
                              format_version=1, maximum_bytes=1000)


def test_excel_serials_and_duration_kind_survive_projection():
    book = Workbook()
    sheet = book.active
    sheet.title = "Payroll"
    sheet.append(["ID", "Date", "Clock", "Duration"])
    sheet.append(["001", 1, 0.5, 1.125])
    sheet["B2"].number_format, sheet["C2"].number_format, sheet["D2"].number_format = "yyyy-mm-dd", "hh:mm:ss", "[h]:mm:ss"
    output = io.BytesIO()
    book.save(output)
    book.close()
    def exact(xml):
        for coord, raw in ((b"B2", b"1.0000000000000001"), (b"C2", b"0.5000000000000001"), (b"D2", b"1.1250000000000001")):
            xml = re.sub(rb'(<c r="' + coord + rb'"[^>]*>).*?(</c>)',
                         lambda m: m.group(1) + b"<v>" + raw + b"</v>" + m.group(2), xml)
        return xml
    value = _decode(extract_canonical_file(_rewrite(output.getvalue(), exact), "serial.xlsx"))
    cells = {cell["column"]: cell["value"] for cell in value["rows"][0]["cells"]}
    assert (cells[2]["kind"], cells[2]["excel_serial"]) == ("date", "1.0000000000000001")
    assert cells[2]["epoch"] == cells[3]["epoch"] == cells[4]["epoch"] == "1899-12-30T00:00:00"
    assert (cells[3]["kind"], cells[3]["excel_serial"]) == ("time", "0.5000000000000001")
    assert cells[4]["kind"] == "duration" and cells[4]["value"] == "1.1250000000000001"
    assert canonical_source_unit(value, unit_name="Payroll")["rows"][0]["cells"]["4"].kind == "duration"


def test_formula_header_metadata_and_hidden_column_are_retained():
    book = Workbook()
    sheet = book.active
    sheet.append(["ID", "Amount"])
    sheet.append(["001", 2])
    sheet.column_dimensions["B"].hidden = True
    output = io.BytesIO()
    book.save(output)
    book.close()
    def formula_header(xml):
        return re.sub(rb'<c r="A1"[^>]*>.*?</c>',
                      b'<c r="A1" t="str"><f>"Employee ID"</f><v>Employee ID</v></c>', xml)
    result = extract_canonical_file(_rewrite(output.getvalue(), formula_header), "hidden.xlsx")
    value = _decode(result)
    header = value["header"][0]
    assert header["label"] == "Employee ID"
    assert header["formula"]["text"] == '"Employee ID"'
    assert header["formula"]["cache"]["usable"] is True
    assert value["header"][1]["hidden"] is True
    assert result["inspection"]["hidden_columns"] == [2]


def test_xlsx_empty_physical_row_inside_data_extent_is_preserved():
    data = _xlsx([["ID", "Amount"], ["001", 2], ["002", 3], ["003", 4]])
    def insert_empty_row(xml):
        return re.sub(rb'<row r="3">.*?</row>', b'<row r="3"></row>', xml)
    value = _decode(extract_canonical_file(_rewrite(data, insert_empty_row), "blank-row.xlsx"))
    rows = value["rows"]
    assert [row["row"] for row in rows] == [2, 3, 4]
    assert [cell["value"]["kind"] for cell in rows[1]["cells"]] == ["blank", "blank"]


@pytest.mark.parametrize("shape", ["hidden", "chart"])
def test_hidden_or_chart_sheet_counts_as_an_extra_workbook_sheet(shape):
    book = Workbook()
    sheet = book.active
    sheet.append(["ID", "Amount"])
    sheet.append(["001", 2])
    if shape == "hidden":
        book.create_sheet("Hidden").sheet_state = "hidden"
    else:
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        book.create_chartsheet("Chart").add_chart(chart)
    output = io.BytesIO()
    book.save(output)
    book.close()
    with pytest.raises(ImportFailure) as exc:
        extract_canonical_file(output.getvalue(), "extra-sheet.xlsx")
    assert exc.value.code == "IMPORT_WORKBOOK_SHEET_COUNT_INVALID"


@pytest.mark.parametrize("kind", ["merged", "multiple_tables", "table_header_below_row_one"])
def test_explicit_workbook_table_shapes_that_conflict_with_row_one_are_rejected(kind):
    book = Workbook()
    sheet = book.active
    if kind == "merged":
        sheet.append(["ID", "Amount"])
        sheet.append(["001", 2])
        sheet.merge_cells("A1:B1")
    elif kind == "multiple_tables":
        sheet.append(["ID", "Amount", "Other ID", "Other Amount"])
        sheet.append(["001", 2, "002", 3])
        sheet.add_table(Table(displayName="FirstTable", ref="A1:B2"))
        sheet.add_table(Table(displayName="SecondTable", ref="C1:D2"))
    else:
        sheet.append(["Report title", None])
        sheet.append(["ID", "Amount"])
        sheet.append(["001", 2])
        sheet.add_table(Table(displayName="LaterHeader", ref="A2:B3"))
    output = io.BytesIO()
    book.save(output)
    book.close()
    with pytest.raises(ImportFailure) as exc:
        extract_canonical_file(output.getvalue(), "table.xlsx")
    assert exc.value.code in {"IMPORT_MERGED_CELLS_UNSUPPORTED", "IMPORT_MULTIPLE_TABLES_UNSUPPORTED",
                              "IMPORT_HEADER_ROW_INVALID"}


@pytest.mark.parametrize("cell_xml", [
    b'<c r="B2" t="z"><f>1+1</f><v>2</v></c>',
    b'<c r="B2" t="inlineStr"><f>1+1</f><is><t>2</t></is></c>',
])
def test_unknown_or_inline_string_formula_cache_types_are_rejected(cell_xml):
    data = _xlsx([["ID", "Value"], ["001", "=1+1"]])
    altered = _rewrite(data, lambda xml: re.sub(rb'<c r="B2"[^>]*>.*?</c>', cell_xml, xml))
    with pytest.raises(ImportFailure):
        extract_canonical_file(altered, "unsupported-cache.xlsx")


def test_error_valued_header_is_rejected_and_cached_error_data_keeps_error_type():
    data = _xlsx([["ID", "Value"], ["001", "=NA()"]])
    error_header = _rewrite(data, lambda xml: re.sub(rb'<c r="A1"[^>]*>.*?</c>',
        b'<c r="A1" t="e"><v>#N/A</v></c>', xml))
    with pytest.raises(ImportFailure) as exc:
        extract_canonical_file(error_header, "error-header.xlsx")
    assert exc.value.code == "IMPORT_HEADER_INVALID"
    error_value = _rewrite(data, lambda xml: re.sub(rb'<c r="B2"[^>]*>.*?</c>',
        b'<c r="B2" t="e"><f>NA()</f><v>#N/A</v></c>', xml))
    cell = _decode(extract_canonical_file(error_value, "error-cell.xlsx"))["rows"][0]["cells"][1]
    assert cell["value"] == {"kind": "error", "value": "#N/A"}
    assert cell["formula"]["cache"]["usable"] is True


def test_decoder_rechecks_current_row_policy_after_storage():
    result = extract_canonical_file(b"ID,Amount\n001,2\n002,3\n", "daily.csv")
    policy = replace(DEFAULT_IMPORT_RESOURCE_POLICY, max_rows_per_unit=2)
    with pytest.raises(ImportFailure):
        _decode(result, resource_policy=policy)


def test_excel_near_midnight_serial_is_not_rounded_into_next_work_date():
    book = Workbook()
    sheet = book.active
    sheet.append(["ID", "WorkDate"])
    sheet.append(["001", 1.999999999999999])
    sheet["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    output = io.BytesIO()
    book.save(output)
    book.close()
    def exact(xml):
        return re.sub(rb'(<c r="B2"[^>]*>).*?(</c>)',
                      lambda m: m.group(1) + b"<v>1.999999999999999999</v>" + m.group(2), xml)
    cell = _decode(extract_canonical_file(_rewrite(output.getvalue(), exact), "midnight.xlsx"))["rows"][0]["cells"][1]["value"]
    assert cell["kind"] == "date"
    assert cell["excel_serial"] == "1.999999999999999999"
    assert cell["epoch"] == "1899-12-30T00:00:00"
    assert cell["value"].startswith("1900-01-01T23:59:59")


def test_headerless_only_tail_rows_do_not_create_mapped_quantity_candidates():
    from app.imports.batch_service import _materialize_batch

    raw = b"DriverNumber,WorkDate,Amount,\n001,2026-10-20,1,\n,,,unheaded-tail\n"
    extracted = extract_canonical_file(raw, "daily.csv")
    assert extracted["inspection"]["row_count"] == 3
    assert extracted["inspection"]["imported_row_count"] == 1
    stored_file = {
        "filesourceid": 1, "fileordinal": 1, "filesizebytes": len(raw),
        "inspectionmetadata": {"decompressed_bytes": 0}, "contentpurgedatutc": None,
        "canonicalpayload": extracted["canonical_payload"],
        "canonicalbytecount": extracted["canonical_byte_count"],
        "canonicalsha256": extracted["canonical_sha256"], "canonicalformatversion": 1,
        "originalfilename": "daily.csv",
    }
    unit = {
        "sourceunitid": 1, "filesourceid": 1, "fileordinal": 1, "unitname": "CSV",
        "unitkind": "CSV", "sourcelayoutid": 1, "isselected": True,
        "headerrownumber": 1, "dateorder": None, "datesource": {},
        "inspectionmetadata": extracted["inspection"],
    }
    def selector(header):
        return {"header": header}
    layout = {"sourcelayoutid": 1, "structurefingerprint": "test",
              "inspectionmetadata": extracted["inspection"], "columnmapping": None}
    pending_rule = {"paydefinitionid": 8, "payrollperioddefinitionid": 8,
                    "rulepayload": {"payload_version": 1, "transform": "Direct",
                                    "inputs": [selector("Amount")],
                                    "shared_columns": {"DriverNumber": selector("DriverNumber"),
                                                       "WorkDate": selector("WorkDate")}}}
    policy = replace(DEFAULT_IMPORT_RESOURCE_POLICY, max_data_rows=1, max_candidates=1)
    rows = _materialize_batch([stored_file], [unit], [layout], policy, pending_rules=[pending_rule])
    assert len(rows) == 1 and rows[0]["source_unit_id"] == 1
