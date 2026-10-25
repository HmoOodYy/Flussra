from datetime import date, timedelta
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.imports.errors import ImportFailure
from app.imports.template import generate_template
from app.payroll.definition_calculation import PeriodDefinition


def _definition(index=1, *, code=None, name=None):
    return PeriodDefinition(
        payroll_period_definition_id=index,
        pay_definition_id=index + 100,
        rate_definition_id=index + 200,
        code=code or f"PAY{index:02d}",
        name=name or f"Pay {index:02d}",
        input_type="Number",
        unit="hours",
        calculation_method="PerUnit",
        calculation_method_version=1,
        rate_shape="Single",
        is_active=True,
    )


def _context(*, definitions=None, driver_count=1):
    start = date(2026, 10, 20)
    days = [start + timedelta(days=offset) for offset in range(7)]
    roster = [
        {
            "driverassignmentid": index + 1,
            "driverid": index + 1,
            "iseligibleforperiod": True,
            "eligibilityreasoncode": "EligibleByWindow",
            "eligibleworkdates": days,
            "drivernumber": f"{index + 1:05d}",
            "drivername": f"Driver {index + 1}",
        }
        for index in range(driver_count)
    ]
    return {
        "definitions": [_definition()] if definitions is None else definitions,
        "days": [{"workdate": workday} for workday in days],
        "lines": [],
        "entries": [],
        "roster": roster,
        "live_days": {(row["driverassignmentid"], day) for row in roster for day in days},
    }


def test_template_serializes_formula_like_driver_identifiers_as_literal_text():
    context = _context()
    context["roster"][0]["drivernumber"] = "=0012"
    context["roster"][0]["drivername"] = '=HYPERLINK("https://example.test","click")'
    data = generate_template(context)

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        cells = next(workbook["Payroll work"].iter_rows(min_row=2, max_row=2))
        assert [cell.value for cell in cells[:3]] == [
            "=0012", '=HYPERLINK("https://example.test","click")', "2026-10-20",
        ]
        assert all(cell.data_type == "s" and cell.number_format == "@" for cell in cells[:3])
    finally:
        workbook.close()
    with ZipFile(BytesIO(data)) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in sheet_xml and b"<f " not in sheet_xml


def test_template_refuses_period_without_operational_targets():
    with pytest.raises(ImportFailure) as raised:
        generate_template(_context(definitions=[]))
    assert raised.value.code == "IMPORT_TEMPLATE_NO_TARGETS"


@pytest.mark.parametrize("definitions", [
    [_definition(1, code="foo"), _definition(2, code="FOO")],
    [_definition(1, code="REG", name="Driver name")],
    [_definition(1, code="REG", name="Regular"),
     _definition(2, code="OT", name="Pay Item: REG")],
])
def test_template_refuses_generated_headers_that_cannot_map_exactly(definitions):
    with pytest.raises(ImportFailure) as raised:
        generate_template(_context(definitions=definitions))
    assert raised.value.code == "IMPORT_TEMPLATE_HEADERS_AMBIGUOUS"
    assert raised.value.status_code == 422


def test_template_refuses_more_than_fifty_columns_without_truncating():
    with pytest.raises(ImportFailure) as raised:
        generate_template(_context(definitions=[_definition(index) for index in range(1, 49)]))
    assert raised.value.code == "IMPORT_TEMPLATE_LIMIT"


def test_template_refuses_more_than_twenty_thousand_rows_without_truncating():
    with pytest.raises(ImportFailure) as raised:
        generate_template(_context(driver_count=2_858))
    assert raised.value.code == "IMPORT_TEMPLATE_LIMIT"
