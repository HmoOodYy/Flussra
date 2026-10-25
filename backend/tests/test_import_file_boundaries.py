"""Inclusive product limits, actual worksheet content and safe package handling."""
import asyncio
import io
import re
import struct
import warnings
import zipfile

import pytest
from openpyxl import Workbook

from app.imports import ImportFailure, inspect_file, receive_file
from tests.test_import_file_inspection import _rewrite_xlsx, _workbook_bytes
from tests.test_import_upload_transport import _multipart, _request


def test_ten_megabyte_csv_is_usable_and_one_more_byte_is_refused():
    row = b",".join([b"x" * 1999] * 50) + b"\n"
    data = row * 100
    assert len(data) == 10_000_000
    result = inspect_file(data, "exact.csv")
    assert (result["row_count"], result["column_count"]) == (100, 50)
    with pytest.raises(ImportFailure) as exc:
        inspect_file(data + b"x", "over.csv")
    assert exc.value.code == "IMPORT_FILE_TOO_LARGE"


def test_csv_rectangle_at_one_million_cells_and_beyond_row_and_column_bounds():
    row = b",".join([b"1"] * 50) + b"\n"
    result = inspect_file(row * 20_000, "dense.csv")
    assert (result["row_count"], result["column_count"], result["cell_count"]) == (20_000, 50, 1_000_000)
    # With both dimension caps enforced, exceeding a million necessarily
    # exceeds a dimension first; there is no independent >1M valid rectangle.
    for data, code in ((row * 20_001, "IMPORT_ROW_LIMIT_EXCEEDED"),
                       (b",".join([b"1"] * 51), "IMPORT_COLUMN_LIMIT_EXCEEDED")):
        with pytest.raises(ImportFailure) as exc:
            inspect_file(data, "over.csv")
        assert exc.value.code == code


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_cell_limit_measures_utf8_bytes_including_ignored_text(extension):
    def source(text):
        return text.encode("utf-8") if extension == "csv" else _workbook_bytes([["Driver Number", text]])

    exact = "\U0001f600" * 2000
    assert len(exact.encode()) == 8000
    assert inspect_file(source(exact), f"cell.{extension}")["column_count"] >= 1
    with pytest.raises(ImportFailure) as exc:
        inspect_file(source(exact + "a"), f"cell.{extension}")
    assert exc.value.code == "IMPORT_CELL_TEXT_LIMIT_EXCEEDED"


def test_xlsx_fifty_megabytes_actual_uncompressed_is_accepted_and_one_more_refused():
    base = _workbook_bytes([["Driver Number", "Date"]])
    with zipfile.ZipFile(io.BytesIO(base)) as archive:
        size = sum(info.file_size for info in archive.infolist())
    padding = b"0" * (50_000_000 - size)
    exact = _rewrite_xlsx(base, {}, {"custom/padding.bin": padding})
    with zipfile.ZipFile(io.BytesIO(exact)) as archive:
        assert sum(info.file_size for info in archive.infolist()) == 50_000_000
    assert inspect_file(exact, "exact.xlsx")["row_count"] == 1
    with pytest.raises(ImportFailure) as exc:
        inspect_file(_rewrite_xlsx(base, {}, {"custom/padding.bin": padding + b"1"}), "over.xlsx")
    assert exc.value.code == "IMPORT_XLSX_DECOMPRESSED_LIMIT_EXCEEDED"


def _sparse_workbook(row, column):
    wb = Workbook()
    wb.active.title = "Small"
    wb.active.append(["Driver Number", "Date"])
    second = wb.create_sheet("Large")
    second.cell(row, column, "x")
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def test_actual_selected_extents_count_gaps_and_ignore_declared_dimensions():
    data = _sparse_workbook(20_000, 50)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("xl/worksheets/sheet2.xml")
    rewritten, count = re.subn(rb'<dimension ref="[^"]+"\s*/>', b'<dimension ref="A1"/>', xml)
    assert count == 1
    result = inspect_file(_rewrite_xlsx(data, {"xl/worksheets/sheet2.xml": rewritten}),
                          "sparse.xlsx", sheet_name="Large")
    assert (result["row_count"], result["column_count"], result["cell_count"]) == (20_000, 50, 1_000_000)


@pytest.mark.parametrize(("row", "column", "code"), [(20_001, 1, "IMPORT_ROW_LIMIT_EXCEEDED"), (1, 51, "IMPORT_COLUMN_LIMIT_EXCEEDED")])
def test_unselected_sheet_shape_does_not_reduce_the_selected_sheet_capacity(row, column, code):
    data = _sparse_workbook(row, column)
    assert inspect_file(data, "scope.xlsx", sheet_name="Small")["row_count"] == 1
    with pytest.raises(ImportFailure) as exc:
        inspect_file(data, "scope.xlsx", sheet_name="Large")
    assert exc.value.code == code


def test_unselected_sheet_text_is_still_checked():
    data = _workbook_bytes([["Header"]])
    wb = Workbook()
    wb.active.append(["Header"])
    wb.create_sheet("Ignored").cell(1, 1, "x" * 8001)
    output = io.BytesIO()
    wb.save(output)
    with pytest.raises(ImportFailure) as exc:
        inspect_file(output.getvalue(), "ignored.xlsx")
    assert exc.value.code == "IMPORT_CELL_TEXT_LIMIT_EXCEEDED"
    assert inspect_file(data, "small.xlsx")["row_count"] == 1


def test_hidden_sheets_count_toward_ten_sheet_cap():
    wb = Workbook()
    wb.active.append(["Header"])
    for index in range(9):
        wb.create_sheet(f"Hidden{index}").sheet_state = "hidden"
    output = io.BytesIO()
    wb.save(output)
    assert len(inspect_file(output.getvalue(), "ten.xlsx")["sheets"]) == 10
    wb.create_sheet("Eleventh").sheet_state = "hidden"
    output = io.BytesIO()
    wb.save(output)
    with pytest.raises(ImportFailure) as exc:
        inspect_file(output.getvalue(), "eleven.xlsx")
    assert exc.value.code == "IMPORT_SHEET_LIMIT_EXCEEDED"


def test_formula_cache_is_used_for_header_without_recalculation():
    data = _workbook_bytes([["=1+1", "00123"]])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
    xml, count = re.subn(rb'(<f>1\+1</f>)<v\s*/>', rb'\1<v>999</v>', xml)
    assert count == 1
    result = inspect_file(_rewrite_xlsx(data, {"xl/worksheets/sheet1.xml": xml}), "cache.xlsx", header_row_number=1)
    assert result["header_cells"] == ["999", "00123"]
    assert result["formula_cached_value_count"] == 1


def test_duplicate_encrypted_and_malformed_zip_are_refused():
    data = _workbook_bytes([["Header"]])
    output = io.BytesIO(data)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        with zipfile.ZipFile(output, "a") as archive:
            archive.writestr("xl/workbook.xml", b"<bad/>")
    with pytest.raises(ImportFailure) as exc:
        inspect_file(output.getvalue(), "duplicate.xlsx")
    assert exc.value.code == "IMPORT_XLSX_DUPLICATE_MEMBER"
    encrypted = bytearray(data)
    start = encrypted.index(b"PK\x01\x02")
    flags = struct.unpack_from("<H", encrypted, start + 8)[0]
    struct.pack_into("<H", encrypted, start + 8, flags | 1)
    with pytest.raises(ImportFailure) as exc:
        inspect_file(bytes(encrypted), "encrypted.xlsx")
    assert exc.value.code == "IMPORT_XLSX_ENCRYPTED"
    with pytest.raises(ImportFailure):
        inspect_file(data[:100], "truncated.xlsx")


def test_csv_bom_tabs_explicit_header_and_malformed_quoting():
    result = inspect_file(b"\xef\xbb\xbfNotes\n001\tDate\n", "bom.csv", header_row_number=2)
    assert result["encoding"] == "utf-8-sig" and result["delimiter"] == "\t"
    assert result["header_cells"] == ["001", "Date"]
    with pytest.raises(ImportFailure) as exc:
        inspect_file(b'"unterminated', "bad.csv")
    assert exc.value.code == "IMPORT_CSV_MALFORMED"


def test_header_fallback_requires_text_and_preserves_unicode_filename():
    result = inspect_file(_workbook_bytes([[1, 2], ["Notes", "Other"]]), "numbers.xlsx")
    assert result["header_row_number"] == 2
    assert result["header_cells"] == ["Notes", "Other"]
    filename = "\u0628\u064a\u0627\u0646\u0627\u062a.csv"
    request = _request(_multipart(b"Header\n001\n", filename=filename))
    assert asyncio.run(receive_file(request)).filename == filename


@pytest.mark.parametrize("data", [b"%PDF-1.7\n", b"GIF89a", b"\x01binary", b"PK\x03\x04workbook"])
def test_binary_signatures_are_not_text_csv(data):
    with pytest.raises(ImportFailure) as exc:
        inspect_file(data, "renamed.csv")
    assert exc.value.code == "IMPORT_CSV_BINARY_CONTENT"


def test_consumed_bytes_ignore_false_content_length_and_disconnect_cleans_memory():
    data = b"Header\n001\n"
    request = _request(_multipart(data), chunk_sizes=(1, 3, 5, 7))
    request.scope["headers"].append((b"content-length", b"1"))
    assert asyncio.run(receive_file(request)).data == data
    # No spool/temp file exists: invocation-owned buffers become unreachable
    # on rejection; a subsequent invocation remains independent and succeeds.
    with pytest.raises(ImportFailure):
        asyncio.run(receive_file(_request(_multipart(data)[:-15])))
    assert asyncio.run(receive_file(_request(_multipart(data)))).data == data
