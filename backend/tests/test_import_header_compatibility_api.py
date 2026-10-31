from __future__ import annotations

from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from sqlalchemy import text

from tests.compensation_api_fixtures import (  # noqa: F401 - register isolated API fixtures
    p3c_application,
    p3c_database_engine,
    p3c_http_client,
)
from tests.compensation_database_fixtures import p3b_database  # noqa: F401
from tests.payroll_runtime_fixtures import build_payroll_tenant, open_period_with_grid

pytestmark = pytest.mark.asyncio


@pytest.fixture(name="tenant")
def payroll_tenant(p3b_dsn):
    return build_payroll_tenant(p3b_dsn)


def _multipart(files):
    boundary = "import-header-boundary"
    parts = []
    for filename, data in files:
        parts.extend((
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
            "Content-Type: application/octet-stream\r\n\r\n".encode(), data, b"\r\n"))
    parts.append(f"--{boundary}--\r\n".encode())
    return boundary, b"".join(parts)


async def _upload(client, tenant, period, files):
    boundary, body = _multipart(files)
    return await client.post(
        f"/payroll/periods/{period['payroll_period_id']}/imports",
        content=body,
        headers={**tenant.admin, "Content-Type": f"multipart/form-data; boundary={boundary}",
                 "Idempotency-Key": str(uuid4())},
    )


def _xlsx(headers, row):
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def _xlsx_with_hidden_and_styled_extent(headers, row):
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    sheet.append(row)
    sheet.column_dimensions["D"].hidden = True
    sheet["Z100"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


async def _period_counts(engine, period_id, company_id):
    async with engine.connect() as db:
        return tuple((await db.execute(text("""
            SELECT
                (SELECT count(*) FROM import.PayrollWorkImports WHERE PayrollPeriodID=:pid),
                (SELECT count(*) FROM import.FileSources f JOIN import.PayrollWorkImports w
                    ON w.ImportID=f.ImportID AND w.CompanyID=f.CompanyID WHERE w.PayrollPeriodID=:pid),
                (SELECT count(*) FROM import.FileCanonicalData c JOIN import.PayrollWorkImports w
                    ON w.ImportID=c.ImportID AND w.CompanyID=c.CompanyID WHERE w.PayrollPeriodID=:pid),
                (SELECT count(*) FROM payroll.PayrollDraftLines WHERE PayrollPeriodID=:pid),
                (SELECT count(*) FROM payroll.PayrollPeriodAuditEvidenceEvents WHERE PayrollPeriodID=:pid),
                (SELECT count(*) FROM import.PayrollWorkRuleVersions WHERE CompanyID=:cid),
                (SELECT count(*) FROM import.PayrollWorkPendingRules WHERE CompanyID=:cid),
                (SELECT count(*) FROM import.PayrollWorkRuleHeads WHERE CompanyID=:cid)
        """), {"pid": period_id, "cid": company_id})).one())


async def test_partial_overlap_reports_combined_max_multiplicity_and_duplicate_provenance(
    p3c_client, p3c_engine, tenant,
):
    period, _grid = await open_period_with_grid(p3c_client, p3c_engine, tenant)
    first = b"DriverNumber,Hours,Hours\n001,1,2\n"
    second = b" drivernumber ,Hours,Bonus\n001,3,4\n"
    response = await _upload(p3c_client, tenant, period, [("first.csv", first), ("second.csv", second)])

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["header_structure"]["available"] is True
    assert body["header_structure"]["combined_column_count"] == 4
    assert [item["count"] for item in body["header_structure"]["headed_column_counts"]] == [3, 3]
    assert len(body["layouts"]) == 2
    first_layout = next(item for item in body["layouts"] if item["inspection"]["duplicate_header_signatures"])
    duplicate = first_layout["inspection"]["duplicate_header_signatures"][0]
    assert set(duplicate) == {"signature", "count", "columns"}
    assert duplicate["count"] == 2
    assert duplicate["columns"] == [2, 3]
    page = await p3c_client.get(
        f"/payroll/imports/{body['import_id']}?header_layout_id={first_layout['layout_id']}&header_limit=3",
        headers=tenant.admin)
    assert page.status_code == 200, page.text
    items = page.json()["header_page"]["items"]
    assert [item["source_column"] for item in items] == [1, 2, 3]
    assert [item["label"] for item in items] == ["DriverNumber", "Hours", "Hours"]
    assert items[1]["occurrence_id"] != items[2]["occurrence_id"]
    assert items[1]["duplicate_signature"] == items[2]["duplicate_signature"]


async def test_header_page_exposes_server_normalized_label_for_unique_and_duplicate_headers(
    p3c_client, p3c_engine, tenant,
):
    period, _grid = await open_period_with_grid(p3c_client, p3c_engine, tenant)
    response = await _upload(p3c_client, tenant, period, [
        ("unicode-a.csv", "Straße,Hours,Hours\n001,1,2\n".encode()),
        ("unicode-b.csv", b"STRASSE,Hours,Bonus\n001,3,4\n"),
    ])
    assert response.status_code == 200, response.text
    body = response.json()
    assert len(body["layouts"]) == 2
    pages = []
    for layout in body["layouts"]:
        page = await p3c_client.get(
            f"/payroll/imports/{body['import_id']}?header_layout_id={layout['layout_id']}&header_limit=3",
            headers=tenant.admin)
        assert page.status_code == 200, page.text
        pages.append(page.json()["header_page"]["items"])
    assert pages[0][0]["normalized_label"] == "strasse"
    assert pages[1][0]["normalized_label"] == "strasse"
    assert pages[0][1]["normalized_label"] == pages[0][2]["normalized_label"] == "hours"
    assert pages[0][1]["duplicate_signature"] == pages[0][2]["duplicate_signature"]


async def test_mixed_csv_xlsx_counts_headed_hidden_columns_but_not_headerless_or_styled_extent(
    p3c_client, p3c_engine, tenant,
):
    period, _grid = await open_period_with_grid(p3c_client, p3c_engine, tenant)
    headers = ["DriverNumber", "", "WorkDate", "Region", "Hours"]
    workbook = _xlsx_with_hidden_and_styled_extent(
        headers, ["001", "hidden input", "2026-10-20", "East", 2])
    csv = b"DriverNumber,,WorkDate,Region,Hours\n002,unheaded,2026-10-21,West,3\n"

    response = await _upload(p3c_client, tenant, period, [("day.xlsx", workbook), ("day.csv", csv)])

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["header_structure"]["combined_column_count"] == 4
    assert len(body["layouts"]) == 1
    units = [file["units"][0] for file in body["files"]]
    assert all(unit["inspection"]["headed_column_count"] == 4 for unit in units)
    assert units[0]["inspection"]["hidden_columns"] == [4]
    assert units[1]["inspection"]["headerless_columns"] == [2]
    assert body["layouts"][0]["inspection"]["header_count"] == 5


@pytest.mark.parametrize(("combined_width", "configured_limit", "expected_status"), [
    (40, 40, 200), (41, 40, 422), (40, 39, 422), (41, 41, 200),
])
async def test_combined_header_allowance_exact_boundary_and_runtime_overrides(
    p3c_client, p3c_app, p3c_engine, tenant, monkeypatch, combined_width, configured_limit, expected_status,
):
    period, _grid = await open_period_with_grid(p3c_client, p3c_engine, tenant)
    monkeypatch.setattr(p3c_app.state.settings, "IMPORT_MAX_COMBINED_COLUMNS", configured_limit)
    left_count = (combined_width - 1) // 2
    right_count = combined_width - 1 - left_count
    left = ["DriverNumber", *[f"Left{index}" for index in range(left_count)]]
    right = ["DriverNumber", *[f"Right{index}" for index in range(right_count)]]
    files = [("left.csv", (",".join(left) + "\n" + ",".join(["1"] * len(left)) + "\n").encode()),
             ("right.csv", (",".join(right) + "\n" + ",".join(["1"] * len(right)) + "\n").encode())]

    response = await _upload(p3c_client, tenant, period, files)

    assert response.status_code == expected_status, response.text
    if expected_status == 200:
        assert response.json()["header_structure"]["combined_column_count"] == combined_width
    else:
        assert response.json()["detail"]["code"] == "IMPORT_COMBINED_COLUMN_LIMIT_EXCEEDED"


async def test_pairwise_disjoint_headers_reject_before_row_extraction_or_staging(
    p3c_client, p3c_engine, tenant, monkeypatch,
):
    import app.imports.batch_service as batch_service

    period, _grid = await open_period_with_grid(p3c_client, p3c_engine, tenant)
    before = await _period_counts(p3c_engine, period["payroll_period_id"], tenant.company_id)

    def forbidden_extract(*_args, **_kwargs):
        raise AssertionError("canonical row extraction ran before header compatibility passed")

    monkeypatch.setattr(batch_service, "extract_canonical_file", forbidden_extract)
    response = await _upload(p3c_client, tenant, period, [
        ("one.csv", b"DriverNumber\n001\n"),
        ("bridge.csv", b"DriverNumber,Hours\n001,2\n"),
        ("three.csv", b"Hours\n2\n"),
    ])

    after = await _period_counts(p3c_engine, period["payroll_period_id"], tenant.company_id)
    assert response.status_code == 422, response.text
    assert response.json()["detail"]["code"] == "IMPORT_HEADER_STRUCTURE_INCOMPATIBLE"
    assert after == before


async def test_distinct_layout_resource_cap_is_checked_before_row_extraction(
    p3c_client, p3c_app, p3c_engine, tenant, monkeypatch,
):
    import app.imports.batch_service as batch_service

    period, _grid = await open_period_with_grid(p3c_client, p3c_engine, tenant)
    monkeypatch.setattr(p3c_app.state.settings, "IMPORT_MAX_LAYOUTS", 1)
    before = await _period_counts(p3c_engine, period["payroll_period_id"], tenant.company_id)
    def forbidden_extract(*_args, **_kwargs):
        raise AssertionError("canonical row extraction ran before the layout resource cap")

    monkeypatch.setattr(batch_service, "extract_canonical_file", forbidden_extract)
    response = await _upload(p3c_client, tenant, period, [
        ("first.csv", b"DriverNumber,Hours\n001,1\n"),
        ("second.csv", b"DriverNumber,Bonus\n001,2\n"),
    ])
    after = await _period_counts(p3c_engine, period["payroll_period_id"], tenant.company_id)

    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "IMPORT_LAYOUT_LIMIT_EXCEEDED"
    assert after == before


async def test_lowered_policy_blocks_mapping_and_preview_but_metadata_keeps_truthful_count(
    p3c_client, p3c_app, p3c_engine, tenant, monkeypatch,
):
    import app.imports.batch_service as batch_service

    period, _grid = await open_period_with_grid(p3c_client, p3c_engine, tenant)
    response = await _upload(p3c_client, tenant, period, [
        ("partial-a.csv", b"DriverNumber,WorkDate,Hours\n001,2026-10-20,1\n"),
        ("partial-b.csv", b"DriverNumber,WorkDate,Bonus\n001,2026-10-20,2\n"),
    ])
    assert response.status_code == 200, response.text
    body = response.json()
    monkeypatch.setattr(p3c_app.state.settings, "IMPORT_MAX_COMBINED_COLUMNS", 3)

    metadata = await p3c_client.get(f"/payroll/imports/{body['import_id']}", headers=tenant.admin)
    assert metadata.status_code == 200, metadata.text
    assert metadata.json()["header_structure"]["combined_column_count"] == 4

    mapping = await p3c_client.put(
        f"/payroll/imports/{body['import_id']}/mapping",
        json={"expected_revision": body["revision"], "shared_columns": {}, "targets": []}, headers=tenant.admin)
    assert mapping.status_code == 422
    assert mapping.json()["detail"]["code"] == "IMPORT_COMBINED_COLUMN_LIMIT_EXCEEDED"

    def forbidden_decode(*_args, **_kwargs):
        raise AssertionError("preview decoded canonical rows before checking the current header allowance")

    monkeypatch.setattr(batch_service, "decode_canonical_file", forbidden_decode)
    preview = await p3c_client.post(f"/payroll/imports/{body['import_id']}/preview",
        json={"expected_revision": body["revision"]}, headers=tenant.admin)
    assert preview.status_code == 422
    assert preview.json()["detail"]["code"] == "IMPORT_COMBINED_COLUMN_LIMIT_EXCEEDED"


async def test_calendar_looking_labels_are_accepted_as_headers_not_day_targets(
    p3c_client, p3c_engine, tenant,
):
    period, _grid = await open_period_with_grid(p3c_client, p3c_engine, tenant)
    headers = ["DriverNumber", "Monday", "Tuesday", "Jan", "Feb", "2026-10-20"]
    payload = (",".join(headers) + "\n" + ",".join(["001"] * len(headers)) + "\n").encode()
    response = await _upload(p3c_client, tenant, period, [("calendar-labels.csv", payload)])

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["header_structure"]["combined_column_count"] == len(headers)
    from app.imports.mapping import targets
    from app.payroll.period_definitions import list_period_definitions

    async with p3c_engine.connect() as db:
        definitions = await list_period_definitions(period["payroll_period_id"], tenant.company_id, db, active_only=True)
    quantity_target = next(item["target"] for item in targets(definitions)
                           if item["target"].startswith("Definition:"))
    ordinary_quantity = await p3c_client.put(
        f"/payroll/imports/{body['import_id']}/mapping",
        json={"expected_revision": body["revision"], "shared_columns": {
            "DriverNumber": {"header": "DriverNumber"}}, "targets": [{
                "target": quantity_target, "transform": "Direct", "inputs": [{"header": "Monday"}],
                "expected_head_revision": 0}]}, headers=tenant.admin)
    assert ordinary_quantity.status_code == 200, ordinary_quantity.text
    invalid_day_target = await p3c_client.put(
        f"/payroll/imports/{body['import_id']}/mapping",
        json={"expected_revision": ordinary_quantity.json()["revision"], "shared_columns": {
            "DriverNumber": {"header": "DriverNumber"}}, "targets": [{
                "target": "Monday", "transform": "Direct", "inputs": [{"header": "Tuesday"}],
                "expected_head_revision": 0}]}, headers=tenant.admin)
    assert invalid_day_target.status_code == 422
    assert invalid_day_target.json()["detail"]["code"] == "IMPORT_MAPPING_REQUEST_INVALID"
    current = await p3c_client.get(f"/payroll/imports/{body['import_id']}", headers=tenant.admin)
    assert current.status_code == 200, current.text
    assert current.json()["revision"] == ordinary_quantity.json()["revision"]
    current = await p3c_client.get(f"/payroll/imports/{body['import_id']}", headers=tenant.admin)
    assert current.json()["revision"] == ordinary_quantity.json()["revision"]


async def test_regroup_endpoint_refuses_without_changing_source_setup(
    p3c_client, p3c_engine, tenant,
):
    period, _grid = await open_period_with_grid(p3c_client, p3c_engine, tenant)
    response = await _upload(p3c_client, tenant, period, [
        ("daily.csv", b"DriverNumber,WorkDate,Hours\n001,2026-10-20,1\n")])
    assert response.status_code == 200, response.text
    body = response.json()
    source_unit_id = body["files"][0]["units"][0]["source_unit_id"]
    path = f"/payroll/imports/{body['import_id']}/layouts"

    stale = await p3c_client.post(path, json={"expected_revision": body["revision"] + 1,
        "source_unit_ids": [source_unit_id], "label": "legacy regroup"}, headers=tenant.admin)
    assert stale.status_code == 409
    rejected = await p3c_client.post(path, json={"expected_revision": body["revision"],
        "source_unit_ids": [source_unit_id], "label": "legacy regroup"}, headers=tenant.admin)
    assert rejected.status_code == 422
    assert rejected.json()["detail"]["code"] == "IMPORT_LAYOUT_REGROUP_UNSUPPORTED"
    current = await p3c_client.get(f"/payroll/imports/{body['import_id']}", headers=tenant.admin)
    assert current.status_code == 200
    assert current.json()["revision"] == body["revision"]
    assert len(current.json()["layouts"]) == 1
